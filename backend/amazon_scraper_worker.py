"""Short-lived Amazon fetch/export worker.

The process intentionally does no work until its parent sends one newline-delimited
message over stdin. The parent first assigns it to a Windows Job Object, so there is
no network or Excel activity before process ownership is established. Secrets such
as proxy credentials are only present in that private stdin message.
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path
from typing import Any

if __name__ == "__main__":
    # Mark only the short-lived child process; importing this module in the app/test harness must
    # not switch the parent process into worker mode.
    os.environ["AMAZON_SCRAPER_WORKER"] = "1"


def _fetch(request: dict[str, Any]) -> dict[str, Any]:
    from .external_effects import assert_allowed
    assert_allowed('amazon_collection')
    from .amazon_scraper import _fetch_product

    result = _fetch_product(request["item"], request.get("proxy"), _NeverCancelled())
    images: list[list[str]] = []
    image_dir = Path(request["scratch"])
    image_dir.mkdir(parents=True, exist_ok=True)
    if result.get("status") == "ready":
        for index, (data, mime) in enumerate(result.get("images", []), 1):
            if index > 10:
                break
            name = f"image-{index:02d}.bin"
            (image_dir / name).write_bytes(data)
            images.append([name, mime])
        result["images"] = []
    return {"result": result, "images": images}


class _NeverCancelled:
    def is_set(self) -> bool:
        return False


def _export(request: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.utils import get_column_letter

    root = Path(request["root"])
    job = request["job"]
    job_dir = root / job["id"]
    ready = [item for item in job["items"] if item.get("status") == "ready"]
    max_images = max((int(item.get("image_count", 0)) for item in ready), default=0)
    wb = Workbook()
    ws = wb.active
    ws.title = "Amazon Listing 图片"
    ws.append(["ASIN", "商品标题"] + [f"图片 {index}" for index in range(1, max_images + 1)])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 48
    for column in range(3, 3 + max_images):
        ws.column_dimensions[get_column_letter(column)].width = 22
    for row_number, item in enumerate(ready, 2):
        ws.cell(row_number, 1, item["asin"])
        ws.cell(row_number, 2, item.get("title", ""))
        ws.row_dimensions[row_number].height = 112
        folder = job_dir / "images"
        paths = sorted(folder.glob(f"{item['asin']}_*"))
        for image_index, path in enumerate(paths[:10], 3):
            image = ExcelImage(path)
            image.width = 144
            image.height = 144
            ws.add_image(image, f"{get_column_letter(image_index)}{row_number}")
    target = Path(request["target"])
    wb.save(target)
    wb.close()
    return {"ok": True}


def main() -> int:
    line = sys.stdin.readline()
    if not line:
        return 2
    try:
        request = json.loads(line)
        if request.get("action") == "fetch":
            response = _fetch(request)
        elif request.get("action") == "export":
            response = _export(request)
        else:
            response = {"error": "unsupported worker action"}
        sys.stdout.write(json.dumps(response, ensure_ascii=False, separators=(",", ":")) + "\n")
        sys.stdout.flush()
        return 0
    except Exception:
        # Keep diagnostics private; exceptions can include URLs and transport details.
        sys.stdout.write(json.dumps({"error": "worker operation failed"}) + "\n")
        sys.stdout.flush()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
