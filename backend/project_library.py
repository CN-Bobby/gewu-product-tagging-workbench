"""Read-only previews for the small set of finished project deliverables."""
from __future__ import annotations

import hashlib
import io
import re
import threading
import zipfile
from collections import OrderedDict
from urllib.parse import quote
from xml.etree import ElementTree
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image
from starlette.responses import Response

from . import project_files
from .db import ProjectArtifact, Session
from .importer import wps_images

FOLDER_ORDER = ("source", "data_x", "reviews", "voc", "keywords", "other")
FOLDER_LABELS = {"source": "原始上传数据", "data_x": "打标结果",
                 "reviews": "评论文件", "voc": "VOC 报告",
                 "keywords": "关键词研究", "other": "其他资源"}
WORKBOOK_TYPES = {
    "source": {"dataset"},
    "data_x": {"data_x_snapshot"},
    "reviews": {"amazon_reviews", "sorftime_mcp_review_export"},
}
MAX_WORKBOOK_BYTES = 100 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 500 * 1024 * 1024
MAX_ROWS = 100_001
MAX_COLUMNS = 500
MAX_HTML_BYTES = 20 * 1024 * 1024
_DISPIMG = re.compile(r'DISPIMG\(\s*"([^"]+)"', re.IGNORECASE)
_IMAGE_CACHE_MAX_ENTRIES = 4
_IMAGE_CACHE_MAX_BYTES = 32 * 1024 * 1024
_image_cache: OrderedDict[tuple[str, str, str], tuple[dict, dict[str, bytes], int]] = OrderedDict()
_image_cache_bytes = 0
_image_cache_lock = threading.RLock()


def _extension(item: dict[str, Any]) -> str:
    return Path(str(item.get("name") or "")).suffix.casefold()


def _folder_for(item: dict[str, Any]) -> str | None:
    kind = item.get("kind")
    source_type = item.get("source_type")
    ext = _extension(item)
    if kind in WORKBOOK_TYPES and source_type in WORKBOOK_TYPES[kind] and ext == ".xlsx":
        if kind == "reviews" and source_type == "amazon_reviews":
            source_id = str(item.get("source_id") or "")
            if not source_id.endswith(":result.xlsx"):
                return None
        return str(kind)
    if kind == "voc" and source_type == "amazon_voc_output" and ext == ".html":
        return "voc"
    if kind == "keywords" and source_type == "keyword_research" and ext == ".xlsx":
        return "keywords"
    # Direct project uploads are visible in the category the user selected.
    # Their file extension determines whether this module can preview them.
    if (kind in FOLDER_ORDER and source_type in {"manual_upload", "project_upload"}
            and ext in project_files.MANAGED_EXTENSIONS):
        return str(kind)
    return None


def _public_item(item: dict[str, Any], folder: str) -> dict[str, Any]:
    downloadable = bool(item.get("downloadable"))
    is_html = _extension(item) == ".html"
    source_type = item.get("source_type")
    source_id = item.get("source_id")
    source_identity = item.get("source_identity")
    if not source_identity and source_type and source_id:
        source_identity = f"{source_type}:{source_id}"
    metadata = item.get("metadata") or {}
    source_task_url = metadata.get("source_task_url")
    if not source_task_url and source_type == "keyword_research" and source_id:
        source_task_url = f"#/keyword/run/{quote(str(source_id), safe='')}"
    return {
        "id": item["id"], "name": item["name"], "kind": folder,
        "created_at": item.get("created_at"), "size": int(item.get("size") or 0),
        "sha256": item.get("sha256"),
        "previewable": downloadable and (is_html or _extension(item) == ".xlsx"),
        "preview_reason": None if downloadable and (is_html or _extension(item) == ".xlsx")
                         else "此文件类型暂不支持在线预览",
        "downloadable": downloadable, "source_type": source_type, "source_id": source_id,
        "source_identity": source_identity, "source_task_url": source_task_url,
    }


def list_library(workspace_id: str) -> dict[str, Any]:
    project = project_files.public_project(workspace_id)
    # A legacy external index and its backfilled local copy describe one deliverable.
    # Prefer the verified local copy without revealing the external absolute path.
    unique: dict[tuple[str, str, str, str], tuple[str, dict[str, Any]]] = {}
    for artifact in project_files.list_artifacts(workspace_id):
        folder = _folder_for(artifact)
        if folder:
            key = (folder, str(artifact.get("source_type") or ""),
                   str(artifact.get("source_id") or ""), str(artifact.get("sha256") or ""))
            prior = unique.get(key)
            if prior is None or (not prior[1].get("downloadable") and artifact.get("downloadable")):
                unique[key] = (folder, artifact)
    grouped: dict[str, list[dict[str, Any]]] = {kind: [] for kind in FOLDER_ORDER}
    for folder, artifact in unique.values():
        grouped[folder].append(_public_item(artifact, folder))
    for items in grouped.values():
        items.sort(key=lambda item: (str(item.get("created_at") or ""), item["id"]), reverse=True)
    folders = [{"kind": kind, "label": FOLDER_LABELS[kind],
                "count": len(grouped[kind]), "items": grouped[kind]}
               for kind in FOLDER_ORDER]
    return {"project": {"id": project["id"], "name": project["name"]}, "folders": folders}


def _artifact(workspace_id: str, artifact_id: str) -> tuple[dict[str, Any], Path, str]:
    with Session() as session:
        record = session.get(ProjectArtifact, artifact_id)
        if not record or record.workspace_id != workspace_id:
            raise ValueError("项目资料不存在或不属于当前项目")
        public = project_files._public(record)
    folder = _folder_for(public)
    if not folder:
        raise ValueError("该项目资料不支持此预览")
    if not public.get("downloadable"):
        raise ValueError("历史文件仅登记索引，不支持预览")
    try:
        path = project_files.artifact_path(workspace_id, artifact_id)
    except FileNotFoundError as exc:
        raise ValueError("项目资料文件丢失") from exc
    except OSError as exc:
        raise ValueError("项目资料文件无法读取") from exc
    return public, path, folder


def _checked_zip(path: Path) -> None:
    if path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ValueError("工作簿超过 100 MB，无法预览")
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("工作簿解压后超过 500 MB，无法预览")
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValueError("工作簿损坏或格式无效") from exc


def _load_workbook(path: Path, *, data_only: bool):
    _checked_zip(path)
    try:
        return load_workbook(path, read_only=False, data_only=data_only, keep_links=False)
    except (OSError, zipfile.BadZipFile, InvalidFileException, ValueError) as exc:
        raise ValueError("工作簿损坏或格式无效") from exc


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _image_id(artifact_sha: str, sheet: str, row: int, column: int,
              index: int, data: bytes) -> str:
    identity = f"{artifact_sha}\0{sheet}\0{row}\0{column}\0{index}\0".encode()
    return hashlib.sha256(identity + hashlib.sha256(data).digest()).hexdigest()[:32]


def _build_image_catalog(path: Path, item: dict[str, Any]):
    """Parse workbook drawing and WPS images once into row references and opaque IDs."""
    workbook = _load_workbook(path, data_only=False)
    try:
        references: dict[str, dict[int, list[dict[str, Any]]]] = {}
        raw_by_id: dict[str, bytes] = {}
        try:
            with zipfile.ZipFile(path) as archive:
                cell_images = wps_images(archive)
        except (OSError, KeyError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
            raise ValueError("工作簿中的 WPS 内嵌图片结构损坏") from exc

        for worksheet in workbook.worksheets:
            if worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS:
                raise ValueError("工作簿超过预览限制（100000 行、500 列）")
            sheet_refs = references.setdefault(worksheet.title, {})
            for image in getattr(worksheet, "_images", ()):
                anchor = getattr(image, "anchor", None)
                marker = getattr(anchor, "_from", None)
                if marker is None:
                    continue
                row, column = marker.row + 1, marker.col + 1
                try:
                    data = image._data()
                except Exception:
                    continue
                image_id = _image_id(item["sha256"], worksheet.title, row, column,
                                     len(sheet_refs.get(row, [])), data)
                sheet_refs.setdefault(row, []).append({"column_index": column, "image_id": image_id})
                raw_by_id[image_id] = data

            for row in worksheet.iter_rows():
                excel_row = row[0].row if row else 0
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    match = _DISPIMG.search(cell.value)
                    if not match or match.group(1) not in cell_images:
                        continue
                    data = cell_images[match.group(1)]
                    column = cell.column
                    image_id = _image_id(item["sha256"], worksheet.title, excel_row, column,
                                         len(sheet_refs.get(excel_row, [])), data)
                    sheet_refs.setdefault(excel_row, []).append(
                        {"column_index": column, "image_id": image_id})
                    raw_by_id[image_id] = data
        return references, raw_by_id
    finally:
        workbook.close()


def _image_catalog(workspace_id: str, artifact_id: str, item: dict[str, Any], path: Path):
    """Use only SHA-scoped catalogs; artifact_path has already revalidated the source."""
    global _image_cache_bytes
    key = (workspace_id, artifact_id, str(item["sha256"]))
    with _image_cache_lock:
        cached = _image_cache.get(key)
        if cached is not None:
            _image_cache.move_to_end(key)
            return cached[0], cached[1]
        references, images = _build_image_catalog(path, item)
        cache_size = sum(len(data) for data in images.values())
        if cache_size <= _IMAGE_CACHE_MAX_BYTES:
            while _image_cache and (
                len(_image_cache) >= _IMAGE_CACHE_MAX_ENTRIES
                or _image_cache_bytes + cache_size > _IMAGE_CACHE_MAX_BYTES
            ):
                _, (_, _, removed_size) = _image_cache.popitem(last=False)
                _image_cache_bytes -= removed_size
            _image_cache[key] = (references, images, cache_size)
            _image_cache_bytes += cache_size
        return references, images


def workbook_page(workspace_id: str, artifact_id: str, *, sheet: str = "",
                  offset: int = 0, limit: int = 50, query: str = "") -> dict[str, Any]:
    item, path, _ = _artifact(workspace_id, artifact_id)
    if offset < 0 or offset > 10_000_000:
        raise ValueError("工作簿页码无效")
    if limit < 1 or limit > 100:
        raise ValueError("每页显示条数需在 1–100 之间")
    if not isinstance(query, str) or len(query) > 300:
        raise ValueError("工作簿搜索内容不能超过 300 个字符")
    workbook = _load_workbook(path, data_only=True)
    try:
        sheets = list(workbook.sheetnames)
        if not sheets:
            raise ValueError("工作簿中没有工作表")
        selected = sheet or sheets[0]
        if selected not in sheets:
            raise ValueError("所选工作表不存在")
        worksheet = workbook[selected]
        if worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS:
            raise ValueError("工作簿超过预览限制（100000 行、500 列）")
        iterator = worksheet.iter_rows(values_only=True)
        raw_headers = next(iterator, ())
        headers = [str(value).strip() if value is not None else f"列 {index + 1}"
                   for index, value in enumerate(raw_headers)]
        while headers and not headers[-1]:
            headers.pop()
        if not headers:
            raise ValueError("工作簿中没有可预览的数据表头或数据行")
        image_references, _ = _image_catalog(workspace_id, artifact_id, item, path)
        image_rows = image_references.get(selected, {})
        needle = query.strip().casefold()
        rows: list[dict[str, Any]] = []
        total = 0
        has_data_rows = False
        for excel_row, raw_row in enumerate(iterator, start=2):
            values = [_cell_value(value) for value in raw_row[:len(headers)]]
            values.extend([None] * (len(headers) - len(values)))
            if not any(value not in (None, "") for value in values):
                continue
            has_data_rows = True
            if needle and not any(needle in str(value).casefold() for value in values if value is not None):
                continue
            if offset <= total < offset + limit:
                images = [{"column": headers[ref["column_index"] - 1], "image_id": ref["image_id"]}
                          for ref in image_rows.get(excel_row, [])
                          if 1 <= ref["column_index"] <= len(headers)]
                rows.append({"excel_row": excel_row, "values": values, "images": images})
            total += 1
        if not has_data_rows:
            raise ValueError("工作簿中没有可预览的数据行")
        return {"sheets": sheets, "sheet": selected, "headers": headers, "rows": rows,
                "total": total, "offset": offset, "limit": limit, "query": query}
    finally:
        workbook.close()


def image_thumbnail(workspace_id: str, artifact_id: str, image_id: str) -> Response:
    item, path, _ = _artifact(workspace_id, artifact_id)
    _, images = _image_catalog(workspace_id, artifact_id, item, path)
    found = images.get(image_id)
    if found is None:
        raise ValueError("工作簿图片不存在")
    try:
        with Image.open(io.BytesIO(found)) as image:
            image.thumbnail((320, 320))
            output = io.BytesIO()
            image.convert("RGBA").save(output, format="PNG", optimize=True)
    except (OSError, ValueError) as exc:
        raise ValueError("工作簿图片损坏或格式无效") from exc
    return Response(output.getvalue(), media_type="image/png",
                    headers={"Cache-Control": "no-store", "X-Content-Type-Options": "nosniff"})


def html_text(workspace_id: str, artifact_id: str) -> str:
    _, path, _ = _artifact(workspace_id, artifact_id)
    if path.stat().st_size > MAX_HTML_BYTES:
        raise ValueError("HTML 文件超过 20 MB，无法预览")
    try:
        return path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as exc:
        raise ValueError("HTML 文件损坏或编码无效") from exc
