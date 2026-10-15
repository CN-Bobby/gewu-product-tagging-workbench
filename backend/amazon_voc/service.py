"""Workbench adapter for the accepted, review-only Amazon VOC runtime."""
from __future__ import annotations

import json
import hashlib
import html
import os
import socket
import threading
import time
from pathlib import Path
from typing import Any

from ..db import STORAGE
from ..keyword_research.credentials import model_target
from ..model_settings import defaults, read as read_model_settings, require_validated
from ..runtime_isolation import CredentialUnavailable
from .runtime.voc_app.config import Settings
from .runtime.voc_app.db import Database, utcnow
from .runtime.voc_app.pipeline import Pipeline
from .runtime.voc_app.provider import ModelError, OpenAICompatibleProvider
from .runtime.voc_app.storage import Storage


MAX_REVIEW_BYTES = 50 * 1024 * 1024
ALLOWED_INPUT_SUFFIXES = {".xlsx", ".xlsm"}
TERMINAL = {"completed", "failed", "cancelled"}


class SafeProvider:
    """Keep provider exceptions from copying an authorization secret into run logs."""
    def __init__(self, settings: Settings):
        self.settings = settings
        self.inner = OpenAICompatibleProvider(settings)

    def call(self, task, model, instructions, payload, schema):
        try:
            return self.inner.call(task, model, instructions, payload, schema)
        except Exception as exc:
            detail = str(exc)
            if self.settings.model_api_key:
                detail = detail.replace(self.settings.model_api_key, "[redacted]")
            raise ModelError(detail[:1500] or type(exc).__name__,
                             retryable=bool(getattr(exc, "retryable", True)),
                             status_code=getattr(exc, "status_code", None)) from None


def _review_source_context(source: dict[str, Any]) -> dict[str, Any]:
    items = source.get("items") if isinstance(source.get("items"), list) else []
    completed = [item for item in items if item.get("status") == "completed"]
    exported = [item for item in items if int(item.get("review_count") or 0) > 0]
    empty = [item for item in items if int(item.get("review_count") or 0) == 0]
    failed = [item for item in items if item.get("status") != "completed"]
    return {
        "source_job_id": str(source.get("id") or ""),
        "source_job_status": str(source.get("status") or ""),
        "workspace_id": str(source.get("workspace_id") or ""),
        "source_dataset_id": str(source.get("source_dataset_id") or ""),
        "source_archive_id": str(source.get("source_archive_id") or ""),
        "requested_asin_count": len(items),
        "completed_asin_count": len(completed),
        "comment_asin_count": len(exported),
        "empty_comment_asin_count": len(empty),
        "failed_asin_count": len(failed),
        "failed_asins": [str(item.get("asin") or "") for item in failed if item.get("asin")],
        "empty_comment_asins": [str(item.get("asin") or "") for item in empty if item.get("asin")],
        "review_count": int(source.get("review_count") or 0),
        "analysis_scope": ("仅分析本次评论任务最终 Excel 实际包含的评论；未完成 ASIN 已写入 Excel 的评论也会纳入，"
                           "并在失败清单中保留。每个 ASIN 最多两次 Both 调用，每次最多 100 条，覆盖最近一年。"
                           "重复调用不是分页，分析结果不代表商品全历史评论。"),
    }


def _append_review_scope(html_path: Path | None, excel_path: Path | None,
                         context: dict[str, Any]) -> None:
    if html_path and html_path.is_file():
        details = [
            ("来源 Reviews 任务", context.get("source_job_id", "")),
            ("任务状态", context.get("source_job_status", "")),
            ("请求 ASIN 数", context.get("requested_asin_count", 0)),
            ("成功完成 ASIN 数", context.get("completed_asin_count", 0)),
            ("有评论并进入分析的 ASIN 数", context.get("comment_asin_count", 0)),
            ("成功但无评论的 ASIN 数", context.get("empty_comment_asin_count", 0)),
            ("失败或未完成 ASIN 数", context.get("failed_asin_count", 0)),
            ("导出评论数", context.get("review_count", 0)),
        ]
        rows = "".join("<tr><th>%s</th><td>%s</td></tr>" % (html.escape(str(name)), html.escape(str(value)))
                       for name, value in details)
        failures = context.get("failed_asins") or []
        failure_list = "、".join(html.escape(str(value)) for value in failures) if failures else "无"
        section = ('<section class="module"><h2>评论采集范围与 ASIN 覆盖</h2><p>%s</p>'
                   '<table><tbody>%s</tbody></table><p><b>失败或未完成 ASIN：</b>%s</p></section>' %
                   (html.escape(str(context.get("analysis_scope") or "")), rows, failure_list))
        content = html_path.read_text(encoding="utf-8")
        if "评论采集范围与 ASIN 覆盖" not in content:
            anchor = "</main>"
            if anchor not in content:
                anchor = "</body>"
            if anchor in content:
                content = content.replace(anchor, section + anchor, 1)
                html_path.write_text(content, encoding="utf-8")

    if excel_path and excel_path.is_file():
        from openpyxl import load_workbook

        workbook = load_workbook(excel_path)
        if "采集范围" in workbook.sheetnames:
            del workbook["采集范围"]
        sheet = workbook.create_sheet("采集范围", 0)
        sheet.append(["项目", "内容"])
        for name, value in (
            ("来源 Reviews 任务", context.get("source_job_id", "")),
            ("任务状态", context.get("source_job_status", "")),
            ("请求 ASIN 数", context.get("requested_asin_count", 0)),
            ("成功完成 ASIN 数", context.get("completed_asin_count", 0)),
            ("有评论并进入分析的 ASIN 数", context.get("comment_asin_count", 0)),
            ("成功但无评论的 ASIN 数", context.get("empty_comment_asin_count", 0)),
            ("失败或未完成 ASIN 数", context.get("failed_asin_count", 0)),
            ("导出评论数", context.get("review_count", 0)),
            ("分析范围", context.get("analysis_scope", "")),
            ("失败或未完成 ASIN", "、".join(context.get("failed_asins") or []) or "无"),
        ):
            sheet.append([name, value])
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 100
        workbook.save(excel_path)


class AmazonVOCService:
    def __init__(self, root: Path | None = None):
        self.root = Path(root or (STORAGE / "amazon-voc")).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.db = Database(self.root / "voc.sqlite3")
        self.storage = Storage(self.root)
        self.db.execute("CREATE TABLE IF NOT EXISTS workbench_model_snapshots ("
                        "run_id TEXT PRIMARY KEY REFERENCES runs(id), model_ref TEXT NOT NULL, "
                        "credential_id TEXT NOT NULL DEFAULT '')")
        self.db.execute("CREATE TABLE IF NOT EXISTS workbench_reviews_handoffs ("
                        "source_job_id TEXT PRIMARY KEY, workbook_sha256 TEXT NOT NULL, "
                        "project_id TEXT NOT NULL REFERENCES projects(id), "
                        "run_id TEXT NOT NULL UNIQUE REFERENCES runs(id), created_at TEXT NOT NULL)")
        self._lock = threading.RLock()
        self._wake = threading.Event()
        self._stop = threading.Event()
        self._thread: threading.Thread | None = None

    def public_model_settings(self) -> dict[str, Any]:
        settings = read_model_settings()
        text_model, _ = defaults()
        providers = []
        for provider in settings.get("providers", []):
            providers.append({
                "name": provider.get("name"),
                "model_refs": provider.get("model_refs", []),
                "credentials": [{key: credential.get(key) for key in
                                 ("id", "name", "enabled", "authorized")}
                                for credential in provider.get("credentials", [])],
            })
        selected_provider = next((p for p in providers if p["name"] == text_model.partition("/")[0]), None)
        selected_provider_raw = next((p for p in settings.get("providers", [])
                                      if p.get("name") == text_model.partition("/")[0]), None)
        credential_id = (selected_provider_raw or {}).get("primary_credential_id")
        if selected_provider:
            active_ids = [c["id"] for c in selected_provider["credentials"]
                          if c.get("id") and c.get("enabled") and c.get("authorized")]
            if credential_id not in active_ids:
                credential_id = next(iter(active_ids), "")
        return {"configured": bool(settings.get("configured")), "default_model": text_model,
                "default_credential_id": credential_id or "", "providers": providers}

    def resolve_shared_model(self, model_ref: str | None = None,
                             credential_id: str | None = None) -> dict[str, str]:
        model_ref = str(model_ref or defaults()[0]).strip()
        public = self.public_model_settings()
        provider_name = model_ref.partition("/")[0]
        provider = next((p for p in public["providers"] if p["name"] == provider_name), None)
        if not provider or model_ref not in provider.get("model_refs", []):
            raise CredentialUnavailable("请选择总设置中已配置的文本模型")
        credential_id = str(credential_id or "").strip()
        if not credential_id:
            credential_id = (public.get("default_credential_id")
                             if model_ref == public.get("default_model") else "") or ""
            if not credential_id:
                credential_id = next((c["id"] for c in provider["credentials"]
                                      if c.get("id") and c.get("enabled") and c.get("authorized")), "")
        if not credential_id and provider.get("credentials"):
            raise CredentialUnavailable("请为 VOC 选择总设置中已授权的文本模型凭据")
        base_url, api_key, model = model_target(model_ref, credential_id or None)
        require_validated(model_ref, credential_id or None, "text")
        return {"model_ref": model_ref, "credential_id": credential_id,
                "base_url": base_url, "api_key": api_key, "model": model}

    def create_project(self, body: dict[str, Any]) -> dict[str, Any]:
        name = str(body.get("name") or "").strip()
        marketplace = str(body.get("marketplace") or "US").strip().upper()
        country = str(body.get("country") or marketplace).strip()
        language = str(body.get("language") or "en").strip()
        category = str(body.get("category") or "").strip()
        if not name or len(name) > 120:
            raise ValueError("项目名称必填且不能超过 120 个字符")
        if len(marketplace) > 16 or len(country) > 80 or len(language) > 40 or len(category) > 120:
            raise ValueError("站点、国家、语言或品类设置过长")
        workspace_id = str(body.get("workspace_id") or "").strip() or None
        if workspace_id:
            from ..db import Session, Workspace, require
            with Session() as session:
                require(session, Workspace, workspace_id)
        return self.db.create_project(name, marketplace, country, language, category, workspace_id)

    def create_run(self, project_id: str, filename: str, data: bytes,
                   model_ref: str | None = None, credential_id: str | None = None,
                   source_context: dict[str, Any] | None = None) -> dict[str, Any]:
        suffix = Path(filename or "").suffix.lower()
        if suffix not in ALLOWED_INPUT_SUFFIXES:
            raise ValueError("仅支持 .xlsx 或 .xlsm 评论Excel文件")
        if not data:
            raise ValueError("评论Excel文件为空")
        if len(data) > MAX_REVIEW_BYTES:
            raise ValueError("评论Excel文件最大支持 50 MiB")
        project = self.db.one("SELECT * FROM projects WHERE id=?", (project_id,))
        if not project:
            raise KeyError("项目不存在")
        selected = self.resolve_shared_model(model_ref, credential_id)
        review_path, review_hash = self.storage.save_raw(project_id, filename, data)
        sources = {"reviews": str(review_path), "reviews_sha256": review_hash}
        if source_context:
            sources["review_source_context"] = dict(source_context)
        run = self.db.create_run(project_id, sources)
        self.db.execute("INSERT INTO workbench_model_snapshots(run_id,model_ref,credential_id) VALUES(?,?,?)",
                        (run["id"], selected["model_ref"], selected["credential_id"]))
        self.storage.revision_dir(project_id, run["id"], 1)
        if project.get("workspace_id"):
            from .. import project_files
            project_files.register_file(project["workspace_id"], "voc", review_path,
                source_type="amazon_voc_input", source_id=run["id"], name=Path(filename).name,
                metadata={"voc_project_id": project_id, "run_id": run["id"], "reviews_sha256": review_hash})
        self._wake.set()
        return self.public_run(run)

    def create_run_from_amazon_reviews(self, source_job_id: str,
                                       confirmed_reviewed: bool) -> dict[str, Any]:
        """Create or return the single VOC handoff for a reviewed Amazon Reviews workbook."""
        if confirmed_reviewed is not True:
            raise ValueError("请先确认已检查 Amazon 评论导出的 Excel")

        # Resolve the source through the Reviews service. Never accept a caller path.
        from ..amazon_reviews import service as amazon_reviews_service

        source = amazon_reviews_service.get(source_job_id)
        allowed_statuses = {"completed", "completed_with_errors", "partial"}
        if source.get("status") not in allowed_statuses or not source.get("can_download"):
            raise ValueError("评论任务尚无可导入的最终 Excel")
        if int(source.get("review_count") or 0) <= 0:
            raise ValueError("评论 Excel 中没有可分析的评论")

        path = amazon_reviews_service.output_path(source_job_id)
        if path.stat().st_size > MAX_REVIEW_BYTES:
            raise ValueError("评论 Excel 超过 VOC 的 50 MiB 文件限制")
        data = path.read_bytes()
        if len(data) > MAX_REVIEW_BYTES:
            raise ValueError("评论 Excel 超过 VOC 的 50 MiB 文件限制")
        workbook_sha256 = hashlib.sha256(data).hexdigest()

        with self._lock:
            existing = self.db.one("SELECT * FROM workbench_reviews_handoffs WHERE source_job_id=?",
                                    (source_job_id,))
            if existing:
                if existing["workbook_sha256"] != workbook_sha256:
                    raise RuntimeError("该评论任务的导出 Excel 已变化；请使用新的评论任务创建 VOC 分析")
                project = self.project(existing["project_id"])
                run = self.get_run(existing["run_id"])
                if not project or not run:
                    raise RuntimeError("已有 VOC 交接记录不完整，不能重复创建分析")
                duplicate = True
                created_at = existing["created_at"]
            else:
                site = str(source.get("site") or "US").strip().upper()
                settings = self.public_model_settings()
                model_ref = settings.get("default_model")
                credential_id = settings.get("default_credential_id") or None
                # Validate the shared model before creating persistent VOC records. If the
                # workbench has no authorized default credential, the user can fix settings
                # and retry without accumulating empty projects.
                self.resolve_shared_model(model_ref, credential_id)
                project = self.create_project({
                    "name": f"Amazon VOC · {site} · 评论任务 {source_job_id[:8]}",
                    "marketplace": site, "country": site, "language": "auto", "category": "",
                    "workspace_id": source.get("workspace_id"),
                })
                run = self.create_run(project["id"], path.name, data, model_ref, credential_id,
                                      source_context=_review_source_context(source))
                created_at = utcnow()
                self.db.execute("INSERT INTO workbench_reviews_handoffs "
                                "(source_job_id,workbook_sha256,project_id,run_id,created_at) VALUES(?,?,?,?,?)",
                                (source_job_id, workbook_sha256, project["id"], run["id"], created_at))
                duplicate = False

            handoff = {"source_job_id": source_job_id, "workbook_sha256": workbook_sha256,
                       "project_id": project["id"], "run_id": run["id"],
                       "created_at": created_at, "status": "started"}
            source_summary = amazon_reviews_service.record_voc_handoff(source_job_id, handoff)
            return {"run": run, "project": project,
                    "source_job": {key: source_summary.get(key) for key in
                                   ("id", "status", "site", "review_count", "voc_handoff")},
                    "duplicate": duplicate}

    def model_snapshot(self, run_id: str) -> dict[str, Any]:
        return self.db.one("SELECT run_id,model_ref,credential_id FROM workbench_model_snapshots WHERE run_id=?",
                           (run_id,)) or {}

    def public_run(self, run: dict[str, Any]) -> dict[str, Any]:
        if not run:
            return run
        project_id, run_id = run["project_id"], run["id"]
        revision = int(run.get("current_revision") or 1)
        snapshot = self.model_snapshot(run_id)
        artifacts = [kind for kind in ("html", "excel", "quality")
                     if self.storage.artifact(project_id, run_id, revision, kind)]
        sources = run.get("sources") or {}
        public = {key: value for key, value in run.items() if key not in {"sources", "source_json"}}
        filename = Path(sources.get("reviews") or "reviews.xlsx").name
        digest_prefix = str(sources.get("reviews_sha256") or "")[:12] + "_"
        if digest_prefix != "_" and filename.startswith(digest_prefix):
            filename = filename[len(digest_prefix):]
        public.update({"sources": {"reviews": filename,
                                    "reviews_sha256": sources.get("reviews_sha256", ""),
                                    "review_source_context": sources.get("review_source_context")},
                       "model_ref": snapshot.get("model_ref", ""),
                       "credential_id": snapshot.get("credential_id", ""),
                       "artifacts": artifacts, "storage_scope": "amazon-voc"})
        return public

    def get_run(self, run_id: str) -> dict[str, Any] | None:
        run = self.db.get_run(run_id)
        if not run:
            return None
        public = self.public_run(run)
        public["model_calls"] = self.db.all(
            "SELECT stage,batch_id,model,prompt_version,protocol,latency_ms,input_tokens,output_tokens,retries,status,created_at "
            "FROM model_calls WHERE run_id=? AND revision=? ORDER BY id DESC LIMIT 100", (run_id, public["current_revision"]))
        return public

    def list_runs(self, project_id: str | None = None) -> list[dict[str, Any]]:
        return [self.public_run(run) for run in self.db.list_runs(project_id)]

    def project(self, project_id: str) -> dict[str, Any] | None:
        return self.db.one("SELECT p.*,COUNT(r.id) run_count,MAX(r.updated_at) last_run_at "
                           "FROM projects p LEFT JOIN runs r ON r.project_id=p.id WHERE p.id=? GROUP BY p.id",
                           (project_id,))

    def list_projects(self) -> list[dict[str, Any]]:
        return self.db.list_projects()

    def run_pipeline(self, run_id: str, provider=None) -> dict[str, Any]:
        run = self.db.get_run(run_id)
        if not run:
            raise KeyError("任务不存在")
        selected = self.resolve_shared_model(self.model_snapshot(run_id).get("model_ref"),
                                             self.model_snapshot(run_id).get("credential_id"))
        settings = Settings(data_dir=self.root, database_path=self.db.path,
                            model_base_url=selected["base_url"], model_api_key=selected["api_key"],
                            model_protocol="chat", extractor_model=selected["model"],
                            reviewer_model=selected["model"],
                            **Settings.runtime_tuning_from_env())
        settings.ensure()
        try:
            Pipeline(settings, self.db, self.storage, provider or SafeProvider(settings)).process(run_id)
        except Exception as exc:
            message = str(exc)
            if selected["api_key"]:
                message = message.replace(selected["api_key"], "[redacted]")
            self.db.update_run(run_id, status="failed", error=message[:1500], completed_at=utcnow())
            raise
        result = self.get_run(run_id)
        source_context = (result.get("sources") or {}).get("review_source_context") if result else None
        if source_context:
            revision = int(result.get("current_revision") or 1)
            _append_review_scope(self.storage.artifact(result["project_id"], run_id, revision, "html"),
                                 self.storage.artifact(result["project_id"], run_id, revision, "excel"),
                                 source_context)
        workspace_id = result.get("workspace_id") if result else None
        if workspace_id and result and result.get("status") == "completed":
            from .. import project_files
            revision = int(result.get("current_revision") or 1)
            for kind, filename in (("html", "VOC-report.html"), ("excel", "VOC-report.xlsx"),
                                   ("quality", "VOC-quality.json")):
                path = self.storage.artifact(result["project_id"], run_id, revision, kind)
                if path and path.is_file():
                    project_files.register_file(workspace_id, "voc", path,
                        source_type="amazon_voc_output", source_id=f"{run_id}:rev{revision}:{kind}",
                        name=filename, metadata={"voc_project_id": result["project_id"],
                                                "run_id": run_id, "revision": revision})
        if result and result.get("status") == "completed":
            try:
                from ..product_archive.sync import sync_voc_run
                result["archive_sync"] = sync_voc_run(self, run_id)
            except Exception as exc:  # Archive enrichment must not change VOC completion.
                result["archive_sync"] = {"status": "sync_failed", "error_type": type(exc).__name__}
        return result

    def recover(self) -> None:
        now = utcnow()
        self.db.execute("UPDATE runs SET status='cancelled',cancel_requested=0,worker_id='',completed_at=?,updated_at=? "
                        "WHERE cancel_requested=1 AND status NOT IN ('completed','failed','cancelled')", (now, now))
        self.db.execute("UPDATE runs SET worker_id='',updated_at=? WHERE status NOT IN ('completed','failed','cancelled')",
                        (now,))

    def start_worker(self) -> None:
        with self._lock:
            if self._thread and self._thread.is_alive():
                return
            self.recover()
            self._stop.clear()
            self._thread = threading.Thread(target=self._worker_loop, name="amazon-voc-worker", daemon=True)
            self._thread.start()

    def stop_worker(self) -> None:
        self._stop.set()
        self._wake.set()
        thread = self._thread
        if thread and thread.is_alive():
            thread.join(timeout=5)

    def _worker_loop(self) -> None:
        worker_id = f"{socket.gethostname()}:{os.getpid()}:amazon-voc"
        while not self._stop.is_set():
            # Handoffs hold this lock until the source-job mapping is durable. This
            # keeps the worker from starting model calls before a retry can find the
            # idempotency record.
            with self._lock:
                run_id = self.db.pending_run(worker_id)
            if not run_id:
                self._wake.wait(1.0)
                self._wake.clear()
                continue
            try:
                self.run_pipeline(run_id)
            except Exception:
                # The durable run status carries the safe, user-visible failure.
                pass

    def cancel(self, run_id: str) -> dict[str, Any]:
        run = self.db.get_run(run_id)
        if not run:
            raise KeyError("任务不存在")
        if run["status"] in TERMINAL:
            raise RuntimeError("任务已结束")
        if not run.get("worker_id"):
            self.db.update_run(run_id, status="cancelled", cancel_requested=0, completed_at=utcnow())
        else:
            self.db.update_run(run_id, cancel_requested=1)
        self._wake.set()
        return self.get_run(run_id)

    def retry(self, run_id: str) -> dict[str, Any]:
        run = self.db.get_run(run_id)
        if not run:
            raise KeyError("任务不存在")
        if run["status"] not in {"failed", "cancelled"}:
            raise RuntimeError("只有失败或已取消的任务可以重试")
        self.db.update_run(run_id, status="uploaded", error="", cancel_requested=0,
                           completed_at=None, worker_id="")
        self._wake.set()
        return self.get_run(run_id)

    def artifact_path(self, run_id: str, kind: str) -> Path:
        if kind not in {"html", "excel", "quality"}:
            raise ValueError("不支持的 VOC 结果类型")
        run = self.db.get_run(run_id)
        if not run:
            raise KeyError("任务不存在")
        path = self.storage.artifact(run["project_id"], run_id, run["current_revision"], kind)
        if not path:
            raise FileNotFoundError("结果文件尚未生成")
        return path

    def cache_doc(self, run_id: str, name: str, default: Any) -> Any:
        run = self.db.get_run(run_id)
        if not run:
            raise KeyError("任务不存在")
        if name not in {"events.json", "segment_model.json", "quality_report.json"}:
            raise ValueError("不支持的 VOC 结果内容")
        path = self.storage.revision_dir(run["project_id"], run_id, run["current_revision"]) / "cache" / name
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default

    def correct_event(self, run_id: str, event_id: str, event: dict[str, Any], note: str = "") -> dict[str, Any]:
        run = self.db.get_run(run_id)
        if not run:
            raise KeyError("任务不存在")
        if run["status"] != "completed":
            raise RuntimeError("只有已完成任务可创建修订")
        old_revision, new_revision = run["current_revision"], run["current_revision"] + 1
        target = self.storage.clone_revision(run["project_id"], run_id, old_revision, new_revision)
        events_path = target / "cache" / "events.json"
        document = json.loads(events_path.read_text(encoding="utf-8"))
        events = document.get("events") or []
        before = next((row for row in events if row.get("event_id") == event_id), None)
        if not before:
            raise KeyError("评论事件不存在")
        after = dict(event)
        for key in ("event_id", "review_id", "asin", "rating", "review_date", "review_text"):
            after[key] = before.get(key)
        document["events"] = [after if row.get("event_id") == event_id else row for row in events]
        events_path.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")
        for name in ("tag_dictionary.json", "segment_model.json", "analysis.json", "evidence_index.json",
                     "journey_matrix.json", "quality_report.json", "validate_report.json"):
            cache_path = target / "cache" / name
            if cache_path.exists():
                cache_path.unlink()
        for path in (target / "output").glob("*"):
            if path.is_file():
                path.unlink()
        self.db.execute("INSERT INTO corrections(run_id,from_revision,to_revision,event_id,before_json,after_json,note,created_at) "
                        "VALUES(?,?,?,?,?,?,?,?)", (run_id, old_revision, new_revision, event_id,
                        json.dumps(before, ensure_ascii=False), json.dumps(after, ensure_ascii=False),
                        str(note)[:500], utcnow()))
        self.db.update_run(run_id, current_revision=new_revision, status="normalizing", error="",
                           completed_at=None, worker_id="", cancel_requested=0)
        self._wake.set()
        return {"run_id": run_id, "revision": new_revision, "status": "normalizing"}


service = AmazonVOCService()
