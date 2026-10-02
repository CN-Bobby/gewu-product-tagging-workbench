# -*- coding: utf-8 -*-
"""Initialize a VOC task from one Amazon review workbook."""
import argparse
import datetime
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common
from openpyxl import load_workbook

FIELD_ALIASES = {
    "asin": ["asin", "来源asin", "产品asin", "asin码"],
    "title": ["产品标题", "商品标题", "标题", "title", "商品名称"],
    "rank": ["top50排名", "top100排名", "排名", "rank", "bsr排名", "榜单排名"],
    "url": ["商品链接", "链接", "url", "商品url"],
    "sentiment": ["情感", "sentiment", "正负"],
    "rating": ["评分", "星级", "rating", "star", "评论星级"],
    "review_title": ["评论标题", "评论标题(EN)", "评论标题(中文)", "review title"],
    "review_content": ["评论内容", "评论正文", "评论内容(EN)", "评论内容(中文)",
                       "评论原文", "评论文本", "review", "review text", "review body",
                       "review content", "comment", "comment text", "content", "text", "正文"],
    "variant": ["变体属性", "变体", "variant", "尺码颜色", "size/color"],
    "review_date": ["评论日期", "日期", "date", "评论时间"],
}

SHEET_ALIASES = {
    "review_detail": ["全部评论", "评论明细", "评论详情", "评论", "reviews", "review", "明细"],
}


def pick_sheet(wb, key):
    """按别名识别工作表，返回第一个命中的 sheet 或 None。"""
    targets = [common.norm(a) for a in SHEET_ALIASES[key]]
    # Prefer exact detail-sheet names even when summary tabs appear first.
    for ws in wb.worksheets:
        if common.norm(ws.title) in targets:
            return ws
    for ws in wb.worksheets:
        nt = common.norm(ws.title)
        if any(token in nt for token in ("统计", "汇总", "summary", "stats")):
            continue
        specific = [target for target in targets if target not in {common.norm("评论"), common.norm("review")}]
        if any(target in nt for target in specific if target):
            return ws
    return None


def build_col_map(header, override=None):
    """header: list[str] 表头。返回 {规范字段: 列索引}。"""
    over = override or {}
    colmap = {}
    norm_header = [common.norm(h) for h in header]
    for field, aliases in FIELD_ALIASES.items():
        cands = list(aliases) + list(over.get(field, []))
        for idx, nh in enumerate(norm_header):
            if nh and nh in [common.norm(c) for c in cands]:
                colmap[field] = idx
                break
    return colmap


def read_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], {}
    header = ["" if h is None else str(h) for h in rows[0]]
    return rows[1:], header


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return int(v)
    return str(v).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--reviews", required=True, help="唯一输入：评论Excel工作簿")
    ap.add_argument("--workdir", required=True)
    ap.add_argument("--marketplace", default="US")
    ap.add_argument("--country", default="", help="目标国家/地区；留空时使用 marketplace")
    ap.add_argument("--language", default="", help="评论主要语言；留空时由分析阶段识别")
    ap.add_argument("--category", default="", help="目标品类；不得由案例常量写死")
    ap.add_argument("--segment-min-review-rate", type=float, default=0.005,
                    help="正式细分最低独立评论出现率，默认 0.5%%")
    ap.add_argument("--config", default="")
    ap.add_argument("--task-id", default="", help="留空则自动生成")
    args = ap.parse_args()

    if os.path.splitext(args.reviews)[1].lower() not in {".xlsx", ".xlsm"}:
        ap.error("--reviews 仅支持 .xlsx 或 .xlsm 评论工作簿")

    common.ensure_dirs(args.workdir)
    override = common.load_aliases(args.config) if args.config else {}

    wb = load_workbook(args.reviews, data_only=True)
    ws_detail = pick_sheet(wb, "review_detail")
    review_rows = []
    if ws_detail is None and len(wb.worksheets) == 1:
        ws_detail = wb.worksheets[0]
    if ws_detail is not None:
        body, header = read_rows(ws_detail)
        cm = build_col_map(header, override.get("reviews", {}))
        missing = [field for field in ("asin", "review_content") if field not in cm]
        if missing:
            raise ValueError("评论工作表缺少必需列：%s；识别到表头：%s" % (", ".join(missing), header))
        for r in body:
            review_rows.append({
                "asin": clean(r[cm["asin"]]) if "asin" in cm else "",
                "title": clean(r[cm["title"]]) if "title" in cm else "",
                "rank": clean(r[cm["rank"]]) if "rank" in cm else "",
                "url": clean(r[cm["url"]]) if "url" in cm else "",
                "sentiment": clean(r[cm["sentiment"]]) if "sentiment" in cm else "",
                "rating": clean(r[cm["rating"]]) if "rating" in cm else "",
                "review_title": clean(r[cm["review_title"]]) if "review_title" in cm else "",
                "content": clean(r[cm["review_content"]]) if "review_content" in cm else "",
                "variant": clean(r[cm["variant"]]) if "variant" in cm else "",
                "review_date": clean(r[cm["review_date"]]) if "review_date" in cm else "",
            })

    wb.close()
    if not review_rows:
        raise ValueError("未读取到有效评论；评论明细是唯一必需输入")

    # ---- 产品主表 ----
    asin_set = {}
    review_count_by_asin = {}
    for row in review_rows:
        if not row["asin"]:
            continue
        review_count_by_asin[row["asin"]] = review_count_by_asin.get(row["asin"], 0) + 1
        asin_set.setdefault(row["asin"], {"has_reviews": False, "title": row["title"]})
        asin_set[row["asin"]]["has_reviews"] = True
        if row["title"]:
            asin_set[row["asin"]]["title"] = row["title"]

    products = []
    for asin, info in sorted(asin_set.items()):
        products.append({
            "asin": asin,
            "title": info.get("title", ""),
            "rank": "",
            "has_reviews": info["has_reviews"],
            "review_count": review_count_by_asin.get(asin, 0),
        })

    # 用评论明细回填排名
    rank_map = {}
    for row in review_rows:
        if row["asin"] and row["rank"]:
            rank_map[row["asin"]] = row["rank"]
    for p in products:
        p["rank"] = rank_map.get(p["asin"], "")

    task_id = args.task_id or f"{args.marketplace}_VOC_{datetime.datetime.now().strftime('%Y%m%d')}"
    master = {
        "task_id": task_id,
        "marketplace": args.marketplace,
        "country": args.country or args.marketplace,
        "language": args.language,
        "category": args.category,
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "source_files": {"reviews": os.path.abspath(args.reviews)},
        "analysis_config": {
            "segment_min_review_rate": max(0.0, min(1.0, args.segment_min_review_rate)),
            "segment_min_review_count": 2,
        },
        "stats": {
            "review_asin_count": len([p for p in products if p["has_reviews"]]),
            "total_asin_count": len(products),
            "review_count": len(review_rows),
        },
        "products": products,
    }
    common.write_json(common.cache_path(args.workdir, "master.json"), master)
    common.write_json(common.cache_path(args.workdir, "raw_reviews.json"), {"rows": review_rows})

    print(f"[load_input] 任务ID={task_id}")
    print(f"  评论 ASIN={master['stats']['review_asin_count']}  总 ASIN={master['stats']['total_asin_count']}")
    print(f"  评论行={master['stats']['review_count']}")


if __name__ == "__main__":
    main()
