# -*- coding: utf-8 -*-
"""Deterministic delivery checks for the one-call VOC V2 report."""
import argparse
import glob
import html
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common


EXPECTED_HEADINGS = ["01 · 数据与方法概览", "02 · 全部细分总览", "03 · 全部细分详情",
                     "04 · 细分横向比较", "05 · 动态用户旅程", "06 · 探索性信号"]
EXPECTED_SHEETS = ["任务与方法说明", "产品主表", "清洗后评论", "评论处理覆盖", "语义事件", "动态标签词典",
                   "事件标签映射", "正式与探索细分", "细分需求", "动态旅程", "证据索引", "质量门禁"]


def _normal(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def validate(workdir):
    checks = []

    def add(name, passed, detail=None):
        checks.append({"name": name, "pass": bool(passed), "detail": detail if detail is not None else ""})

    cache = os.path.join(workdir, "cache")
    output = os.path.join(workdir, "output")
    analysis = common.read_json(common.cache_path(workdir, "analysis.json"))
    master = common.read_json(common.cache_path(workdir, "master.json"))
    cleaned = common.read_json(common.cache_path(workdir, "cleaned_reviews.json")).get("rows") or []
    events = common.read_json(common.cache_path(workdir, "events.json")).get("events") or []
    processing = common.read_json(common.cache_path(workdir, "review_processing.json")).get("reviews") or []
    segments = analysis.get("segments") or []
    source_files = master.get("source_files") or {}

    add("单个评论Excel是唯一来源", set(source_files) == {"reviews"}, sorted(source_files))
    add("单次分析模式为 VOC V2", analysis.get("analysis_mode") == "full_voc_v2", analysis.get("analysis_mode"))
    add("HTML和Excel六模块分析数据一致", (analysis.get("counts") or {}).get("reviews") == len(cleaned)
        and (analysis.get("counts") or {}).get("events") == len(events),
        {"reviews": len(cleaned), "events": len(events)})
    add("评论处理覆盖行数一致", len(processing) == len(cleaned), {"processing": len(processing), "reviews": len(cleaned)})
    add("分析证据范围仅为上传评论", (analysis.get("method") or {}).get("evidence_scope") == "uploaded_review_workbook_only")
    forbidden_payload = {"recommendations", "external_evidence", "external_research_threshold"}.intersection(analysis)
    add("分析数据不含外部补全或产品建议", not forbidden_payload, sorted(forbidden_payload))

    html_files = glob.glob(os.path.join(output, "*.html"))
    xlsx_files = glob.glob(os.path.join(output, "*.xlsx"))
    add("HTML与Excel均已生成", bool(html_files and xlsx_files), {"html": len(html_files), "excel": len(xlsx_files)})
    headings, html_counts, html_text = [], None, ""
    if html_files:
        with open(html_files[0], "r", encoding="utf-8") as handle:
            html_text = handle.read()
        headings = [re.sub(r"<[^>]+>", "", value).strip()
                    for value in re.findall(r"<h2[^>]*>(.*?)</h2>", html_text, re.S)]
        match = re.search(r'data-counts="(\d+),(\d+),(\d+)"', html_text)
        html_counts = tuple(map(int, match.groups())) if match else None
    add("HTML六模块名称与顺序", headings == EXPECTED_HEADINGS, headings)
    add("HTML其余模块使用卡片布局", all(token in html_text for token in (
        'class="overview-grid"', 'class="comparison-grid"', 'class="journey-stage-grid"',
        'class="v2-insight-grid"')))
    add("HTML中英文报告字段已呈现", all(token in html_text for token in (
        'lang="en"', "细分横向比较", "跨细分共性", "探索性信号", "报告限制")))
    add("HTML评论/事件/细分计数匹配", html_counts == (len(cleaned), len(events), len(segments)),
        {"expected": [len(cleaned), len(events), len(segments)], "actual": html_counts})

    workbook_sheets, excel_counts = [], None
    if xlsx_files:
        from openpyxl import load_workbook
        workbook = load_workbook(xlsx_files[0], read_only=True, data_only=True)
        workbook_sheets = workbook.sheetnames
        try:
            excel_counts = (workbook["清洗后评论"].max_row - 1,
                            workbook["语义事件"].max_row - 1,
                            workbook["正式与探索细分"].max_row - 1)
        except KeyError:
            excel_counts = None
        workbook.close()
    add("Excel十二表名称与顺序", workbook_sheets == EXPECTED_SHEETS, workbook_sheets)
    add("Excel评论/事件/细分计数匹配", excel_counts == (len(cleaned), len(events), len(segments)),
        {"expected": [len(cleaned), len(events), len(segments)], "actual": excel_counts})

    forbidden_terms = [term for term in ("AI产品建议", "外部补全", "外部证据", "Alexa", "Customer Say")
                       if term.casefold() in html_text.casefold()]
    add("HTML不包含外部补全或产品建议", not forbidden_terms, forbidden_terms)
    review_texts = sorted({_normal(row.get("content")) for row in cleaned
                           if len(_normal(row.get("content"))) >= 40}, key=len, reverse=True)
    html_quotes = [_normal(html.unescape(value)) for value in re.findall(r'data-quote="([^"]*)"', html_text)]
    complete_review_embeds = [index for index, text in enumerate(review_texts)
                              if text in html_quotes or text in _normal(html_text)]
    add("HTML不嵌入完整评论原文", not complete_review_embeds, complete_review_embeds[:10])

    all_green = all(row["pass"] for row in checks)
    report = {"schema": "amazon-voc-v2-delivery-validation/v1", "all_green": all_green,
              "passed": sum(1 for row in checks if row["pass"]), "total": len(checks), "checks": checks}
    common.write_json(common.cache_path(workdir, "voc_v2_delivery_report.json"), report)

    quality_path = common.cache_path(workdir, "quality_report.json")
    quality = common.read_json(quality_path) if os.path.exists(quality_path) else {"gates": []}
    gates = [gate for gate in quality.get("gates") or [] if gate.get("gate_id") != "V2-OUTPUT"]
    gates.append({"gate_id": "V2-OUTPUT", "name": "HTML/Excel 输出一致性", "pass": all_green,
                  "blocking": True, "failed_count": sum(1 for row in checks if not row["pass"]),
                  "checks": checks})
    quality["gates"] = gates
    quality["all_automatic_gates_green"] = bool(quality.get("all_automatic_gates_green")) and all_green
    # An independent semantic review is still required before a user accepts the VOC findings.
    quality["acceptance_ready"] = False
    common.write_json(quality_path, quality)
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workdir", required=True)
    args = parser.parse_args()
    report = validate(args.workdir)
    for check in report["checks"]:
        print("[validate-v2] %s - %s" % ("PASS" if check["pass"] else "FAIL", check["name"]))
    if not report["all_green"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
