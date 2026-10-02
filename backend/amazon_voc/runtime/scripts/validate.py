# -*- coding: utf-8 -*-
"""Final V4 delivery validation."""
import argparse
import glob
import html
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common
from validate_methodology import validate_methodology


EXPECTED_HEADINGS = ["01 · 数据与方法概览", "02 · 全部细分总览", "03 · 全部细分详情",
                     "04 · 细分横向比较", "05 · 动态用户旅程", "06 · 探索性信号"]
EXPECTED_SHEETS = ["任务与方法说明", "产品主表", "清洗后评论", "评论处理覆盖", "语义事件", "动态标签词典",
                   "事件标签映射", "正式与探索细分", "细分需求", "动态旅程", "证据索引", "质量门禁"]
FORBIDDEN_CACHE_FILES = {"raw_alexa.json", "cleaned_alexa.json", "external_enrichment.json",
                         "manual_external_sources.json", "recommendations.json"}


def find_full_review_embeds(events, html_text, minimum_chars=20):
    """Detect complete review text, while ignoring short phrase collisions in labels/other quotes."""
    normalize = lambda value: re.sub(r"\s+", " ", str(value or "")).strip().casefold()
    embedded_quotes = {
        normalize(html.unescape(value))
        for value in re.findall(r'data-quote="([^"]*)"', html_text)
    }
    hits = []
    for item in events:
        source = str(item.get("review_text") or "")
        source_normalized = normalize(source)
        if not source_normalized:
            continue
        if source_normalized in embedded_quotes:
            hits.append(item.get("event_id"))
            continue
        if len(source_normalized) >= minimum_chars and (
                source in html_text or html.escape(source, quote=True) in html_text):
            hits.append(item.get("event_id"))
    return hits


def validate(workdir):
    checks = []

    def add(name, ok, detail=""):
        checks.append({"check": name, "pass": bool(ok), "detail": str(detail)})

    method = validate_methodology(workdir)
    add("方法门禁", method.get("all_green"), "%s/%s" % (method.get("passed"), method.get("total")))
    analysis_path = common.cache_path(workdir, "analysis.json")
    analysis = common.read_json(analysis_path) if os.path.exists(analysis_path) else {}
    master_path = common.cache_path(workdir, "master.json")
    master = common.read_json(master_path) if os.path.exists(master_path) else {}
    add("唯一分析源", analysis.get("schema") == "amazon-voc-analysis/v4", analysis.get("schema"))
    add("唯一评论工作簿来源", set((master.get("source_files") or {}).keys()) == {"reviews"},
        sorted((master.get("source_files") or {}).keys()))
    add("评论工作簿证据范围", (analysis.get("method") or {}).get("evidence_scope") == "uploaded_review_workbook_only",
        (analysis.get("method") or {}).get("evidence_scope"))
    forbidden_keys = {"recommendations", "external_evidence", "external_research_threshold"}
    found_keys = sorted(forbidden_keys.intersection(analysis))
    add("分析结果仅含评论洞察", not found_keys, found_keys)
    segments = analysis.get("segments") or []
    ids = [item.get("segment_id") for item in segments]
    add("细分ID唯一", bool(ids) and len(ids) == len(set(ids)), ids)
    forbidden_analysis_keys = {"recommendations", "external_evidence", "external_research_threshold"}
    found_analysis_keys = forbidden_analysis_keys.intersection(analysis)
    found_segment_keys = sorted({key for item in segments for key in forbidden_analysis_keys.intersection(item)})
    add("分析结构仅含评论洞察", not found_analysis_keys and not found_segment_keys,
        {"analysis": sorted(found_analysis_keys), "segments": found_segment_keys})
    cache_dir = os.path.join(workdir, "cache")
    forbidden_files = sorted(name for name in FORBIDDEN_CACHE_FILES if os.path.exists(os.path.join(cache_dir, name)))
    add("缓存仅保留评论分析产物", not forbidden_files, forbidden_files)

    html_files = glob.glob(os.path.join(workdir, "output", "*.html"))
    xlsx_files = glob.glob(os.path.join(workdir, "output", "*.xlsx"))
    add("双格式输出", bool(html_files) and bool(xlsx_files), "html=%s xlsx=%s" % (len(html_files), len(xlsx_files)))
    if html_files:
        with open(html_files[0], "r", encoding="utf-8") as handle:
            text = handle.read()
        headings = [re.sub(r"<[^>]+>", "", value).strip()
                    for value in re.findall(r"<h2[^>]*>(.*?)</h2>", text, re.S)]
        add("HTML六模块", headings == EXPECTED_HEADINGS, headings)
        duplicates = [segment_id for segment_id in ids if text.count('data-segment-id="%s"' % segment_id) != 1]
        add("所有细分只渲染一次", not duplicates, duplicates)
        events = common.read_json(common.cache_path(workdir, "events.json")).get("events") or []
        raw_hits = find_full_review_embeds(events, text)
        add("HTML不嵌完整评论", not raw_hits, raw_hits[:8])
        forbidden_terms = [term for term in ("AI建议", "AI产品建议", "外部补全", "外部证据", "Alexa", "Customer Say") if term in text]
        add("HTML只呈现评论洞察", not forbidden_terms, forbidden_terms)
    if xlsx_files:
        from openpyxl import load_workbook
        workbook = load_workbook(xlsx_files[0], read_only=True, data_only=True)
        add("Excel十二张评论洞察表", workbook.sheetnames == EXPECTED_SHEETS, workbook.sheetnames)
        workbook.close()
    cache_dir = os.path.join(workdir, "cache")
    forbidden_files = sorted(name for name in ("raw_alexa.json", "cleaned_alexa.json", "external_enrichment.json",
                                                   "manual_external_sources.json", "recommendations.json")
                             if os.path.exists(os.path.join(cache_dir, name)))
    add("无非评论输入或补全产物", not forbidden_files, forbidden_files)
    if xlsx_files:
        from openpyxl import load_workbook
        workbook = load_workbook(xlsx_files[0], read_only=True, data_only=True)
        sheetnames = workbook.sheetnames
        workbook.close()
        add("Excel十二张评论洞察表", sheetnames == EXPECTED_SHEETS, sheetnames)

    passed = sum(1 for item in checks if item["pass"])
    report = {"schema": "amazon-voc-validation/v4", "passed": passed, "total": len(checks),
              "all_green": passed == len(checks), "checks": checks}
    common.write_json(common.cache_path(workdir, "validate_report.json"), report)
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workdir", required=True)
    args = parser.parse_args()
    report = validate(args.workdir)
    for item in report["checks"]:
        print("[validate] %s %s - %s" % ("PASS" if item["pass"] else "FAIL", item["check"], item["detail"]))
    if not report["all_green"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
