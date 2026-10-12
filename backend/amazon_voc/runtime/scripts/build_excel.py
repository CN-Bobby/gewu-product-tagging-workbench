# -*- coding: utf-8 -*-
"""Build the V4 fixed twelve-sheet review evidence workbook."""
import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common
from compile_analysis import compile_analysis
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADER_FILL = PatternFill("solid", fgColor="244A73")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def scalar(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else value


def write_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([scalar(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(48, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
        for cell in column[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return sheet


def build_excel(workdir):
    path = common.cache_path(workdir, "analysis.json")
    data = common.read_json(path) if os.path.exists(path) else compile_analysis(workdir)
    master = common.read_json(common.cache_path(workdir, "master.json"))
    cleaned = common.read_json(common.cache_path(workdir, "cleaned_reviews.json"))
    events = common.read_json(common.cache_path(workdir, "events.json")).get("events") or []
    evidence = common.read_json(common.cache_path(workdir, "evidence_index.json"))
    processing_path = common.cache_path(workdir, "review_processing.json")
    processing = common.read_json(processing_path) if os.path.exists(processing_path) else {"reviews": []}
    dictionary_path = common.cache_path(workdir, "tag_dictionary.json")
    dictionary = common.read_json(dictionary_path) if os.path.exists(dictionary_path) else {"tags": [], "mappings": []}
    quality_path = common.cache_path(workdir, "quality_report.json")
    quality = common.read_json(quality_path) if os.path.exists(quality_path) else {"gates": []}

    workbook = Workbook()
    workbook.remove(workbook.active)
    task = data.get("task") or {}
    write_sheet(workbook, "任务与方法说明", ["项目", "内容"], [
        ["任务ID", task.get("task_id")], ["目标站点", task.get("marketplace")],
        ["目标国家", task.get("country")], ["品类", task.get("category")],
        ["有效评论", task.get("review_count")], ["评论ASIN", task.get("review_asin_count")],
        ["数据来源", "本次上传的单个评论Excel工作簿"],
        ["证据范围", "分析事实、引文与评论频率只来自清洗后的评论行"],
        ["细分主键", "人群 × 场景 × 动机 × 任务"],
        ["正式细分频率门槛", data.get("method", {}).get("minimum_formal_segment_reviews")],
        ["类KANO边界", data.get("method", {}).get("kano_boundary")],
        ["事件数量", (data.get("counts") or {}).get("events", len(events))],
        ["动态标签数量", (data.get("counts") or {}).get("tags", len(dictionary.get("tags") or []))],
        ["未归并原始表达标签数量", (data.get("counts") or {}).get("raw_unmerged_tags", 0)],
        ["模型", data.get("method", {}).get("models")],
        ["Prompt版本", data.get("method", {}).get("prompt_versions")],
        ["分析模式", data.get("analysis_mode", "multi_stage_v4")],
        ["样本总览（中英）", data.get("overview")],
        ["方法摘要（中英）", data.get("method_note")],
        ["人工语义抽查", data.get("method", {}).get("human_semantic_review", "" )],
        ["主事件评论互斥归属", "细分主归属评论数按每条评论的主事件统计；细分涉及评论可因多事件而重叠。"],
    ])
    write_sheet(workbook, "产品主表",
                ["ASIN", "标题", "排名", "评论数"],
                [[item.get("asin"), item.get("title"), item.get("rank"), item.get("review_count")]
                 for item in master.get("products") or []])
    write_sheet(workbook, "清洗后评论",
                ["评论ID", "ASIN", "评分", "日期", "变体", "情感", "标题", "原文"],
                [[item.get("review_id"), item.get("asin"), item.get("rating"), item.get("review_date"),
                  item.get("variant"), item.get("sentiment"), item.get("review_title"), item.get("content")]
                 for item in cleaned.get("rows") or []])
    write_sheet(workbook, "评论处理覆盖", ["评论ID", "处理状态", "原因", "事件数", "主事件ID", "互斥主细分ID", "主旅程阶段ID"],
                [[item.get("review_id"), item.get("status"), item.get("reason"), item.get("event_count"),
                  item.get("primary_event_id"), item.get("primary_segment_id"), item.get("journey_stage_id")]
                 for item in processing.get("reviews") or []])
    write_sheet(workbook, "语义事件",
                ["事件ID", "评论ID", "ASIN", "证据片段", "人群", "Buyer", "角色关系", "物理限制", "专业度",
                 "场景", "地点", "环境压力", "社交情绪背景", "旅程动作", "明确动机", "推理动机", "推理依据",
                 "核心任务", "期望结果", "产品对象", "部件", "属性", "结果", "痛点", "痒点", "情感", "严重度",
                 "Workaround", "补偿动作", "策略", "恢复目标", "事实推理边界", "置信度", "细分ID", "旅程阶段ID"],
                [[item.get("event_id"), item.get("review_id"), item.get("asin"), item.get("evidence_quote"),
                  common.nested(item, "audience", "user"), common.nested(item, "audience", "buyer"),
                  common.nested(item, "audience", "relationship"), common.nested(item, "audience", "physical_constraints"),
                  common.nested(item, "audience", "skill_level"), common.nested(item, "scenario", "context"),
                  common.nested(item, "scenario", "place"), common.nested(item, "scenario", "environmental_stressors"),
                  common.nested(item, "scenario", "social_emotional_context"), common.nested(item, "scenario", "journey_action"),
                  common.nested(item, "motivation", "explicit"), common.nested(item, "motivation", "inferred"),
                  common.nested(item, "motivation", "inference_basis"), common.nested(item, "task", "core_job"),
                  common.nested(item, "task", "desired_outcomes"), common.nested(item, "experience", "object"),
                  common.nested(item, "experience", "component"), common.nested(item, "experience", "attribute"),
                  common.nested(item, "experience", "result"), common.nested(item, "experience", "pain"),
                  common.nested(item, "experience", "itch"), common.nested(item, "experience", "sentiment"),
                  common.nested(item, "experience", "severity"), common.nested(item, "workaround", "exists"),
                  common.nested(item, "workaround", "standardized_action"), common.nested(item, "workaround", "strategy"),
                  common.nested(item, "workaround", "recovery_goal"), item.get("fact_inference"), item.get("confidence"),
                  item.get("segment_id"), item.get("journey_stage_id")]
                 for item in events])

    write_sheet(workbook, "动态标签词典",
                ["标签ID", "命名空间", "标准名称（中文）", "Standard name (English)",
                 "定义（中文）", "Definition (English)", "父标签", "原始变体", "包含例", "排除例",
                 "支持事件ID", "独立评论数", "置信度", "状态", "归一状态"],
                [[item.get("tag_id"), item.get("namespace"), item.get("canonical_name"), item.get("canonical_name_en"),
                  item.get("definition"), item.get("definition_en"),
                  item.get("parent_tag_id"), item.get("raw_variants"), item.get("include_examples"),
                  item.get("exclude_examples"), item.get("support_event_ids"), item.get("independent_review_count"),
                  item.get("confidence"), item.get("status"), item.get("normalization_status", "normalized")]
                 for item in dictionary.get("tags") or []])
    write_sheet(workbook, "事件标签映射", ["事件ID", "字段路径", "原始值", "标签ID", "映射方式"],
                [[item.get("event_id"), item.get("field_path"), item.get("raw_value"), item.get("tag_id"),
                  item.get("mapping_method", "model_normalized")]
                 for item in dictionary.get("mappings") or []])

    segment_rows = []
    need_rows = []
    for segment in data.get("segments") or []:
        key = segment.get("segment_key") or {}
        metrics = segment.get("metrics") or {}
        key_en = segment.get("segment_key_en") or {}
        segment_rows.append([segment.get("segment_id"), segment.get("name"), segment.get("name_en"), segment.get("status"),
                             key.get("audience"), key_en.get("audience"), key.get("scenario"), key_en.get("scenario"),
                             key.get("motivation"), key_en.get("motivation"), key.get("task"), key_en.get("task"),
                             metrics.get("review_count"), metrics.get("review_coverage"), metrics.get("event_count"),
                             metrics.get("asin_count"), metrics.get("primary_review_count"),
                             segment.get("persona_summary"), segment.get("persona_summary_en"),
                             segment.get("jtbd_statement"), segment.get("jtbd_statement_en"),
                             segment.get("merge_reason"), segment.get("merge_reason_en"),
                             segment.get("split_reason"), segment.get("split_reason_en")])
        for dimension, value in (segment.get("profile_synthesis") or {}).items():
            value = value or {}
            need_rows.append([segment.get("segment_id"), segment.get("name"), segment.get("name_en"),
                              dimension, value.get("zh"), value.get("en"), "模型基于评论画像抽象合并 / LLM synthesis"])
        for item in segment.get("need_priorities") or []:
            need_rows.append([segment.get("segment_id"), segment.get("name"), segment.get("name_en"), "类KANO推断 / Kano-like inference",
                              item.get("need"), item.get("basis"), item.get("priority_type"), item.get("confidence")])
        if segment.get("decision_value"):
            value = segment.get("decision_value") or {}
            need_rows.append([segment.get("segment_id"), segment.get("name"), segment.get("name_en"), "决策价值 / Decision value",
                              value.get("zh"), value.get("en"), "评论证据上的综合 / Review-based synthesis"])
    write_sheet(workbook, "正式与探索细分",
                ["细分ID", "名称（中文）", "Segment name (English)", "状态 / Status",
                 "人群（中文）", "Audience (English)", "场景（中文）", "Scenario (English)",
                 "动机（中文）", "Motivation (English)", "任务（中文）", "Task (English)",
                 "涉及评论（可重叠）", "评论出现率", "事件数", "ASIN数", "互斥主归属评论",
                 "画像摘要（中文）", "Persona summary (English)", "JTBD（中文）", "Core job (English)",
                 "合并理由（中文）", "Merge rationale (English)", "拆分理由（中文）", "Split rationale (English)"], segment_rows)
    for field, label in (("segment_comparison", "细分差异与取舍 / Segment comparison"),
                         ("cross_cutting", "跨细分共性与冲突 / Cross-segment patterns"),
                         ("exploratory_signals", "探索性信号 / Exploratory signal"),
                         ("limitations", "报告限制 / Limitation")):
        for item in data.get(field) or []:
            need_rows.append(["", "", "", label, item.get("zh"), item.get("en"), "VOC V2 中英双语综合"])
    write_sheet(workbook, "细分需求",
                ["细分ID", "细分名称（中文）", "Segment name (English)", "维度 / Dimension",
                 "模型归纳（中文）", "Model synthesis (English)", "边界与来源 / Basis"], need_rows)

    stage_names = {item.get("stage_id"): item for item in data.get("journey", {}).get("stages") or []}
    write_sheet(workbook, "动态旅程",
                ["细分ID", "阶段ID", "阶段（中文）", "Stage (English)", "阶段定义（中文）", "Stage definition (English)",
                 "用户需要（中文）", "User need (English)", "摩擦或结果（中文）", "Friction or outcome (English)",
                 "状态 / Status", "独立评论", "事件数", "动作", "目标", "痛点", "痒点", "Workaround"],
                [[item.get("segment_id"), item.get("stage_id"),
                  (stage_names.get(item.get("stage_id")) or {}).get("name"),
                  (stage_names.get(item.get("stage_id")) or {}).get("name_en"),
                  (stage_names.get(item.get("stage_id")) or {}).get("definition"),
                  (stage_names.get(item.get("stage_id")) or {}).get("definition_en"),
                  common.nested(stage_names.get(item.get("stage_id")) or {}, "user_need", "zh"),
                  common.nested(stage_names.get(item.get("stage_id")) or {}, "user_need", "en"),
                  common.nested(stage_names.get(item.get("stage_id")) or {}, "friction_or_outcome", "zh"),
                  common.nested(stage_names.get(item.get("stage_id")) or {}, "friction_or_outcome", "en"), item.get("status"),
                  item.get("review_count"), item.get("event_count"), item.get("actions"), item.get("goals"),
                  item.get("pains"), item.get("itches"), item.get("workarounds")]
                 for item in data.get("journey", {}).get("cells") or []])
    write_sheet(workbook, "证据索引",
                ["事件ID", "评论ID", "ASIN", "评分", "日期", "证据片段", "完整原文"],
                [[event_id, item.get("review_id"), item.get("asin"), item.get("rating"), item.get("review_date"),
                  item.get("evidence_quote"), item.get("review_text")] for event_id, item in evidence.items()])
    quality_rows = []
    for gate in quality.get("gates") or []:
        for item in gate.get("checks") or []:
            quality_rows.append([gate.get("gate_id"), gate.get("name"), gate.get("blocking"),
                                 item.get("name"), item.get("pass"), item.get("detail") or item.get("examples") or item.get("missing_count")])
    write_sheet(workbook, "质量门禁", ["层级", "门禁", "阻断", "检查", "通过", "详情"], quality_rows)

    common.ensure_dirs(workdir)
    output = common.output_path(workdir, "VOC_消费者需求分析_%s.xlsx" % (task.get("task_id") or "TASK"))
    workbook.save(output)
    return output


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workdir", required=True)
    args = parser.parse_args()
    print("[build-excel] %s" % build_excel(args.workdir))


if __name__ == "__main__":
    main()
