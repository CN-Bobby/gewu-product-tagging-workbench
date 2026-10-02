from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path


NAMESPACES = [
    "audience", "scenario", "motivation", "task", "object", "component", "attribute",
    "result", "pain", "itch", "desired_outcome", "workaround_action",
    "workaround_strategy", "journey_action",
]

FIELD_NAMESPACES = {
    "audience.user": "audience", "audience.buyer": "audience", "audience.relationship": "audience",
    "audience.physical_constraints": "audience", "audience.skill_level": "audience", "audience.role_needs": "audience",
    "scenario.context": "scenario", "scenario.place": "scenario", "scenario.environmental_stressors": "scenario",
    "scenario.social_emotional_context": "scenario", "scenario.trigger": "scenario",
    "scenario.journey_action": "journey_action", "motivation.explicit": "motivation",
    "motivation.inferred": "motivation", "task.core_job": "task", "task.functional_job": "task",
    "task.emotional_job": "task", "task.social_job": "task", "task.desired_outcomes": "desired_outcome",
    "experience.object": "object", "experience.component": "component", "experience.attribute": "attribute",
    "experience.result": "result", "experience.pain": "pain", "experience.itch": "itch",
    "workaround.standardized_action": "workaround_action", "workaround.strategy": "workaround_strategy",
}


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def iter_values(event):
    for path, namespace in FIELD_NAMESPACES.items():
        value = event
        for key in path.split("."):
            value = value.get(key) if isinstance(value, dict) else None
        values = value if isinstance(value, list) else [value]
        for raw in values:
            if str(raw or "").strip():
                yield path, str(raw).strip(), namespace


def _gate(gate_id, name, passed, checks, blocking=True):
    return {"gate_id": gate_id, "name": name, "pass": bool(passed), "blocking": blocking,
            "checks": checks, "failed_count": sum(1 for row in checks if not row.get("pass"))}


def evaluate(workdir: Path, include_outputs=False):
    cache = Path(workdir) / "cache"
    read = lambda name, default: json.loads((cache / name).read_text(encoding="utf-8")) if (cache / name).exists() else default
    cleaned = read("cleaned_reviews.json", {"rows": []}).get("rows") or []
    master = read("master.json", {})
    events_doc = read("events.json", {})
    events = events_doc.get("events") or []
    processing = read("review_processing.json", {"reviews": []}).get("reviews") or []
    dictionary = read("tag_dictionary.json", {"tags": [], "mappings": []})
    tags, mappings = dictionary.get("tags") or [], dictionary.get("mappings") or []
    segments_doc = read("segment_model.json", {})
    semantic_review = read("semantic_review.json", {"selected_event_ids": [], "decisions": []})
    segments = segments_doc.get("segments") or []
    journey = segments_doc.get("journey") or {}
    review_by_id = {r.get("review_id"): r for r in cleaned}
    event_ids = [e.get("event_id") for e in events]

    structure_checks = [
        {"name": "events schema", "pass": events_doc.get("schema") == "amazon-voc-events/v4"},
        {"name": "event ids unique", "pass": bool(event_ids) and len(event_ids) == len(set(event_ids)) and all(event_ids)},
        {"name": "required event groups", "pass": all(set(e) >= {"audience", "scenario", "motivation", "task", "experience", "workaround"} for e in events)},
        {"name": "generation metadata", "pass": all((events_doc.get("metadata") or {}).get(k) for k in ("model", "prompt_version", "generated_at"))},
    ]
    processing_by_id = {row.get("review_id"): row for row in processing}
    coverage_checks = [
        {"name": "one workbook is the only source", "pass": set((master.get("source_files") or {}).keys()) == {"reviews"},
         "detail": sorted((master.get("source_files") or {}).keys())},
        {"name": "every cleaned review accounted", "pass": set(review_by_id) == set(processing_by_id),
         "detail": {"cleaned": len(review_by_id), "processed": len(processing_by_id)}},
        {"name": "valid processing status", "pass": all(row.get("status") in ("已生成事件", "无可用需求信号", "失败") for row in processing)},
        {"name": "failed reviews absent", "pass": not any(row.get("status") == "失败" for row in processing)},
    ]
    bad_evidence = []
    for event in events:
        source = (review_by_id.get(event.get("review_id")) or {}).get("content", "")
        if not event.get("evidence_quote") or _norm(event.get("evidence_quote")) not in _norm(source):
            bad_evidence.append(event.get("event_id"))
    evidence_checks = [{"name": "quotes locate in source", "pass": not bad_evidence, "examples": bad_evidence[:20]}]

    semantic_bad = []
    for event in events:
        inferred = (event.get("motivation") or {}).get("inferred")
        basis = (event.get("motivation") or {}).get("inference_basis") or []
        if inferred and not basis:
            semantic_bad.append(event.get("event_id"))
        if event.get("review_status") in ("unsupported_inference", "non_atomic"):
            semantic_bad.append(event.get("event_id"))
    selected_ids = set(semantic_review.get("selected_event_ids") or [])
    decided_ids = {row.get("event_id") for row in semantic_review.get("decisions") or []}
    semantic_checks = [
        {"name": "atomic and inference-bounded", "pass": not semantic_bad,
         "examples": list(dict.fromkeys(semantic_bad))[:20]},
        {"name": "selected semantic reviews completed",
         "pass": selected_ids <= decided_ids and bool(semantic_review.get("complete")) and
                 not semantic_review.get("missing_event_ids") and not semantic_review.get("unresolved_quote_event_ids"),
         "missing_count": len(selected_ids - decided_ids),
         "unresolved_quote_count": len(semantic_review.get("unresolved_quote_event_ids") or [])},
        {"name": "review sample exists", "pass": bool(selected_ids) if events else True,
         "detail": {"selected": len(selected_ids), "events": len(events)}},
    ]

    tag_ids = [t.get("tag_id") for t in tags]
    mapping_counter = Counter((m.get("event_id"), m.get("field_path"), _norm(m.get("raw_value"))) for m in mappings)
    mapping_keys = set(mapping_counter)
    expected = {(e.get("event_id"), path, _norm(raw)) for e in events for path, raw, _ in iter_values(e)}
    supported = {eid for tag in tags for eid in tag.get("support_event_ids") or []}
    tag_by_id = {t.get("tag_id"): t for t in tags}
    dictionary_checks = [
        {"name": "tag ids unique", "pass": bool(tag_ids) and len(tag_ids) == len(set(tag_ids)) and all(tag_ids)},
        {"name": "tag fields complete", "pass": all(t.get("namespace") in NAMESPACES and t.get("canonical_name") and t.get("definition") and t.get("support_event_ids") for t in tags)},
        {"name": "all values mapped", "pass": expected <= mapping_keys, "missing_count": len(expected - mapping_keys)},
        {"name": "one mapping per value", "pass": all(mapping_counter[key] == 1 for key in expected),
         "duplicate_count": sum(max(0, mapping_counter[key] - 1) for key in expected)},
        {"name": "mappings reference tags", "pass": all(m.get("tag_id") in set(tag_ids) for m in mappings)},
        {"name": "mapping namespace matches field", "pass": all(
            (tag_by_id.get(m.get("tag_id")) or {}).get("namespace") == FIELD_NAMESPACES.get(m.get("field_path"))
            for m in mappings)},
        {"name": "no orphan tags", "pass": set(event_ids) >= supported and all(t.get("support_event_ids") for t in tags)},
    ]

    segment_members = [eid for s in segments for eid in s.get("member_event_ids") or []]
    stage_members = [a.get("event_id") for a in journey.get("event_assignments") or []]
    assignment_checks = [
        {"name": "one segment per event", "pass": Counter(segment_members) == Counter({eid: 1 for eid in event_ids})},
        {"name": "one journey stage per event", "pass": Counter(stage_members) == Counter({eid: 1 for eid in event_ids})},
        {"name": "segment key is four-dimensional", "pass": all(set(s.get("segment_key") or {}) == {"audience", "scenario", "motivation", "task"} for s in segments)},
    ]

    output_checks = [{"name": "output consistency deferred", "pass": True, "detail": "编译完成后复核"}]
    if include_outputs:
        analysis = read("analysis.json", {})
        html_files = list((Path(workdir) / "output").glob("*.html"))
        xlsx_files = list((Path(workdir) / "output").glob("*.xlsx"))
        html_text = html_files[0].read_text(encoding="utf-8") if html_files else ""
        html_counts = re.search(r'data-counts="(\d+),(\d+),(\d+)"', html_text)
        expected_counts = (len(cleaned), len(events), len(segments))
        html_tuple = tuple(map(int, html_counts.groups())) if html_counts else None
        excel_tuple = None
        excel_sheets = []
        if xlsx_files:
            from openpyxl import load_workbook
            book = load_workbook(xlsx_files[0], read_only=True, data_only=True)
            excel_sheets = book.sheetnames
            excel_tuple = (book["清洗后评论"].max_row - 1, book["语义事件"].max_row - 1,
                           book["正式与探索细分"].max_row - 1)
            book.close()
        html_headings = [re.sub(r"<[^>]+>", "", value).strip()
                         for value in re.findall(r"<h2[^>]*>(.*?)</h2>", html_text, re.S)]
        expected_headings = ["01 · 数据与方法概览", "02 · 全部细分总览", "03 · 全部细分详情",
                             "04 · 细分横向比较", "05 · 动态用户旅程", "06 · 探索性信号"]
        expected_sheets = ["任务与方法说明", "产品主表", "清洗后评论", "评论处理覆盖", "语义事件", "动态标签词典",
                           "事件标签映射", "正式与探索细分", "细分需求", "动态旅程", "证据索引", "质量门禁"]
        forbidden = [term for term in ("AI建议", "外部补全", "外部证据", "Alexa", "Customer Say") if term in html_text]
        forbidden_outputs = {"recommendations", "external_evidence", "external_research_threshold"}.intersection(analysis)
        legacy_files = {"raw_alexa.json", "cleaned_alexa.json", "external_enrichment.json",
                        "manual_external_sources.json", "recommendations.json"}
        present_legacy_files = sorted(name for name in legacy_files if (cache / name).exists())
        output_checks = [
            {"name": "analysis counts", "pass": len(analysis.get("segments") or []) == len(segments)},
            {"name": "review-only evidence scope", "pass": (analysis.get("method") or {}).get("evidence_scope") == "uploaded_review_workbook_only"},
            {"name": "analysis has no external/advice payload", "pass": not forbidden_outputs, "detail": sorted(forbidden_outputs)},
            {"name": "HTML counts", "pass": html_tuple == expected_counts, "detail": {"expected": expected_counts, "actual": html_tuple}},
            {"name": "Excel counts", "pass": excel_tuple == expected_counts, "detail": {"expected": expected_counts, "actual": excel_tuple}},
            {"name": "both formats exist", "pass": bool(html_files and xlsx_files)},
            {"name": "six HTML modules", "pass": html_headings == expected_headings, "detail": html_headings},
            {"name": "twelve Excel sheets", "pass": excel_sheets == expected_sheets, "detail": excel_sheets},
            {"name": "no external/advice output", "pass": not forbidden and not present_legacy_files,
             "detail": {"html_terms": forbidden, "legacy_cache_files": present_legacy_files}},
        ]

    gates = [
        _gate(1, "结构门禁", all(c["pass"] for c in structure_checks), structure_checks),
        _gate(2, "处理覆盖", all(c["pass"] for c in coverage_checks), coverage_checks),
        _gate(3, "证据门禁", all(c["pass"] for c in evidence_checks), evidence_checks),
        _gate(4, "语义门禁", all(c["pass"] for c in semantic_checks), semantic_checks),
        _gate(5, "词典门禁", all(c["pass"] for c in dictionary_checks), dictionary_checks),
        _gate(6, "细分与旅程门禁", all(c["pass"] for c in assignment_checks), assignment_checks),
        _gate(7, "输出一致性", all(c["pass"] for c in output_checks), output_checks),
    ]
    report = {"schema": "amazon-voc-quality/v4", "all_green": all(g["pass"] for g in gates),
              "precompile_green": all(g["pass"] for g in gates[:6]),
              "counts": {"reviews": len(cleaned), "events": len(events), "segments": len(segments), "tags": len(tags)},
              "gates": gates}
    (cache / "quality_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report
