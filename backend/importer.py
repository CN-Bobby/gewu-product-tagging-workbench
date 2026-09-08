"""Preserve source bytes and values, including WPS DISPIMG cell-image relationships."""
import hashlib
import io
import json
import posixpath
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET
import openpyxl
from PIL import Image
from sqlalchemy import select, func
from .db import Session, Dataset, Product, Asset, Workspace, STORAGE, uid, obj, audit, require

ALIASES = {
 'asin': ['ASIN', 'asin', 'ASIN码'], 'sku': ['SKU', 'sku'],
 'title': ['商品标题', '标题', '商品名称', '产品名称', 'Product Title', 'Title', 'title', 'product_name'],
 'bullets': ['产品卖点', '五点描述', 'Bullet Points', 'bullets'],
 'parameters': ['详细信息', '详细参数', '参数', 'Details'], 'brand': ['品牌', 'Brand'],
 'parent_asin': ['父ASIN', 'Parent ASIN'], 'price': ['价格($)', '价格', 'Price'],
 'sales': ['月销量', 'Sales'], 'rating': ['评分', 'Rating'],
 'reviews': ['评分数', '评论数'], 'image_url': ['商品主图', '主图链接', '图片链接', 'Image URL'],
 'category': ['类目路径', '小类目', 'Category'], 'url': ['商品详情页链接', '详情页链接', 'URL'],
}
def clean(v):
    if isinstance(v, (datetime, date)): return v.isoformat()
    return v

def image_bytes(data):
    if len(data) > 25 * 1024 * 1024: raise ValueError('单张图片超过 25 MB')
    with Image.open(io.BytesIO(data)) as im:
        im.verify()
    digest = hashlib.sha256(data).hexdigest()
    with Image.open(io.BytesIO(data)) as im: ext = (im.format or 'PNG').lower()
    target = STORAGE / 'assets' / f'{digest}.{ext}'
    target.parent.mkdir(exist_ok=True)
    if not target.exists(): target.write_bytes(data)
    return str(target), digest

def wps_images(z):
    if 'xl/cellimages.xml' not in z.namelist(): return {}
    rel = ET.fromstring(z.read('xl/_rels/cellimages.xml.rels'))
    targets = {x.get('Id'):posixpath.normpath(posixpath.join('xl', x.get('Target', ''))).lstrip('/') for x in rel if x.get('TargetMode') != 'External'}
    result = {}
    for node in ET.fromstring(z.read('xl/cellimages.xml')):
        name = next((x.get('name') for x in node.iter() if x.tag.endswith('}cNvPr')), None)
        rid = next((x.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') for x in node.iter() if x.tag.endswith('}blip')), None)
        if name and rid in targets and targets[rid] in z.namelist(): result[name] = z.read(targets[rid])
    return result

def parse(data, sheet=None):
    if len(data) > 100 * 1024 * 1024: raise ValueError('表格超过 100 MB')
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(x.file_size for x in z.infolist()) > 500 * 1024 * 1024: raise ValueError('表格解压后过大')
        cells = wps_images(z)
        rich = any('richData/' in n for n in z.namelist())
    book = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    if sheet and sheet not in book.sheetnames: raise ValueError('工作表不存在')
    ws = book[sheet] if sheet else book.active
    if ws.max_row > 100001 or ws.max_column > 500: raise ValueError('支持至多 100000 行、500 列')
    headers = [str(v) if v is not None else f'未命名列_{i+1}' for i,v in enumerate(next(ws.values))]
    if len(set(headers)) != len(headers): raise ValueError('表头有重复列名，请先为重复列命名')
    mapping = {key: next((h for h in aliases if h in headers), None) for key, aliases in ALIASES.items()}
    floating = {}
    for im in ws._images:
        if hasattr(im.anchor, '_from'): floating.setdefault(im.anchor._from.row + 1, []).append(im._data())
    rows, warnings = [], []
    if rich: warnings.append('检测到 Excel richData 图片；当前仅保证 WPS DISPIMG 和浮动图片解析，请核对图片覆盖率。')
    for row in ws.iter_rows(min_row=2):
        values = [clean(c.value) for c in row]
        if all(x is None for x in values): continue
        images = list(floating.get(row[0].row, []))
        for cell in row:
            if isinstance(cell.value, str) and 'DISPIMG(' in cell.value:
                match = re.search(r'DISPIMG\(\s*"([^"]+)"', cell.value)
                if match and match[1] in cells: images.append(cells[match[1]])
                else: warnings.append(f'{cell.coordinate} 的内嵌图片未找到')
        raw = dict(zip(headers, values))
        fields = {}
        for key, header in mapping.items():
            if header:
                cell = row[headers.index(header)]
                fields[key] = clean(cached[ws.title][cell.coordinate].value) if cell.data_type == 'f' else raw[header]
        anchor = str(fields.get('asin') or fields.get('sku') or f'ROW-{row[0].row}')
        rows.append({'row': row[0].row, 'anchor':anchor, 'raw':raw, 'fields':fields, 'images':images})
    if not rows: raise ValueError('工作表中没有商品数据')
    dupes = [k for k,v in Counter(r['anchor'] for r in rows).items() if v > 1]
    if dupes: warnings.append(f'存在重复商品标识，保留独立行：{", ".join(dupes[:10])}')
    if not mapping['asin'] and not mapping['sku']: warnings.append('缺少 ASIN / SKU，使用原始行号作为临时标识')
    from .field_catalog import infer
    fields=infer(headers,[r['raw'] for r in rows])
    for f in fields:
        if f['status']=='recognized' and f['role'] in mapping and not mapping[f['role']]:mapping[f['role']]=f['name']
    for r in rows:
        for key,h in mapping.items():
            if h and key not in r['fields']:r['fields'][key]=r['raw'].get(h)
        r['anchor']=str(r['fields'].get('asin') or r['fields'].get('sku') or r['fields'].get('parent_asin') or f"ROW-{r['row']}")
    return {'sheets':book.sheetnames, 'sheet':ws.title, 'headers':headers, 'mapping':mapping, 'rows':rows, 'warnings':warnings}

def inspect(data, sheet=None):
    p = parse(data, sheet)
    return {**{k:v for k,v in p.items() if k!='rows'}, 'count':len(p['rows']), 'embedded_images':sum(len(r['images']) for r in p['rows']), 'sample':[{k:v for k,v in r.items() if k!='images'} for r in p['rows'][:3]]}

def import_file(data, filename, sheet=None, workspace_id=None, mapping=None):
    p = parse(data, sheet)
    if mapping:
        if any(v not in p['headers'] for v in mapping.values() if v): raise ValueError('映射列不存在')
        p['mapping'].update(mapping)
        for row in p['rows']:
            for key, header in mapping.items(): row['fields'][key] = row['raw'].get(header) if header else None
            row['anchor'] = str(row['fields'].get('asin') or row['fields'].get('sku') or f"ROW-{row['row']}")
    ident = uid()
    original = STORAGE / 'originals' / f'{ident}.xlsx'
    original.parent.mkdir(exist_ok=True)
    original.write_bytes(data)
    with Session.begin() as s:
        if workspace_id: workspace = require(s, Workspace, workspace_id)
        else:
            workspace = Workspace(name=Path(filename).stem)
            s.add(workspace); s.flush()
        version = (s.scalar(select(func.max(Dataset.version)).where(Dataset.workspace_id==workspace.id)) or 0)+1
        dataset = Dataset(id=ident, workspace_id=workspace.id, name=Path(filename).stem, version=version, original_file=str(original), sha256=hashlib.sha256(data).hexdigest(), sheet=p['sheet'], headers=p['headers'], mapping=p['mapping'], warnings=p['warnings'])
        s.add(dataset); s.flush()
        for row in p['rows']:
            product = Product(dataset_id=ident, **{k:row[k] for k in ['row','anchor','raw','fields']})
            s.add(product); s.flush()
            seen=set()
            for data_image in row['images']:
                path, digest=image_bytes(data_image)
                if digest not in seen: s.add(Asset(product_id=product.id,source='excel',path=path,sha256=digest))
                seen.add(digest)
            url=row['fields'].get('image_url')
            if isinstance(url,str) and url.startswith(('https://','http://')):
                s.add(Asset(product_id=product.id, source='url', url=url))
        audit(s, ident, 'dataset_imported', {'rows':len(p['rows']), 'sha256':dataset.sha256,'format':'WPS DISPIMG / DrawingML'})
        return obj(dataset)

def attach_images(dataset_id, files):
    matched, unmatched=0, []
    with Session.begin() as s:
        ds=require(s,Dataset,dataset_id)
        if ds.stage!='DATA_READY': raise ValueError('开始分析后图片输入已冻结；请导入新数据版本补充图片')
        products=s.scalars(select(Product).where(Product.dataset_id==dataset_id)).all()
        for name, data in files:
            stem=Path(name).stem
            targets=[p for p in products if stem==p.anchor or stem.startswith(p.anchor+'_') or stem.startswith(p.anchor+'-')]
            if not targets: unmatched.append(name); continue
            path,digest=image_bytes(data)
            for p in targets:
                if not s.scalar(select(Asset).where(Asset.product_id==p.id, Asset.sha256==digest)):
                    s.add(Asset(product_id=p.id,source='upload',path=path,sha256=digest)); matched+=1
        audit(s,dataset_id,'images_attached',{'matched':matched,'unmatched':unmatched})
    return {'matched':matched,'unmatched':unmatched}
