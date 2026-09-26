"""Amazon Listing image collection into an importable GeWu workbook.

The collector owns only its job state and temporary image bytes under the configured
storage root. It never writes into the source repository or an arbitrary output path.
"""
from __future__ import annotations

import io
import ipaddress
import json
import tempfile
import os
import random
import re
import shutil
import socket
import ssl
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from .db import STORAGE

MARKETPLACES: dict[str, str] = {
    "US": "amazon.com", "UK": "amazon.co.uk", "DE": "amazon.de", "FR": "amazon.fr",
    "IT": "amazon.it", "ES": "amazon.es", "CA": "amazon.ca", "JP": "amazon.co.jp",
    "IN": "amazon.in", "MX": "amazon.com.mx", "BR": "amazon.com.br", "AU": "amazon.com.au",
    "AE": "amazon.ae", "NL": "amazon.nl", "SE": "amazon.se", "PL": "amazon.pl",
    "SG": "amazon.sg", "EG": "amazon.eg", "SA": "amazon.sa", "TR": "amazon.com.tr",
}
ASIN_RE = re.compile(r"^B[0-9A-Z]{9}$", re.I)
ASIN_IN_URL = re.compile(r"/(?:dp|gp/product|product)/([A-Z0-9]{10})(?:[/?]|$)", re.I)
IMAGE_HOST_SUFFIXES = (".media-amazon.com", ".ssl-images-amazon.com")
MAX_IMAGES = 10
MAX_IMAGE_BYTES = 15 * 1024 * 1024
MAX_CONCURRENT_JOBS = 2
MAX_ACTIVE_PAGE_FETCHES = 4
MAX_INPUT_BYTES = 128 * 1024
MAX_NEW_ITEMS = 100
CAPTCHA_TEXT = ("robot check", "enter the characters you see", "type the characters you see", "captcha")
PROFILES = {
    "stable": {"workers": 1, "delay_min": 6.0, "delay_max": 10.0},
    "balanced": {"workers": 2, "delay_min": 2.0, "delay_max": 5.0},
    "fast": {"workers": 4, "delay_min": 1.0, "delay_max": 2.0},
}
MAX_RESTRICTED_RESULTS = 3
RESTRICTED_COOLDOWN_SECONDS = 30
_JOB_LOCKS_GUARD = threading.RLock()
_JOB_LOCKS: dict[str, threading.RLock] = {}


def _allowed_marketplace(host: str) -> str | None:
    host = host.lower().rstrip(".")
    for code, domain in MARKETPLACES.items():
        if host in (domain, "www." + domain):
            return code
    return None


def _safe_input_label(token: str) -> str:
    """Keep a useful validation hint without retaining URL credentials or queries."""
    try:
        parsed = urllib.parse.urlsplit(token)
        if parsed.hostname:
            return f"{parsed.hostname}{parsed.path}"[:200]
    except ValueError:
        pass
    return token[:80]


def _scrub_error(message: str, proxies: list[str] | None = None) -> str:
    text = str(message or "采集失败")
    for proxy in proxies or []:
        parsed = urllib.parse.urlsplit(proxy)
        if parsed.username is not None:
            credentials = f"{parsed.username}:{parsed.password or ''}"
            text = text.replace(credentials, "[代理凭据]").replace(urllib.parse.unquote(credentials), "[代理凭据]")
        text = text.replace(proxy, "[代理]")
    text = re.sub(r"(?i)(https?://)[^/@\s]+:[^/@\s]+@", r"\1[凭据]@", text)
    return text[:500]


def _public_dns(host: str) -> None:
    try:
        addresses = socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise ValueError(f"Amazon 商品域名无法解析：{exc}") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if not ip.is_global:
            raise ValueError("商品域名解析到非公网地址")


def parse_inputs(text: str, marketplace: str = "US") -> dict[str, Any]:
    """Parse pasted ASINs and Amazon URLs, retaining invalids for user review."""
    text = str(text or "")
    if len(text.encode("utf-8")) > MAX_INPUT_BYTES:
        raise ValueError("输入内容过大：最多 128 KiB，请分批提交")
    marketplace = str(marketplace or "US").upper()
    if marketplace not in MARKETPLACES:
        raise ValueError("不支持的 Amazon 站点")
    tokens = [value.strip() for value in re.split(r"[\s,;，；、]+", str(text or "")) if value.strip()]
    valid: list[dict[str, str]] = []
    invalid: list[str] = []
    seen: set[tuple[str, str]] = set()
    for token in tokens:
        if ASIN_RE.fullmatch(token):
            code = marketplace
            asin = token.upper()
            url = f"https://www.{MARKETPLACES[code]}/dp/{asin}"
        else:
            try:
                parsed = urllib.parse.urlsplit(token)
            except ValueError:
                invalid.append(_safe_input_label(token)); continue
            code = _allowed_marketplace(parsed.hostname or "") if parsed.scheme.lower() == "https" else None
            match = ASIN_IN_URL.search(parsed.path or "")
            if not code or not match or parsed.username or parsed.password:
                invalid.append(_safe_input_label(token)); continue
            asin = match.group(1).upper()
            url = f"https://www.{MARKETPLACES[code]}{parsed.path}"
        key = (code, asin)
        if key not in seen:
            seen.add(key)
            valid.append({"marketplace": code, "asin": asin, "url": url})
            if len(valid) > MAX_NEW_ITEMS:
                raise ValueError("单次最多提交 100 个去重后的商品，请分批提交")
    return {"items": valid, "invalid": invalid, "duplicates_removed": len(tokens) - len(valid) - len(invalid)}


def _safe_image_url(url: str) -> bool:
    try:
        parsed = urllib.parse.urlsplit(url)
    except ValueError:
        return False
    host = (parsed.hostname or "").lower().rstrip(".")
    return parsed.scheme.lower() == "https" and not parsed.username and not parsed.password and any(host.endswith(suffix) for suffix in IMAGE_HOST_SUFFIXES)


class _SafeRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        if not _safe_image_url(newurl):
            raise urllib.error.HTTPError(newurl, code, "图片重定向目标不受支持", headers, fp)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def _proxy_pool() -> list[str]:
    single = os.environ.get("AMAZON_SCRAPER_PROXY", "").strip()
    proxy_file = os.environ.get("AMAZON_SCRAPER_PROXY_FILE", "").strip()
    values = [single] if single else []
    path = Path(proxy_file).expanduser() if proxy_file else (STORAGE / "amazon-scraper" / "proxies.txt")
    if path.is_file():
        values.extend(line.strip() for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip() and not line.lstrip().startswith("#"))
    return list(dict.fromkeys(values))


def runtime_status() -> dict[str, Any]:
    try:
        proxy_count = len(_proxy_pool())
    except Exception:
        proxy_count = 0
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as playwright:
            binaries = {
                "msedge": (Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")) / "Microsoft/Edge/Application/msedge.exe",
                           Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Microsoft/Edge/Application/msedge.exe"),
                "chrome": (Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Google/Chrome/Application/chrome.exe",
                           Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")) / "Google/Chrome/Application/chrome.exe"),
            }
            channel = next((name for name, paths in binaries.items() if any(path.is_file() for path in paths)), None)
            chromium = Path(playwright.chromium.executable_path)
            installed = chromium.is_file()
            return {"available": bool(channel or installed), "playwright_installed": True,
                    "browser": channel or ("chromium" if installed else None),
                    "proxy_configured": proxy_count > 0, "proxy_count": proxy_count,
                    "message": "可使用系统浏览器" if channel else ("可使用 Playwright Chromium" if installed else "请运行 scripts/prepare-amazon-scraper-runtime.ps1 安装 Chromium")}
    except Exception as exc:
        return {"available": False, "playwright_installed": False, "browser": None,
                "proxy_configured": proxy_count > 0, "proxy_count": proxy_count,
                "message": f"缺少 Playwright 运行组件：{exc}"}


def _launch_browser(playwright, proxy: str | None):
    options: dict[str, Any] = {"headless": True}
    if proxy:
        parsed_proxy = urllib.parse.urlsplit(proxy)
        if parsed_proxy.scheme not in ("http", "https", "socks5") or not parsed_proxy.hostname:
            raise ValueError("代理 URL 格式无效")
        server = f"{parsed_proxy.scheme}://{parsed_proxy.hostname}:{parsed_proxy.port or (1080 if parsed_proxy.scheme == 'socks5' else 8080)}"
        proxy_config: dict[str, str] = {"server": server}
        if parsed_proxy.username is not None:
            proxy_config["username"] = urllib.parse.unquote(parsed_proxy.username)
            proxy_config["password"] = urllib.parse.unquote(parsed_proxy.password or "")
        options["proxy"] = proxy_config
    edge = Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")) / "Microsoft/Edge/Application/msedge.exe"
    edge2 = Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Microsoft/Edge/Application/msedge.exe"
    chrome = Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Google/Chrome/Application/chrome.exe"
    chrome2 = Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")) / "Google/Chrome/Application/chrome.exe"
    if edge.is_file() or edge2.is_file():
        return playwright.chromium.launch(channel="msedge", **options)
    if chrome.is_file() or chrome2.is_file():
        return playwright.chromium.launch(channel="chrome", **options)
    if Path(playwright.chromium.executable_path).is_file():
        return playwright.chromium.launch(**options)
    raise RuntimeError("没有可用浏览器；请运行 scripts/prepare-amazon-scraper-runtime.ps1")


def _gallery_urls(page) -> list[str]:
    raw = page.evaluate("""() => {
      const out=[];
      const add=(value)=>{ if(!value) return; try { const u=new URL(value,location.href); if(u.protocol==='https:') out.push(u.href); } catch {} };
      const hero=document.querySelector('#landingImage, #imgTagWrapperId img');
      if(hero){ add(hero.getAttribute('data-old-hires')); const d=hero.getAttribute('data-a-dynamic-image'); if(d){try{const obj=JSON.parse(d); const best=Object.entries(obj).sort((a,b)=>(b[1]?.[0]||0)*(b[1]?.[1]||0)-(a[1]?.[0]||0)*(a[1]?.[1]||0)); if(best[0]) add(best[0][0]);}catch{}} add(hero.currentSrc||hero.src); }
      for(const img of document.querySelectorAll('#altImages li img, #imageBlock img')){ add(img.getAttribute('data-old-hires')); add(img.getAttribute('src')); add(img.getAttribute('data-src')); }
      return out;
    }""")
    result: list[str] = []
    seen: set[str] = set()
    for value in raw:
        if not _safe_image_url(value):
            continue
        value = re.sub(r"\._[A-Za-z0-9,+_-]+_\.", ".", value)
        parsed = urllib.parse.urlsplit(value)
        identity = re.sub(r"\._.*", "", parsed.path)
        if identity not in seen:
            seen.add(identity); result.append(value)
        if len(result) >= MAX_IMAGES:
            break
    return result


def _fetch_product(item: dict[str, str], proxy: str | None, cancel: threading.Event) -> dict[str, Any]:
    if cancel.is_set():
        return {"status": "cancelled", "error": "用户取消"}
    _public_dns(urllib.parse.urlsplit(item["url"]).hostname or "")
    from playwright.sync_api import sync_playwright
    with sync_playwright() as playwright:
        browser = _launch_browser(playwright, proxy)
        try:
            context = browser.new_context(locale="en-US", user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/132.0.0.0 Safari/537.36",
                                          extra_http_headers={"Accept-Language": "en-US,en;q=0.9"})
            page = context.new_page()
            checked_hosts: set[str] = set()
            def guard_request(route):
                parsed_request = urllib.parse.urlsplit(route.request.url)
                host = (parsed_request.hostname or "").lower().rstrip(".")
                owned = any(host == domain or host.endswith("." + domain) for domain in MARKETPLACES.values())
                image_host = any(host.endswith(suffix) for suffix in IMAGE_HOST_SUFFIXES)
                if parsed_request.scheme != "https" or not (owned or image_host):
                    route.abort(); return
                try:
                    if host not in checked_hosts:
                        _public_dns(host); checked_hosts.add(host)
                    route.continue_()
                except Exception:
                    route.abort()
            page.route("**/*", guard_request)
            response = page.goto(item["url"], wait_until="domcontentloaded", timeout=30_000)
            final_host = urllib.parse.urlsplit(page.url).hostname or ""
            if _allowed_marketplace(final_host) is None:
                return {"status": "failed", "error": "Amazon 页面跳转到未登记域名"}
            status_code = response.status if response else 0
            if status_code in (429, 503):
                return {"status": "restricted", "error": f"Amazon 暂时限制访问（HTTP {status_code}）"}
            page.wait_for_timeout(900)
            text = (page.locator("body").inner_text(timeout=5_000) or "")[:6000].lower()
            if any(value in text for value in CAPTCHA_TEXT):
                return {"status": "needs_attention", "restricted": True, "error": "检测到 CAPTCHA，请稍后重试"}
            try:
                title_locator = page.locator("#productTitle").first
                title_locator.wait_for(state="visible", timeout=8_000)
                title = (title_locator.inner_text(timeout=3_000) or "").strip()
            except Exception:
                title = ""
            urls = _gallery_urls(page)
            if not urls:
                return {"status": "failed", "error": "页面未找到 Listing 图片"}
            images: list[tuple[bytes, str]] = []
            for url in urls:
                if cancel.is_set():
                    return {"status": "cancelled", "error": "用户取消"}
                if not _safe_image_url(url):
                    continue
                headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/132.0.0.0 Safari/537.36",
                           "Referer": item["url"]}
                handlers: list[Any] = [_SafeRedirect(), urllib.request.HTTPSHandler(context=ssl.create_default_context())]
                if proxy:
                    handlers.append(urllib.request.ProxyHandler({"http": proxy, "https": proxy}))
                opener = urllib.request.build_opener(*handlers)
                try:
                    with opener.open(urllib.request.Request(url, headers=headers), timeout=30) as response_image:
                        mime = response_image.headers.get_content_type().lower()
                        if mime not in ("image/jpeg", "image/png"):
                            continue
                        length = response_image.headers.get("Content-Length")
                        if length and int(length) > MAX_IMAGE_BYTES:
                            continue
                        data = response_image.read(MAX_IMAGE_BYTES + 1)
                    if len(data) > MAX_IMAGE_BYTES:
                        continue
                    with Image.open(io.BytesIO(data)) as image:
                        image.verify()
                    images.append((data, mime))
                except Exception:
                    continue
            if not images:
                return {"status": "failed", "error": "Listing 图片均无法下载或无法解析"}
            return {"status": "ready", "title": title[:1000], "images": images}
        finally:
            browser.close()


class AmazonScraperService:
    def __init__(self, root: Path | None = None, engine: Callable[..., dict[str, Any]] | None = None):
        self.root = (root or (STORAGE / "amazon-scraper")).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self._use_worker_process = engine is None
        self.engine = engine or _fetch_product
        self.lock = threading.RLock()
        self._job_locks: dict[str, threading.RLock] = {}
        self._fetch_condition = threading.Condition(threading.RLock())
        self._fetch_active = 0
        self._fetch_waiters: dict[str, int] = {}
        self._fetch_order: list[str] = []
        self._restricted_jobs: set[str] = set()
        self._last_start_at: dict[str, float] = {}
        self._child_processes: dict[str, set[tuple[Any, Any]]] = {}
        self.threads: dict[str, threading.Thread] = {}
        self.cancel_events: dict[str, threading.Event] = {}
        self.active_job_ids: set[str] = set()
        self.shutting_down = False
        self.proxy_rotation = 0
        self.recover()

    def _path(self, job_id: str) -> Path:
        if not re.fullmatch(r"[0-9a-f]{32}", job_id):
            raise ValueError("采集任务编号无效")
        return self.root / job_id / "job.json"

    def _read(self, job_id: str) -> dict[str, Any]:
        path = self._path(job_id)
        # Serialize readers with the atomic replacement writer. Windows may briefly deny
        # a concurrent open/replace on the same file; the per-job RLock also permits _save
        # to reuse this helper while already holding the lock.
        with self._job_lock(job_id):
            if not path.is_file():
                raise ValueError("采集任务不存在")
            job = json.loads(path.read_text(encoding="utf-8"))
            # Old checkpoints predate optimistic concurrency fields. Reading them is
            # non-mutating; their first successful write upgrades them to revision/generation 1/0.
            job.setdefault("revision", 0)
            job.setdefault("attempt_generation", 0)
            return job

    def _job_lock(self, job_id: str) -> threading.RLock:
        key = str((self.root / job_id).resolve())
        with _JOB_LOCKS_GUARD:
            return _JOB_LOCKS.setdefault(key, threading.RLock())

    def _save(self, job: dict[str, Any], expected_generation: int | None = None,
              allowed_statuses: tuple[str, ...] | None = None) -> bool:
        path = self._path(job["id"])
        path.parent.mkdir(parents=True, exist_ok=True)
        with self._job_lock(job["id"]):
            previous = json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}
            generation = int(previous.get("attempt_generation", 0))
            expected = int(job.get("attempt_generation", 0)) if expected_generation is None else expected_generation
            if previous and generation != expected:
                return False
            source_revision = int(job.get("revision", 0))
            if previous and source_revision != int(previous.get("revision", 0)):
                return False
            if allowed_statuses is not None and previous.get("status") not in allowed_statuses:
                return False
            saved = dict(job)
            saved["revision"] = int(previous.get("revision", 0)) + 1
            saved["attempt_generation"] = int(saved.get("attempt_generation", generation))
            fd, temp_name = tempfile.mkstemp(prefix=f".{job['id']}.{uuid.uuid4().hex}.", suffix=".tmp", dir=path.parent)
            try:
                with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
                    json.dump(saved, stream, ensure_ascii=False, indent=2)
                    stream.flush()
                    os.fsync(stream.fileno())
                # Windows can transiently reject replacement while a non-cooperating
                # reader (for example an indexer) has the destination open. Keep the
                # unique temp file and atomic replace semantics, retrying briefly before
                # surfacing a persistent lock/permission error.
                for attempt in range(6):
                    try:
                        os.replace(temp_name, path)
                        break
                    except PermissionError:
                        if attempt == 5:
                            raise
                        time.sleep(0.02 * (attempt + 1))
            finally:
                try:
                    os.unlink(temp_name)
                except FileNotFoundError:
                    pass
            job.update(saved)
            return True

    def preview(self, text: str, marketplace: str = "US", profile: str = "balanced") -> dict[str, Any]:
        if profile not in PROFILES:
            raise ValueError("访问档位无效")
        parsed = parse_inputs(text, marketplace)
        return {**parsed, "runtime": runtime_status(), "profile": profile,
                "estimated_pages": len(parsed["items"]), "max_images_per_product": MAX_IMAGES,
                **self._access_limits()}

    @staticmethod
    def _access_limits() -> dict[str, Any]:
        return {"max_items_per_job": MAX_NEW_ITEMS, "max_input_bytes": MAX_INPUT_BYTES,
                "max_active_page_fetches": MAX_ACTIVE_PAGE_FETCHES,
                "max_concurrent_jobs": MAX_CONCURRENT_JOBS,
                "profile_page_caps": {name: settings["workers"] for name, settings in PROFILES.items()}}

    def start(self, text: str, marketplace: str = "US", profile: str = "balanced") -> dict[str, Any]:
        if profile not in PROFILES:
            raise ValueError("访问档位无效")
        parsed = parse_inputs(text, marketplace)
        if not parsed["items"]:
            raise ValueError("没有有效的 Amazon ASIN 或商品链接")
        health = runtime_status()
        if not health["available"]:
            raise ValueError(health["message"])
        job_id = uuid.uuid4().hex
        now = time.time()
        job = {"id": job_id, "status": "queued", "profile": profile, "created_at": now, "queued_at": now,
               "updated_at": now, "revision": 0, "attempt_generation": 0,
               "items": [{**item, "status": "pending", "title": "", "image_count": 0, "error": ""} for item in parsed["items"]],
               "invalid": parsed["invalid"], "duplicates_removed": parsed["duplicates_removed"],
               "restricted_results": 0, "output_file": None, "proxy_count": health["proxy_count"]}
        with self.lock:
            self._save(job)
            self._dispatch_queued_locked()
        return self.get(job_id)

    def _queued_job_ids_locked(self) -> list[str]:
        queued = []
        for path in self.root.glob("*/job.json"):
            try:
                # Always read through the per-job lock. A raw read handle can keep
                # job.json open on Windows while a concurrent checkpoint calls os.replace.
                job_id = path.parent.name
                job = self._read(job_id)
                job_id = str(job.get("id", ""))
                if job.get("status") == "queued" and job_id not in self.active_job_ids:
                    queued.append((float(job.get("queued_at", job.get("created_at", 0))), job_id))
            except Exception:
                continue
        return [job_id for _, job_id in sorted(queued)]

    def _dispatch_queued_locked(self) -> None:
        if self.shutting_down:
            return
        for job_id in self._queued_job_ids_locked():
            if len(self.active_job_ids) >= MAX_CONCURRENT_JOBS:
                break
            self._launch(job_id)

    def _launch(self, job_id: str) -> bool:
        with self.lock:
            if self.shutting_down or job_id in self.active_job_ids or len(self.active_job_ids) >= MAX_CONCURRENT_JOBS:
                return False
            job = self._read(job_id)
            if job.get("status") != "queued":
                return False
            cancel = threading.Event()
            thread = threading.Thread(target=self._run, args=(job_id, cancel), name=f"amazon-scraper-{job_id[:8]}", daemon=True)
            self.cancel_events[job_id] = cancel
            self.threads[job_id] = thread
            self.active_job_ids.add(job_id)
            try:
                thread.start()
            except Exception as exc:
                self.active_job_ids.discard(job_id)
                self.cancel_events.pop(job_id, None)
                job["status"] = "failed"; job["error"] = _scrub_error(str(exc)); job["updated_at"] = time.time(); self._save(job)
                return False
            return True

    def list_jobs(self) -> list[dict[str, Any]]:
        jobs = []
        queue_positions, active_job_count = self._queue_snapshot()
        for path in sorted(self.root.glob("*/job.json"), key=lambda p: p.stat().st_mtime, reverse=True)[:30]:
            try:
                job = self._read(path.parent.name)
                jobs.append(self._public(job, queue_positions, active_job_count))
            except Exception:
                continue
        return jobs

    def _queue_snapshot(self) -> tuple[dict[str, int], int]:
        with self.lock:
            queue_positions = {job_id: index for index, job_id in enumerate(self._queued_job_ids_locked(), 1)}
            return queue_positions, len(self.active_job_ids)

    def _public(self, job: dict[str, Any], queue_positions: dict[str, int] | None = None,
                active_job_count: int | None = None) -> dict[str, Any]:
        if queue_positions is None or active_job_count is None:
            queue_positions, active_job_count = self._queue_snapshot()
        items = [{key: item.get(key) for key in ("marketplace", "asin", "status", "title", "image_count", "error")} for item in job["items"]]
        summary = {key: job.get(key) for key in ("id", "status", "profile", "created_at", "updated_at", "invalid", "duplicates_removed", "restricted_results", "output_file", "proxy_count", "error")}
        summary["revision"] = int(job.get("revision", 0))
        return summary | self._access_limits() | {
            "items": items, "total": len(items), "completed": sum(item["status"] == "ready" for item in job["items"]),
            "pending": sum(item["status"] == "pending" for item in job["items"]),
            "failed": sum(item["status"] in ("failed", "needs_attention") for item in job["items"]),
            "retryable": sum(item["status"] in ("failed", "needs_attention", "cancelled", "pending") for item in job["items"]),
            "can_download": bool(job.get("output_file")) and (self.root / job["id"] / "result.xlsx").is_file(),
            "queue_position": queue_positions.get(job["id"]), "active_job_count": active_job_count,
            "max_concurrent_jobs": MAX_CONCURRENT_JOBS}

    def get(self, job_id: str) -> dict[str, Any]:
        queue_positions, active_job_count = self._queue_snapshot()
        return self._public(self._read(job_id), queue_positions, active_job_count)

    def cancel(self, job_id: str) -> dict[str, Any]:
        deadline = time.monotonic() + 4.5
        with self.lock:
            with self._job_lock(job_id):
                job = self._read(job_id)
                if job["status"] not in ("queued", "running", "cancelling"):
                    return self._public(job)
                generation = int(job.get("attempt_generation", 0))
                was_active = job_id in self.active_job_ids
                worker_thread = self.threads.get(job_id) if was_active else None
                event = self.cancel_events.get(job_id)
                if event:
                    event.set()
                if job["status"] != "cancelling":
                    for item in job["items"]:
                        if item["status"] in ("pending", "scraping", "cancelled"):
                            item.update(status="pending", error="已中断，可继续")
                    job["attempt_generation"] = generation + 1
                    job["status"] = "cancelling"; job["updated_at"] = time.time()
                    self._save(job, expected_generation=generation)
        with self._fetch_condition:
            self._fetch_condition.notify_all()
        # Release the worker tree and wait for its parent process to exit before reporting
        # the terminal cancellation state. Fencing is already active while this happens.
        stopped = self._terminate_children(job_id, deadline=deadline)
        # Join the exact worker even if the first child snapshot was empty: it may
        # have been between Popen and registry insertion when cancellation arrived.
        if worker_thread and worker_thread.is_alive():
            worker_thread.join(timeout=max(0.0, deadline - time.monotonic()))
        stopped = self._terminate_children(job_id, deadline=deadline) and stopped
        with self.lock:
            children_alive = any(child.poll() is None for child, _ in self._child_processes.get(job_id, ()))
        if not stopped or children_alive or (worker_thread and worker_thread.is_alive()):
            return self.get(job_id)
        with self.lock:
            with self._job_lock(job_id):
                job = self._read(job_id)
                if job.get("status") == "cancelling":
                    generation = int(job.get("attempt_generation", 0))
                    job["status"] = "cancelled"; job["updated_at"] = time.time()
                    self._save(job, expected_generation=generation, allowed_statuses=("cancelling",))
            self._dispatch_queued_locked()
        return self.get(job_id)

    def retry(self, job_id: str) -> dict[str, Any]:
        with self.lock:
            with self._job_lock(job_id):
                job = self._read(job_id)
                if job["status"] in ("queued", "running", "cancelling"):
                    raise ValueError("采集任务仍在运行")
                for item in job["items"]:
                    if item["status"] in ("failed", "needs_attention", "cancelled", "pending"):
                        item.update(status="pending", error="")
                if not any(item["status"] == "pending" for item in job["items"]):
                    raise ValueError("没有可重试的商品")
                now = time.time()
                generation = int(job.get("attempt_generation", 0))
                job["attempt_generation"] = generation + 1
                job["status"] = "queued"; job["queued_at"] = now; job["restricted_results"] = 0
                job["output_file"] = None; job["updated_at"] = now
                self._save(job, expected_generation=generation); self._dispatch_queued_locked()
        return self.get(job_id)

    def output_path(self, job_id: str) -> Path:
        job = self._read(job_id)
        path = self.root / job_id / "result.xlsx"
        if not job.get("output_file") or not path.is_file():
            raise ValueError("采集结果尚未生成")
        return path

    def recover(self) -> None:
        for path in self.root.glob("*/job.json"):
            try:
                job = json.loads(path.read_text(encoding="utf-8"))
                if job.get("status") in ("queued", "running", "cancelling"):
                    for item in job["items"]:
                        if item["status"] == "scraping":
                            item.update(status="pending", error="")
                    job["status"] = "interrupted"; job["updated_at"] = time.time()
                    self._save(job)
            except Exception:
                continue

    def stop_all(self) -> None:
        deadline = time.monotonic() + 10.0
        with self.lock:
            self.shutting_down = True
            events = list(self.cancel_events.values())
            active_ids = list(self.active_job_ids)
            threads = list(self.threads.values())
            queued_ids = self._queued_job_ids_locked()
            job_ids = list(dict.fromkeys(active_ids + queued_ids))
            for job_id in job_ids:
                try:
                    with self._job_lock(job_id):
                        job = self._read(job_id)
                        if job.get("status") not in ("queued", "running", "cancelling"):
                            continue
                        generation = int(job.get("attempt_generation", 0))
                        for item in job["items"]:
                            if item["status"] in ("pending", "scraping", "cancelled"):
                                item.update(status="pending", error="工作台关闭，已中断，可继续")
                        job["attempt_generation"] = generation + 1
                        job["status"] = "interrupted"
                        job["updated_at"] = time.time()
                        self._save(job, expected_generation=generation)
                except Exception:
                    continue
        for event in events:
            event.set()
        with self._fetch_condition:
            self._fetch_condition.notify_all()
        for job_id in job_ids:
            self._terminate_children(job_id, deadline=deadline)
        for thread in threads:
            if thread.is_alive():
                thread.join(timeout=max(0, deadline - time.monotonic()))

    def _terminate_children(self, job_id: str, deadline: float | None = None) -> bool:
        with self.lock:
            children = list(self._child_processes.get(job_id, ()))
        deadline = deadline if deadline is not None else time.monotonic() + 4.0
        # Terminate owned direct children and close each Job Object first; closing the
        # object ends any descendants as well. Then reap each exact Popen handle.
        for child, process_job in children:
            try:
                if child.poll() is None:
                    child.terminate()
                with self.lock:
                    process_job.release()
            except Exception:
                pass
        stopped = True
        for child, _process_job in children:
            try:
                remaining = max(0.0, deadline - time.monotonic())
                if child.poll() is None and remaining:
                    child.wait(timeout=remaining)
                if child.poll() is None:
                    child.kill()
                    remaining = max(0.0, deadline - time.monotonic())
                    if remaining:
                        child.wait(timeout=remaining)
                stopped = stopped and child.poll() is not None
            except subprocess.TimeoutExpired:
                try:
                    child.kill()
                    remaining = deadline - time.monotonic()
                    if remaining > 0:
                        child.wait(timeout=remaining)
                    else:
                        stopped = False
                except Exception:
                    stopped = False
            except Exception:
                stopped = False
        return stopped

    def _choose_proxy(self, pool: list[str]) -> str | None:
        if not pool:
            return None
        with self.lock:
            proxy = pool[self.proxy_rotation % len(pool)]
            self.proxy_rotation += 1
        return proxy

    def _run(self, job_id: str, cancel: threading.Event) -> None:
        generation: int | None = None
        try:
            job = self._read(job_id)
            generation = int(job.get("attempt_generation", 0))
            job["status"] = "running"; job["updated_at"] = time.time()
            if not self._save(job, expected_generation=generation, allowed_statuses=("queued",)):
                return
            settings = PROFILES[job["profile"]]
            pool = _proxy_pool()
            pending = [index for index, item in enumerate(job["items"]) if item["status"] == "pending"]
            consecutive_restrictions = int(job.get("restricted_results", 0))
            downshifted = False
            with ThreadPoolExecutor(max_workers=settings["workers"], thread_name_prefix="amazon-product") as executor:
                while pending and not cancel.is_set():
                    workers = 1 if downshifted else settings["workers"]
                    batch_indexes, pending = pending[:workers], pending[workers:]
                    futures = {}
                    for index in batch_indexes:
                        job = self._read(job_id)
                        item = job["items"][index]
                        item["status"] = "scraping"; job["updated_at"] = time.time()
                        if not self._save(job, expected_generation=generation, allowed_statuses=("running",)):
                            return
                        futures[executor.submit(self._delayed_fetch, job_id, item, self._choose_proxy(pool), cancel,
                                                 settings["delay_min"], settings["delay_max"])] = index
                    results = []
                    restriction_in_wave = False
                    for future in as_completed(futures):
                        index = futures[future]
                        try:
                            result = future.result()
                        except Exception as exc:
                            result = {"status": "failed", "error": _scrub_error(str(exc), pool)}
                        results.append((index, result))
                        if result.get("status") == "restricted" or result.get("restricted"):
                            restriction_in_wave = True
                            downshifted = True
                            with self._fetch_condition:
                                self._restricted_jobs.add(job_id)
                                self._fetch_condition.notify_all()
                            self._terminate_children(job_id)
                    for index, result in results:
                        if not self._is_current(job_id, generation, "running"):
                            return
                        state = result.get("status", "failed")
                        if state == "ready":
                            if not self._commit_success(job_id, generation, index, result):
                                return
                            consecutive_restrictions = 0
                            continue
                        job = self._read(job_id)
                        item = job["items"][index]
                        if state == "restricted" or result.get("restricted"):
                            consecutive_restrictions += 1
                            downshifted = True
                            item.update(status="needs_attention", error=result.get("error", "访问受限"))
                        else:
                            if state == "cancelled" and restriction_in_wave and not cancel.is_set():
                                item.update(status="pending", error="访问受限后暂停本轮派发")
                            else:
                                item.update(status=state if state in ("failed", "needs_attention", "cancelled") else "failed", error=result.get("error", "采集失败"))
                        job["restricted_results"] = consecutive_restrictions; job["updated_at"] = time.time()
                        if not self._save(job, expected_generation=generation, allowed_statuses=("running",)):
                            return
                    if restriction_in_wave:
                        # Cool down even if another in-flight request succeeded and reset
                        # the consecutive counter. Restriction also permanently downshifts
                        # this attempt to one page worker.
                        if consecutive_restrictions >= MAX_RESTRICTED_RESULTS:
                            break
                        if cancel.is_set():
                            break
                        self._cooldown(cancel)
                        if cancel.is_set():
                            break
                        with self._fetch_condition:
                            self._restricted_jobs.discard(job_id)
                            self._fetch_condition.notify_all()
                    # Rebuild from durable item states. Work cancelled by the restriction
                    # gate returns to pending instead of disappearing from the local queue.
                    job = self._read(job_id)
                    pending = [index for index, item in enumerate(job["items"])
                               if item.get("status") == "pending"]
            job = self._read(job_id)
            if int(job.get("attempt_generation", 0)) != generation or job.get("status") != "running":
                return
            if cancel.is_set():
                for item in job["items"]:
                    if item["status"] in ("pending", "scraping"):
                        item.update(status="pending", error="已中断，可继续")
                job["status"] = "cancelled"
            elif consecutive_restrictions >= MAX_RESTRICTED_RESULTS:
                for item in job["items"]:
                    if item["status"] == "scraping": item.update(status="pending", error="")
                job["status"] = "paused"
            else:
                if cancel.is_set():
                    return
                self._build_workbook(job_id, job, cancel=cancel, generation=generation)
                if cancel.is_set() or not self._is_current(job_id, generation, "running"):
                    self._cleanup_scratch(job_id, generation)
                    return
                job["output_file"] = "result.xlsx"
                job["status"] = "completed_with_errors" if any(item["status"] in ("failed", "needs_attention") for item in job["items"]) else "completed"
            job["updated_at"] = time.time(); self._save(job, expected_generation=generation, allowed_statuses=("running",))
        except Exception as exc:
            try:
                job = self._read(job_id)
                if generation is not None and int(job.get("attempt_generation", 0)) == generation and job.get("status") == "running":
                    job["status"] = "failed"; job["error"] = _scrub_error(str(exc)); job["updated_at"] = time.time()
                    self._save(job, expected_generation=generation, allowed_statuses=("running",))
            except Exception:
                pass
        finally:
            with self._fetch_condition:
                self._restricted_jobs.discard(job_id)
                self._last_start_at.pop(job_id, None)
                self._fetch_condition.notify_all()
            try:
                with self._job_lock(job_id):
                    latest = self._read(job_id)
                    if latest.get("status") == "cancelling":
                        latest["status"] = "cancelled"
                        latest["updated_at"] = time.time()
                        self._save(latest, expected_generation=int(latest.get("attempt_generation", 0)),
                                   allowed_statuses=("cancelling",))
            except Exception:
                pass
            with self.lock:
                self.active_job_ids.discard(job_id)
                self.cancel_events.pop(job_id, None)
                self._dispatch_queued_locked()

    def _cooldown(self, cancel: threading.Event) -> None:
        cancel.wait(RESTRICTED_COOLDOWN_SECONDS)

    def _is_current(self, job_id: str, generation: int, status: str) -> bool:
        try:
            job = self._read(job_id)
            return int(job.get("attempt_generation", 0)) == generation and job.get("status") == status
        except Exception:
            return False

    def _cleanup_scratch(self, job_id: str, generation: int | None = None) -> None:
        folder = self.root / job_id / ".scratch"
        if not folder.is_dir():
            return
        if generation is None:
            targets = list(folder.iterdir())
        else:
            targets = list(folder.glob(f"{generation}-*"))
        for target in targets:
            if target.is_dir() and target.parent == folder:
                shutil.rmtree(target, ignore_errors=True)

    def _commit_success(self, job_id: str, generation: int, item_index: int,
                        result: dict[str, Any]) -> bool:
        with self._job_lock(job_id):
            job = self._read(job_id)
            if int(job.get("attempt_generation", 0)) != generation or job.get("status") != "running":
                return False
            item = job["items"][item_index]
            images = result.get("images", [])
            self._save_product_images(job_id, item["asin"], images)
            item.update(status="ready", title=result.get("title", ""), image_count=len(images), error="")
            job["restricted_results"] = 0
            job["updated_at"] = time.time()
            return self._save(job, expected_generation=generation, allowed_statuses=("running",))

    def _save_product_images(self, job_id: str, asin: str, images: list[tuple[bytes, str]]) -> None:
        folder = self.root / job_id / "images"; folder.mkdir(parents=True, exist_ok=True)
        for index, (data, mime) in enumerate(images[:MAX_IMAGES], 1):
            extension = ".jpg" if mime == "image/jpeg" else ".png"
            target = folder / f"{asin}_{index:02d}{extension}"
            fd, temp_name = tempfile.mkstemp(prefix=f".{asin}.{uuid.uuid4().hex}.", suffix=".tmp", dir=folder)
            try:
                with os.fdopen(fd, "wb") as stream:
                    stream.write(data)
                    stream.flush()
                    os.fsync(stream.fileno())
                os.replace(temp_name, target)
            finally:
                try: os.unlink(temp_name)
                except FileNotFoundError: pass

    def _delayed_fetch(self, job_id: str, item: dict[str, str], proxy: str | None, cancel: threading.Event,
                       delay_min: float, delay_max: float) -> dict[str, Any]:
        if not self._acquire_fetch_slot(job_id, cancel):
            return {"status": "cancelled", "error": "访问受限后暂停本轮派发"}
        try:
            if self._use_worker_process:
                return self._fetch_in_worker(job_id, item, proxy, cancel, delay_min, delay_max)
            admitted, _ = self._dispatch_with_interval(job_id, delay_min, delay_max, cancel, lambda: True)
            if not admitted:
                return {"status": "cancelled", "error": "访问受限或用户取消"}
            return self.engine(item, proxy, cancel)
        finally:
            self._release_fetch_slot(job_id)

    def _dispatch_with_interval(self, job_id: str, delay_min: float, delay_max: float,
                                cancel: threading.Event, dispatch: Callable[[], Any]) -> tuple[bool, Any]:
        """Wait without holding the gate, then atomically dispatch and timestamp a start."""
        interval = random.uniform(delay_min, delay_max)
        with self._fetch_condition:
            while True:
                if cancel.is_set() or job_id in self._restricted_jobs or self.shutting_down:
                    return False, None
                now = time.monotonic()
                last_start = self._last_start_at.get(job_id)
                allowed_at = now if last_start is None else last_start + interval
                if now >= allowed_at:
                    result = dispatch()
                    if result is False:
                        return False, result
                    self._last_start_at[job_id] = time.monotonic()
                    return True, result
                self._fetch_condition.wait(allowed_at - now)

    def _acquire_fetch_slot(self, job_id: str, cancel: threading.Event) -> bool:
        with self._fetch_condition:
            self._fetch_waiters[job_id] = self._fetch_waiters.get(job_id, 0) + 1
            if job_id not in self._fetch_order:
                self._fetch_order.append(job_id)
            while True:
                if cancel.is_set() or job_id in self._restricted_jobs or self.shutting_down:
                    self._remove_fetch_waiter(job_id)
                    return False
                if self._fetch_active < MAX_ACTIVE_PAGE_FETCHES and self._fetch_order and self._fetch_order[0] == job_id:
                    self._remove_fetch_waiter(job_id, rotate=True)
                    self._fetch_active += 1
                    return True
                self._fetch_condition.wait(0.1)

    def _remove_fetch_waiter(self, job_id: str, rotate: bool = False) -> None:
        count = self._fetch_waiters.get(job_id, 0)
        if count <= 1:
            self._fetch_waiters.pop(job_id, None)
            try:
                self._fetch_order.remove(job_id)
            except ValueError:
                pass
        else:
            self._fetch_waiters[job_id] = count - 1
            if rotate and job_id in self._fetch_order:
                self._fetch_order.remove(job_id)
                self._fetch_order.append(job_id)
        self._fetch_condition.notify_all()

    def _release_fetch_slot(self, job_id: str) -> None:
        with self._fetch_condition:
            self._fetch_active = max(0, self._fetch_active - 1)
            self._fetch_condition.notify_all()

    def _fetch_in_worker(self, job_id: str, item: dict[str, str], proxy: str | None,
                         cancel: threading.Event, delay_min: float = 0.0,
                         delay_max: float = 0.0) -> dict[str, Any]:
        generation = int(self._read(job_id).get("attempt_generation", 0))
        scratch = self.root / job_id / ".scratch" / f"{generation}-{item['asin']}-{uuid.uuid4().hex}"
        scratch.mkdir(parents=True, exist_ok=False)
        from .agent.runtime import ProcessJob
        process_job = ProcessJob()
        process_job.create()
        command = [sys.executable, "-m", "backend.amazon_scraper_worker"]
        try:
            child = subprocess.Popen(command, cwd=str(Path(__file__).resolve().parent.parent),
                                     stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                                     text=True, encoding="utf-8", errors="replace", bufsize=1,
                                     creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except BaseException:
            process_job.release()
            shutil.rmtree(scratch, ignore_errors=True)
            raise
        child_entry = (child, process_job)
        with self.lock:
            self._child_processes.setdefault(job_id, set()).add(child_entry)
        try:
            assigned = process_job.assign(child.pid)
            if os.name == "nt" and not assigned:
                child.kill(); child.wait(timeout=2)
                raise RuntimeError("Amazon 采集进程无法加入受管 Job Object，已阻止网络访问")
            message = {"action": "fetch", "item": item, "proxy": proxy,
                       "scratch": str(scratch), "generation": generation}
            # This final status check, interval admission, and start signal are atomic
            # with the restriction gate. The child is already attached and waits on stdin.
            def send_start_signal() -> bool:
                with self.lock:
                    latest = self._read(job_id)
                    if (cancel.is_set() or job_id in self._restricted_jobs or self.shutting_down or
                            int(latest.get("attempt_generation", 0)) != generation or latest.get("status") != "running"):
                        return False
                    child.stdin.write(json.dumps(message, ensure_ascii=False) + "\n")
                    child.stdin.flush(); child.stdin.close()
                    return True
            admitted, _ = self._dispatch_with_interval(job_id, delay_min, delay_max, cancel, send_start_signal)
            if not admitted:
                child.kill(); child.wait(timeout=2)
                return {"status": "cancelled", "error": "用户取消或访问受限"}
            while child.poll() is None:
                if cancel.wait(0.05):
                    child.terminate()
                    try: child.wait(timeout=1)
                    except subprocess.TimeoutExpired: child.kill(); child.wait(timeout=1)
                    return {"status": "cancelled", "error": "用户取消"}
            stdout = child.stdout.read()
            if child.returncode != 0:
                # Keep stderr private: it may contain URLs, proxy details, or browser internals.
                return {"status": "failed", "error": "Amazon 采集工作进程异常退出"}
            envelope = json.loads(stdout)
            result = dict(envelope.get("result", {}))
            if result.get("status") == "ready":
                images = []
                for name, mime in envelope.get("images", []):
                    if not re.fullmatch(r"image-\d{2}\.bin", str(name)):
                        continue
                    image_path = scratch / name
                    if image_path.parent != scratch or not image_path.is_file():
                        continue
                    images.append((image_path.read_bytes(), mime))
                result["images"] = images
            return result
        finally:
            with self.lock:
                self._child_processes.get(job_id, set()).discard(child_entry)
                process_job.release()
            shutil.rmtree(scratch, ignore_errors=True)

    def _build_workbook(self, job_id: str, job: dict[str, Any], cancel: threading.Event | None = None,
                        generation: int | None = None) -> None:
        if self._use_worker_process:
            generation = int(job.get("attempt_generation", 0)) if generation is None else generation
            scratch = self.root / job_id / ".scratch" / f"{generation}-export-{uuid.uuid4().hex}"
            scratch.mkdir(parents=True, exist_ok=False)
            from .agent.runtime import ProcessJob
            process_job = ProcessJob(); process_job.create()
            try:
                child = subprocess.Popen([sys.executable, "-m", "backend.amazon_scraper_worker"],
                                         cwd=str(Path(__file__).resolve().parent.parent),
                                         stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                                         text=True, encoding="utf-8", errors="replace", bufsize=1,
                                         creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            except BaseException:
                process_job.release()
                shutil.rmtree(scratch, ignore_errors=True)
                raise
            child_entry = (child, process_job)
            with self.lock:
                self._child_processes.setdefault(job_id, set()).add(child_entry)
            try:
                if os.name == "nt" and not process_job.assign(child.pid):
                    child.kill(); child.wait(timeout=2)
                    raise RuntimeError("Excel 生成进程无法加入受管 Job Object")
                target = scratch / "result.xlsx"
                with self.lock:
                    latest = self._read(job_id)
                    if (cancel is not None and cancel.is_set()) or int(latest.get("attempt_generation", 0)) != generation or latest.get("status") != "running":
                        child.kill(); child.wait(timeout=2)
                        return
                    child.stdin.write(json.dumps({"action": "export", "job": job, "root": str(self.root),
                                                  "target": str(target)}) + "\n")
                    child.stdin.flush(); child.stdin.close()
                while child.poll() is None:
                    if cancel is not None and cancel.wait(0.05):
                        child.terminate()
                        try: child.wait(timeout=1)
                        except subprocess.TimeoutExpired: child.kill(); child.wait(timeout=1)
                        return
                if child.returncode != 0:
                    raise RuntimeError("Excel 工作簿生成进程异常退出")
                if not target.is_file():
                    raise RuntimeError("Excel 工作簿生成进程未写入结果")
                with self._job_lock(job_id):
                    if generation is not None and not self._is_current(job_id, generation, "running"):
                        return
                    os.replace(target, self.root / job_id / "result.xlsx")
                return
            finally:
                with self.lock:
                    self._child_processes.get(job_id, set()).discard(child_entry)
                    process_job.release()
                shutil.rmtree(scratch, ignore_errors=True)
        ready = [item for item in job["items"] if item["status"] == "ready"]
        max_images = max((item["image_count"] for item in ready), default=0)
        wb = Workbook(); ws = wb.active; ws.title = "Amazon Listing 图片"
        ws.append(["ASIN", "商品标题"] + [f"图片 {index}" for index in range(1, max_images + 1)])
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 48
        for column in range(3, 3 + max_images):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(column)].width = 22
        for row_number, item in enumerate(ready, 2):
            ws.cell(row_number, 1, item["asin"]); ws.cell(row_number, 2, item.get("title", "")); ws.row_dimensions[row_number].height = 112
            folder = self.root / job_id / "images"
            paths = sorted(folder.glob(f"{item['asin']}_*"))
            for image_index, path in enumerate(paths[:MAX_IMAGES], 3):
                image = ExcelImage(path); image.width = 144; image.height = 144
                from openpyxl.utils import get_column_letter
                ws.add_image(image, f"{get_column_letter(image_index)}{row_number}")
        target = self.root / job_id / "result.xlsx"
        wb.save(target); wb.close()


service = None if os.environ.get("AMAZON_SCRAPER_WORKER") == "1" else AmazonScraperService()


def recover() -> None:
    if service is not None:
        service.recover()


def stop_all() -> None:
    if service is not None:
        service.stop_all()
