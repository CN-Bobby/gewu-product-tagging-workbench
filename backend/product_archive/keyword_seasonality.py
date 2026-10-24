"""Read the completed keyword report's monthly search series for one archive."""

from __future__ import annotations

import math
import re
from pathlib import Path
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException


_MONTH = re.compile(r"\d{4}-(0[1-9]|1[0-2])\Z")
_WIDE_SHEET = "03_24月搜索量_宽表"
MAX_KEYWORDS = 10
MAX_MONTHS = 36


def read_report(path: Path) -> dict:
    """Return report-order search series; blank months remain null rather than zero."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if _WIDE_SHEET not in workbook:
            return {"status": "trend_sheet_missing", "months": [], "keywords": []}
        sheet = workbook[_WIDE_SHEET]
        header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=3 + MAX_MONTHS,
                                     values_only=True), ())
        columns = []
        seen_months = set()
        for index, value in enumerate(header):
            month = str(value or "").strip()
            if index < 3 or not _MONTH.fullmatch(month) or month in seen_months:
                continue
            seen_months.add(month)
            columns.append((index, month))
        columns.sort(key=lambda item: item[1])
        months = [month for _, month in columns]
        if len(months) < 3:
            return {"status": "no_monthly_data", "months": months, "keywords": []}
        keywords = []
        seen_keywords = set()
        for row in sheet.iter_rows(min_row=2, max_row=1 + MAX_KEYWORDS * 3,
                                   max_col=3 + MAX_MONTHS, values_only=True):
            word = str(row[0] or "").strip()[:160]
            if not word or word.casefold() in seen_keywords:
                continue
            seen_keywords.add(word.casefold())
            values = []
            for index, _month in columns:
                value = row[index] if index < len(row) else None
                valid = (isinstance(value, (int, float)) and not isinstance(value, bool)
                         and math.isfinite(value) and value >= 0)
                values.append(int(value) if valid else None)
            if not any(value is not None for value in values):
                continue
            keywords.append({"keyword": word, "translation": str(row[1] or "").strip()[:160],
                             "type": str(row[2] or "").strip()[:40], "values": values})
            if len(keywords) == MAX_KEYWORDS:
                break
        return {"status": "available" if keywords else "no_monthly_data",
                "months": months, "keywords": keywords}
    finally:
        workbook.close()


def for_archive(archive_id: str, source_id: str = "") -> dict:
    """Only expose a finished keyword artifact verified against this archive."""
    from . import overview_drafts, service
    from .. import project_files

    archive = service._read(archive_id)
    available = [item for item in overview_drafts.list_sources(archive_id)
                 if item.get("type") == "keywords" and item.get("available")]
    if source_id:
        selected = next((item for item in available if item.get("id") == source_id), None)
        if selected is None:
            raise ValueError("关键词报告不属于当前档案或任务尚未完成")
    else:
        selected = max(available, key=lambda item: str(item.get("date") or "")) if available else None
    if selected is None:
        return {"status": "no_keyword_report", "source": None, "months": [], "keywords": [],
                "metric": "monthly_searches"}
    workspace_id = str(archive.get("workspace_id") or "")
    artifact_id = selected["id"].removeprefix("artifact:")
    path = project_files.artifact_path(workspace_id, artifact_id)
    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("关键词报告不是可读取的工作簿")
    try:
        result = read_report(path)
    except (OSError, ValueError, KeyError, TypeError, BadZipFile, InvalidFileException) as exc:
        result = {"status": "read_failed", "reason": type(exc).__name__,
                  "months": [], "keywords": []}
    return {**result, "source": {"id": selected["id"], "name": selected.get("name"),
                                 "date": selected.get("date")},
            "metric": "monthly_searches", "marketplace": archive.get("marketplace"),
            "note": "关键词任务报告的月搜索量；不同词可能覆盖不同商品形态。搜索需求不等于本档案商品销量。"}


def from_selected_materials(archive: dict, materials: list[dict]) -> dict | None:
    """Use the selected report with the broadest monthly keyword coverage."""
    from .. import project_files

    best = None
    for material in materials:
        source_id = str(material.get("source_id") or "")
        if material.get("type") != "keywords" or not source_id.startswith("artifact:"):
            continue
        try:
            path = project_files.artifact_path(str(archive.get("workspace_id") or ""),
                                               source_id.removeprefix("artifact:"))
            result = read_report(path)
        except (OSError, ValueError, KeyError, TypeError, BadZipFile, InvalidFileException):
            continue
        if result["status"] != "available":
            continue
        candidate = {**result, "source_id": source_id}
        if best is None or len(candidate["keywords"]) > len(best["keywords"]):
            best = candidate
    return best
