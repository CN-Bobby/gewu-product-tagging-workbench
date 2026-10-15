"""Read-only views of review and VOC tasks linked to a product archive."""

from __future__ import annotations

from datetime import date, datetime, time
from zipfile import BadZipFile
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..amazon_reviews import service as reviews_service
from .. import project_files
from . import service as archive_service


def _archive_review_jobs(archive_id: str, *, include_sheets: bool = True) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    archive = archive_service._read(archive_id)
    workspace_id = str(archive.get("workspace_id") or "")
    if not workspace_id:
        return archive, []
    jobs = reviews_service.list_jobs(
        workspace_id=workspace_id,
        source_archive_id=archive_id,
        source_dataset_id=str(archive.get("dataset_id") or ""),
        limit=1000,
    )
    visible = []
    fields = ("id", "status", "mode", "site", "created_at", "updated_at", "review_count",
              "total", "completed", "failed", "can_download", "voc_handoff")
    for job in jobs:
        item = {key: job.get(key) for key in fields}
        item["source_kind"] = "amazon_reviews_job"
        sheets = _review_workbook_sheets(archive, item) if include_sheets and job.get("can_download") else []
        item["sheets"] = sheets
        if include_sheets and job.get("can_download"):
            item["workbook_readable"] = bool(sheets)
        visible.append(item)
    for artifact in project_files.list_artifacts(workspace_id, "reviews"):
        if not artifact.get("downloadable"):
            continue
        item = {
            "id": artifact["id"], "status": "completed", "mode": "project_file",
            "site": "项目评论文件", "created_at": artifact.get("created_at"),
            "updated_at": artifact.get("created_at"), "review_count": None,
            "total": None, "completed": None, "failed": None,
            "can_download": True, "voc_handoff": None,
            "source_kind": "project_artifact", "source_type": artifact.get("source_type"),
            "name": artifact.get("name"), "artifact_id": artifact["id"],
            "workspace_id": workspace_id,
        }
        sheets = _review_workbook_sheets(archive, item) if include_sheets else []
        item["sheets"] = sheets
        if include_sheets:
            item["workbook_readable"] = bool(sheets)
        visible.append(item)
    return archive, visible


def _review_workbook_path(archive: dict[str, Any], item: dict[str, Any]):
    if item.get("source_kind") == "project_artifact":
        workspace_id = str(item.get("workspace_id") or archive.get("workspace_id") or "")
        return project_files.artifact_path(workspace_id, str(item.get("artifact_id") or item.get("id") or ""))
    return reviews_service.output_path(str(item.get("id") or ""))


def _review_workbook_sheets(archive: dict[str, Any], item: dict[str, Any]) -> list[str]:
    try:
        path = _review_workbook_path(archive, item)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException):
        return []


def list_review_jobs(archive_id: str) -> list[dict[str, Any]]:
    return _archive_review_jobs(archive_id)[1]


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def review_workbook_page(archive_id: str, job_id: str, *, sheet_name: str = "",
                         offset: int = 0, limit: int = 50, query: str = "") -> dict[str, Any]:
    archive, jobs = _archive_review_jobs(archive_id)
    job = next((item for item in jobs if item.get("id") == job_id), None)
    if not job:
        raise ValueError("评论任务不存在或不属于当前选品档案")
    if not job.get("can_download"):
        raise ValueError("该评论任务尚未生成可查看的 Excel")
    if offset < 0 or offset > 10_000_000:
        raise ValueError("评论表格页码无效")
    if limit < 1 or limit > 100:
        raise ValueError("每页显示条数需在 1–100 之间")
    if len(query) > 300:
        raise ValueError("评论搜索内容不能超过 300 个字符")

    try:
        path = _review_workbook_path(archive, job)
    except (OSError, ValueError, KeyError) as exc:
        raise ValueError("评论 Excel 文件无法读取") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException) as exc:
        raise ValueError("评论 Excel 文件无法读取") from exc
    try:
        available_sheets = list(workbook.sheetnames)
        if not available_sheets:
            raise ValueError("评论 Excel 中没有工作表")
        selected_sheet = sheet_name or available_sheets[0]
        if selected_sheet not in available_sheets:
            raise ValueError("所选工作表不存在")
        worksheet = workbook[selected_sheet]
        iterator = worksheet.iter_rows(values_only=True)
        raw_headers = next(iterator, ())
        headers = [str(value).strip() if value is not None else f"列 {index + 1}"
                   for index, value in enumerate(raw_headers)]
        while headers and not headers[-1]:
            headers.pop()
        if not headers:
            return {"job": job, "sheet": selected_sheet, "sheets": available_sheets,
                    "headers": [], "rows": [], "offset": offset, "limit": limit,
                    "total": 0, "query": query}

        needle = query.strip().casefold()
        rows: list[dict[str, Any]] = []
        total = 0
        for excel_row, raw_row in enumerate(iterator, start=2):
            values = [_cell_value(value) for value in raw_row[:len(headers)]]
            if len(values) < len(headers):
                values.extend([None] * (len(headers) - len(values)))
            if not any(value not in (None, "") for value in values):
                continue
            if needle and not any(needle in str(value).casefold() for value in values if value is not None):
                continue
            if offset <= total < offset + limit:
                rows.append({"excel_row": excel_row, "values": values})
            total += 1
        return {"job": job, "sheet": selected_sheet, "sheets": available_sheets,
                "headers": headers, "rows": rows, "offset": offset, "limit": limit,
                "total": total, "query": query}
    finally:
        workbook.close()


def list_voc_runs(archive_id: str) -> list[dict[str, Any]]:
    archive, review_jobs = _archive_review_jobs(archive_id, include_sheets=False)
    workspace_id = str(archive.get("workspace_id") or "")
    if not workspace_id:
        return []
    review_job_ids = {str(job.get("id") or "") for job in review_jobs}
    linked_run_ids = {str((job.get("voc_handoff") or {}).get("run_id") or "")
                      for job in review_jobs}
    linked_run_ids.discard("")

    from ..amazon_voc.service import service as voc_service

    project_voc_runs = set()
    for artifact in project_files.list_artifacts(workspace_id, "voc"):
        if artifact.get("source_type") == "amazon_voc_output":
            metadata = artifact.get("metadata") or {}
            run_id = str(metadata.get("run_id") or "")
            if not run_id:
                run_id = str(artifact.get("source_id") or "").split(":rev", 1)[0]
            if run_id:
                project_voc_runs.add(run_id)

    result = []
    for run in voc_service.list_runs():
        run_workspace_id = str(run.get("workspace_id") or "")
        if not run_workspace_id and run.get("project_id"):
            voc_project = voc_service.project(str(run["project_id"])) or {}
            run_workspace_id = str(voc_project.get("workspace_id") or "")
        source = ((run.get("sources") or {}).get("review_source_context") or {})
        source_archive_id = str(source.get("source_archive_id") or "")
        source_job_id = str(source.get("source_job_id") or "")
        run_id = str(run.get("id") or "")
        if (run_id in project_voc_runs or run_workspace_id == workspace_id
                or run_id in linked_run_ids or source_archive_id == archive_id
                or source_job_id in review_job_ids):
            result.append({**run, "workspace_id": run_workspace_id or workspace_id})
    return result
