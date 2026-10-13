from pathlib import Path
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font,PatternFill,Alignment
from .db import STORAGE,uid
from . import draft

STATUS={'pending':'待处理','failed':'失败','stale':'需要重算','unknown':'未知','review':'待确认','ready':'已完成'}

def safe(v):
    if isinstance(v,str) and v.startswith(('=','+','-','@')):return "'"+v
    if isinstance(v,(list,dict)):
        import json
        return json.dumps(v,ensure_ascii=False)
    return v

def export(did,selected,version=None,unknown_as_blank=False):
    data=draft.sheet(did,version);cols={c['id']:c for c in data['columns']}
    if len(selected)!=len(set(selected)) or any(c not in cols for c in selected):raise ValueError('导出列无效')
    headers=data['dataset']['headers'];wb=Workbook();ws=wb.active;ws.title='Product_Result'
    names={}
    for index,c in enumerate(data['columns']):
        name=c['definition']['name']
        if sum(x['definition']['name']==name for x in data['columns'])>1:
            legacy=c.get('origin',{}).get('schema_version')
            name+=f'（原规则{legacy}）' if legacy else f'（草稿列{index+1}）'
        names[c['id']]=name
    # An inactive adopted column means "show/export the untouched source field". It is not appended as an AI duplicate.
    selected=[cid for cid in selected if not cols[cid]['definition'].get('adopted_source') or cols[cid]['definition'].get('adoption_active',True)]
    adopted={c['definition']['source_header']:c['id'] for c in data['columns'] if c['id'] in selected and c['definition'].get('adopted_source')}
    appended=[cid for cid in selected if not cols[cid]['definition'].get('adopted_source')]
    ws.append(headers+[names[c] for c in appended])
    detail=wb.create_sheet('AI_Detail');detail.append(['商品标识','标识来源','原始行','标签','状态','原始观察','标准值','人工修正','证据与来源','定义修订','冲突候选'])
    dictionary=wb.create_sheet('Tag_Dictionary');dictionary.append(['标签','稳定ID','修订','定义','候选值','处理方式','依赖列','映射／区间','版本','类别ID','具体属性','判断规则','Skill版本','Skill摘要'])
    import json
    for cid in selected:
        c=cols[cid];d=c['definition'];dictionary.append([names[cid],cid,c['revision'],d['definition'],safe(d.get('options')),d.get('method','model'),safe(d.get('depends_on')),safe(d.get('mapping',d.get('bins'))),version or '当前草稿',d.get('parent_id'),d.get('attribute'),safe(d.get('judgment')),d.get('skill',{}).get('version'),d.get('skill',{}).get('sha256')])
    pic_index=next((i+1 for i,h in enumerate(headers) if h=='图片'),2)
    for index,row in enumerate(data['rows'],2):
        raw=row['raw'];values=[raw.get(h) for h in headers] if isinstance(raw,dict) else list(raw)
        values=[None if isinstance(v,str) and 'DISPIMG(' in v else safe(v) for v in values]
        for cid in selected:
            cell=row['cells'].get(cid,{})
            status=cell.get('status','pending');value=cell.get('value')
            if unknown_as_blank and status=='unknown':
                output=None
            elif cols[cid]['definition'].get('attribute'):
                output='✓' if value is True and status=='ready' else '❌' if value is False and status=='ready' else None
            else:output=safe(value) if status in ('ready','review') else f'【{STATUS.get(status,status)}】'+(str(value) if value is not None else '')
            source_header=cols[cid]['definition'].get('source_header') if cid in adopted.values() else None
            if source_header:values[headers.index(source_header)]=output
            else:values.append(output)
            detail.append([row['anchor'],row.get('identity_source'),row['row'],names[cid],STATUS.get(status,status),safe(cell.get('raw_value')),safe(value),safe(cell.get('manual')),safe(cell.get('evidence')),cell.get('revision'),safe(cell.get('conflict'))])
        ws.append(values);ws.row_dimensions[index].height=78
        for cid in selected:
            status=row['cells'].get(cid,{}).get('status','pending')
            source_header=cols[cid]['definition'].get('source_header') if cid in adopted.values() else None
            offset=headers.index(source_header)+1 if source_header else len(headers)+1+appended.index(cid)
            if cols[cid]['definition'].get('attribute') and status!='ready':
                from openpyxl.comments import Comment
                ws.cell(index,offset).comment=Comment(STATUS.get(status,status)+'；完整状态和依据见 AI_Detail','商品打标工作台')
        asset=next((a for a in row['assets'] if a.get('path') and Path(a['path']).is_file()),None)
        if asset:
            im=Image(asset['path']);ratio=min(96/im.width,96/im.height);im.width*=ratio;im.height*=ratio
            from openpyxl.utils import get_column_letter
            ws.add_image(im,f'{get_column_letter(pic_index)}{index}')
    from openpyxl.utils import get_column_letter
    ws.column_dimensions[get_column_letter(pic_index)].width=15
    for sh in wb:
        sh.freeze_panes='D2' if sh==ws else 'C2';sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244F43')
        for i in range(1,sh.max_column+1):
            if sh!=ws or i!=pic_index:sh.column_dimensions[get_column_letter(i)].width=23
        for row in sh:
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    path=STORAGE/'exports'/f'draft-{uid()}.xlsx';path.parent.mkdir(exist_ok=True);wb.save(path);return path
