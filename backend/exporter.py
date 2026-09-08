"""Application export service: portable Excel images plus three auditable sheets."""
import io
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter
from PIL import Image
from sqlalchemy import select
from .db import *
from .service import effective

def safe(value):
    if isinstance(value,str) and value.startswith(('=','+','-','@')): return "'"+value
    return value

def export_run(run_id):
    with Session.begin() as s:
        run=require(s,Run,run_id)
        if run.status!='completed': raise ValueError('任务未全部完成，不能导出正式结果')
        ds=require(s,Dataset,run.dataset_id);schema=require(s,Schema,run.schema_id)
        tags={t.product_id:t for t in s.scalars(select(Tag).where(Tag.run_id==run_id))}
        trans=s.scalars(select(Transform).where(Transform.run_id==run_id,Transform.status=='approved').order_by(Transform.created_at)).all()
        definitions=schema.definitions
        book=openpyxl.Workbook();result=book.active;result.title='Product_Result'
        detail=book.create_sheet('AI_Detail');dictionary=book.create_sheet('Tag_Dictionary')
        result.append(ds.headers+['打标_图片','打标_状态','打标_版本']+['标签_'+d['name']+' ['+d['key']+']' for d in definitions]+['派生_'+t.output_key for t in trans])
        detail.append(['商品标识','原始行','标签key','标签类型','原始观察','标准值','人工校准值','置信度','来源','证据','图片ID','审核状态','审核备注','任务ID','模型','标签版本'])
        dictionary.append(['标签key','名称','类型','定义与边界','候选值','来源要求','版本','转换来源','映射规则','确认时间'])
        image_col=len(ds.headers)+1
        for p in s.scalars(select(Product).where(Product.dataset_id==ds.id).order_by(Product.row)):
            if p.id not in tags: continue
            t=tags[p.id];values=effective(s,run_id,t)
            raw=[p.raw.get(h) for h in ds.headers]
            raw=['' if isinstance(v,str) and 'DISPIMG(' in v else safe(v) for v in raw]
            result.append(raw+['', '人工已审核' if t.review=='accepted' else '模型建议待复核',f'V{schema.version}']+[safe(values.get(d['key'])) for d in definitions]+[safe(values.get(x.output_key)) for x in trans])
            asset=s.scalar(select(Asset).where(Asset.product_id==p.id,Asset.path.is_not(None)))
            if asset and Path(asset.path).exists():
                with Image.open(asset.path) as im:
                    im.thumbnail((100,100));buf=io.BytesIO();im.convert('RGB').save(buf,format='PNG');buf.seek(0)
                pic=ExcelImage(buf);result.add_image(pic,f'{get_column_letter(image_col)}{result.max_row}');result.row_dimensions[result.max_row].height=78
            for v in t.values:
                d=next(d for d in definitions if d['key']==v['key'])
                detail.append([safe(x) for x in [p.anchor,p.row,v['key'],d['type'],v['raw_value'],v['normalized_value'],t.corrected.get(v['key']) if t.corrected is not None else None,v['confidence'],v['source'],v['evidence'],','.join(v['asset_ids']),t.review,t.review_note,run_id,run.model,f'V{schema.version}']])
            for tr in trans:
                detail.append([p.anchor,p.row,tr.output_key,tr.type,values.get(tr.source_key),values.get(tr.output_key),None,None,'transformation',f'{tr.source_key} → {tr.output_key}; {tr.definition}',None,'rule_approved',None,run_id,run.model,f'V{schema.version}'])
        for d in definitions: dictionary.append([d['key'],d['name'],d['type'],d['definition'],' / '.join(d['options']),d['source'],f'V{schema.version}',None,None,schema.approved_at])
        for t in trans: dictionary.append([t.output_key,t.output_key,t.type,t.definition,None,'transformation',t.id,t.source_key,json.dumps(t.mapping,ensure_ascii=False),t.approved_at])
        for ws in book:
            ws.freeze_panes='D2' if ws==result else 'C2';ws.auto_filter.ref=ws.dimensions
            ws.row_dimensions[1].height=32
            for cell in ws[1]: cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163D36');cell.alignment=Alignment(vertical='center')
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font=Font(name='微软雅黑',size=10);cell.alignment=Alignment(vertical='top',wrap_text=True)
                    if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F1F5F2')
            for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=22
        detail.column_dimensions['J'].width=65;dictionary.column_dimensions['D'].width=65
        result.column_dimensions[get_column_letter(image_col)].width=16
        target=STORAGE/'exports'/f'{run_id}-{uid()[:8]}.xlsx';target.parent.mkdir(exist_ok=True)
        book.save(target)
        audit(s,ds.id,'excel_exported',{'run_id':run_id,'path':target.name,'sheets':book.sheetnames})
        return target
