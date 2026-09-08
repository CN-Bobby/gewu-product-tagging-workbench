"""No model calls. Import supplied source files into a unique isolated database."""
import argparse
import hashlib
import json
import os
import sys
import time
import uuid
from pathlib import Path

root=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser();parser.add_argument('files',nargs='+',type=Path);args=parser.parse_args()
target=root/'storage'/'release-verification'/('data-'+uuid.uuid4().hex)
os.environ['TAG_STORAGE']=str(target);os.environ.pop('DATABASE_URL',None)
sys.path.insert(0,str(root))
from backend.db import init_db
from backend import importer,draft,draft_export
from backend.build_info import build_info
import openpyxl

init_db();report={'build':build_info(),'model_requests':0,'cases':[]}
for path in args.files:
    start=time.perf_counter();content=path.read_bytes();parsed=importer.parse(content)
    ds=importer.import_file(content,path.name);did=ds['id'];sheet=draft.sheet(did)
    assert sheet['dataset']['headers']==parsed['headers']
    assert len(sheet['rows'])==len(parsed['rows'])
    for expected,actual in zip(parsed['rows'],sheet['rows']):assert expected['raw']==actual['raw']
    proposal=draft.propose(did,'明确标识的验收列',[{'kind':'add','definition':{'key':'acceptance_only','name':'验收测试列','type':'FACT','source':'text','definition':'仅用于验证版本和导出，不是商品语义结论。','options':['测试甲','测试乙']}}])
    draft.confirm(did,proposal['id']);sheet=draft.sheet(did);cid=sheet['columns'][0]['id'];pid=sheet['rows'][0]['id']
    draft.correct(did,pid,cid,'测试甲');v1=draft.save_version(did);before=draft.sheet(did,v1['id'])
    draft.correct(did,pid,cid,'测试乙');v2=draft.save_version(did)
    assert before==draft.sheet(did,v1['id'])
    assert draft.compare(did,v1['id'],v2['id'])
    exported=draft_export.export(did,[cid],v1['id']);wb=openpyxl.load_workbook(exported);ws=wb['Product_Result']
    assert ws.max_column==len(parsed['headers'])+1
    assert ws.max_row==len(parsed['rows'])+1
    assert ws.cell(2,ws.max_column).value=='测试甲'
    numbers=0
    for index,row in enumerate(parsed['rows'],2):
        for col,h in enumerate(parsed['headers'],1):
            value=row['raw'][h]
            expected=None if isinstance(value,str) and 'DISPIMG(' in value else draft_export.safe(value)
            # XLSX serializes an empty string as a blank cell; original JSON above remains exact.
            if expected=='':expected=None
            actual=ws.cell(index,col)
            assert actual.value==expected,(index,h,actual.value,expected)
            if isinstance(value,(int,float)) and not isinstance(value,bool):assert actual.data_type=='n';numbers+=1
    image_count=sum(bool(row['images']) for row in parsed['rows'])
    assert len(ws._images)==image_count
    assert wb['Tag_Dictionary'].max_row==2
    assert wb['AI_Detail'].max_row==len(parsed['rows'])+1
    report['cases'].append({'source_sha256':hashlib.sha256(content).hexdigest(),'rows':len(parsed['rows']),'original_columns':len(parsed['headers']),'images':image_count,'numeric_cells':numbers,'v1_immutable':True,'raw_values_preserved':True,'export':str(exported),'seconds':time.perf_counter()-start})
(target/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'report':str(target/'report.json'),'cases':report['cases']},ensure_ascii=True,indent=2))
