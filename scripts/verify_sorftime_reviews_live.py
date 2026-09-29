"""Run one small live Amazon US review workflow in isolated workbench storage.

The Sorftime credential is read from stdin, stored only in the run's isolated
settings file, and cleared when the run exits. Model credentials are resolved
through the workbench's existing provider configuration without being printed.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import socket
import sys
import threading
import time
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any


EXPECTED_REVIEW_HEADERS = [
    "#", "来源站点", "品类节点", "ASIN", "商品标题", "变体", "评论日期", "星级", "情感",
    "评论标题(EN)", "评论内容(EN)", "评论标题(中文)", "评论内容(中文)", "拉取批次",
]
DEFAULT_ASINS = ["B0D5TQJXRN"]
DEFAULT_NODE = "17269957011"


def _safe(value: Any, secret: str = "") -> str:
    message = str(value or "验证失败")
    if secret:
        message = message.replace(secret, "[hidden]")
    message = re.sub(r"(?i)([?&](?:key|api[_-]?key|secret-key)=)[^&\s]+", r"\1[hidden]", message)
    message = re.sub(r"(?i)bearer\s+[^\s,;]+", "Bearer [hidden]", message)
    return message[:500]


def _choose_model(settings: dict[str, Any]) -> tuple[str, str]:
    providers = settings.get("providers") or []
    preferred = str(settings.get("default_model") or (settings.get("defaults") or [""])[0] or "")
    candidates = [preferred] + [ref for provider in providers for ref in provider.get("model_refs", [])
                                if ref != preferred]
    for model in candidates:
        if not model:
            continue
        provider_name = model.split("/", 1)[0]
        provider = next((item for item in providers if item.get("name") == provider_name), None)
        if not provider:
            continue
        credentials = []
        for item in provider.get("credentials", []):
            validation = item.get("validation")
            text_validation = validation.get("text") if isinstance(validation, dict) else None
            if (item.get("id") and item.get("enabled", True) and item.get("authorized", False)
                    and not (isinstance(text_validation, dict) and text_validation.get("status") == "failed")):
                credentials.append(item)
        chosen = next((item for item in credentials if item.get("id") == provider.get("primary_credential_id")),
                      credentials[0] if credentials else None)
        if chosen:
            return model, str(chosen["id"])
    raise RuntimeError("当前工作台没有已授权且可用的翻译模型 Key")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--key-stdin", action="store_true", help="read the Sorftime key from stdin; never pass it as a command argument")
    parser.add_argument("--endpoint", default="https://mcp.sorftime.com", help="Sorftime MCP HTTP endpoint")
    parser.add_argument("--site", default="US")
    parser.add_argument("--node-id", default=DEFAULT_NODE)
    parser.add_argument("--asins", nargs="+")
    parser.add_argument("--resume-run-id", help="resume one failed run in its existing isolated storage")
    parser.add_argument("--timeout-seconds", type=int, default=900)
    args = parser.parse_args()
    if not args.key_stdin:
        parser.error("--key-stdin is required so the API key is not placed in command history")
    if sys.stdin.isatty():
        import getpass
        key = getpass.getpass("Sorftime Key (hidden input): ").strip()
    else:
        key = sys.stdin.buffer.readline().decode("utf-8").rstrip("\r\n")
    if not key or len(key) > 2000:
        parser.error("stdin must contain one Sorftime key line")
    if args.timeout_seconds < 30 or args.timeout_seconds > 3600:
        parser.error("--timeout-seconds must be between 30 and 3600")

    repository = Path(__file__).resolve().parents[1]
    if str(repository) not in sys.path:
        sys.path.insert(0, str(repository))
    runs_root = (repository / "builds" / "sorftime-reviews-live").resolve()
    resumed = bool(args.resume_run_id)
    if resumed:
        run_id = args.resume_run_id
        if not re.fullmatch(r"\d{8}T\d{6}Z-[0-9a-f]{8}", run_id):
            parser.error("--resume-run-id must be a run ID created by this verifier")
        run_root = (runs_root / run_id).resolve()
        if not run_root.is_relative_to(runs_root) or not run_root.is_dir():
            parser.error("the requested isolated run does not exist")
        storage = (run_root / "workbench_data").resolve()
        if not storage.is_relative_to(runs_root) or not storage.is_dir():
            parser.error("the requested run has no isolated workbench storage")
        job_files = list((storage / "amazon-reviews").glob("*/job.json"))
        if len(job_files) != 1:
            parser.error("resume requires exactly one saved Amazon Reviews job")
        job_path = job_files[0]
        job_id = job_path.parent.name
        saved_job = json.loads(job_path.read_text(encoding="utf-8"))
        saved_request = json.loads((job_path.parent / "request.json").read_text(encoding="utf-8"))
        asins = list(dict.fromkeys(str(value).strip().upper() for value in saved_request.get("asins", [])))
        node_id = str(saved_request.get("node_id") or "").strip()
        if (saved_request.get("site") != "US" or saved_request.get("mode") != "node_asins"
                or not node_id or not asins
                or any(not re.fullmatch(r"B[0-9A-Z]{9}", value) for value in asins)
                or saved_job.get("id") != job_id or saved_job.get("asins") != asins
                or saved_job.get("node_id") != node_id or saved_job.get("site") != "US"):
            parser.error("saved run is not a consistent Amazon US node review job")
        if saved_job.get("status") in ("queued", "running", "cancelling"):
            parser.error("saved Amazon Reviews job is still active")
    else:
        run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid.uuid4().hex[:8]
        asins = list(dict.fromkeys(value.strip().upper() for value in (args.asins or DEFAULT_ASINS)))
        node_id = str(args.node_id).strip()
        if not asins or any(not re.fullmatch(r"B[0-9A-Z]{9}", value) for value in asins):
            parser.error("ASINs must be Amazon 10-character ASINs")
        if args.site.upper() != "US":
            parser.error("this live acceptance script is restricted to Amazon US")
        if not node_id or len(node_id) > 100:
            parser.error("--node-id must be a valid Amazon category Node ID")
        run_root = runs_root / run_id
        storage = run_root / "workbench_data"
        run_root.mkdir(parents=True, exist_ok=False)
        storage.mkdir()
    if not storage.resolve().is_relative_to(runs_root):
        raise RuntimeError("isolated storage path escaped the test-run directory")

    os.environ["TAG_STORAGE"] = str(storage)
    os.environ["DATABASE_URL"] = "sqlite:///" + str(storage / "workbench.db")
    dist = repository / "frontend" / "dist"
    if dist.is_dir():
        os.environ["TAG_FRONTEND_DIST"] = str(dist)

    result: dict[str, Any] = {
        "run_id": run_id,
        "site": "US",
        "node_id": node_id,
        "asins_requested": asins,
        "mode": "node_asins",
        "resumed_existing_run": resumed,
        "storage": str(storage),
        "started_at": datetime.now(timezone.utc).isoformat(),
        "status": "running",
    }
    server = None
    server_thread = None
    client = None
    sorftime_settings = None
    isolated_credential_cleared = False
    try:
        import httpx
        import uvicorn
        from openpyxl import load_workbook

        # TAG_STORAGE is set before importing any backend modules.
        from backend import app as workbench_app
        from backend.xunlongjue import settings as sorftime_settings

        sorftime_settings.save(args.endpoint, key)
        socket_handle = socket.socket()
        socket_handle.bind(("127.0.0.1", 0))
        port = socket_handle.getsockname()[1]
        socket_handle.close()
        server = uvicorn.Server(uvicorn.Config(
            workbench_app.app, host="127.0.0.1", port=port,
            log_level="warning", access_log=False,
        ))
        server_thread = threading.Thread(target=server.run, name="reviews-live-verifier", daemon=True)
        server_thread.start()
        startup_deadline = time.monotonic() + 45
        while not server.started and server_thread.is_alive() and time.monotonic() < startup_deadline:
            time.sleep(0.1)
        if not server.started:
            raise RuntimeError("isolated workbench did not start")

        base = f"http://127.0.0.1:{port}"
        client = httpx.Client(base_url=base, timeout=45)
        page = client.get("/")
        page.raise_for_status()
        health = client.get("/api/health")
        health.raise_for_status()
        if health.json().get("storage") != str(storage):
            raise RuntimeError("workbench health response did not point at isolated test storage")
        public_settings_response = client.get("/api/amazon-reviews/settings")
        public_settings_response.raise_for_status()
        public_settings = public_settings_response.json()
        if key in json.dumps(public_settings, ensure_ascii=False):
            raise RuntimeError("API settings response exposed the Sorftime credential")
        if not public_settings.get("sorftime", {}).get("key_configured"):
            raise RuntimeError("isolated Sorftime connection was not configured")
        model, credential_id = _choose_model(public_settings)
        if resumed and (model != saved_job.get("model")
                        or credential_id != saved_job.get("credential_id")):
            raise RuntimeError("the saved job model credential differs from the currently authorized workbench configuration")
        result.update({"model": model, "credential_id": credential_id, "sorftime_configured": True})

        if resumed:
            job_id = saved_job["id"]
            initial_response = client.get(f"/api/amazon-reviews/jobs/{job_id}")
            initial_response.raise_for_status()
            snapshot = initial_response.json()
            if [item.get("asin") for item in snapshot.get("items", [])] != asins:
                raise RuntimeError("saved job ASINs differ from the original request")
            metrics_before_retry = int((snapshot.get("metrics") or {}).get("review_tool_calls") or 0)
        else:
            created_response = client.post("/api/amazon-reviews/jobs", json={
                "mode": "node_asins", "site": "US", "node_id": node_id,
                "asins": asins, "model": model, "credential_id": credential_id,
            })
            created_response.raise_for_status()
            created = created_response.json()
            job_id = created["id"]
            snapshot = created
            metrics_before_retry = 0
        result["job_id"] = job_id
        result["started_at"] = datetime.now(timezone.utc).isoformat()
        if not resumed:
            deadline = time.monotonic() + args.timeout_seconds
            while time.monotonic() < deadline:
                snapshot_response = client.get(f"/api/amazon-reviews/jobs/{job_id}")
                snapshot_response.raise_for_status()
                snapshot = snapshot_response.json()
                if snapshot.get("status") in ("completed", "completed_with_errors", "failed", "cancelled"):
                    break
                time.sleep(2)
            else:
                raise TimeoutError("live job exceeded the configured deadline")
        retry_attempts = 0
        max_retry_attempts = 1 if resumed else 2
        while (snapshot.get("status") != "completed" and retry_attempts < max_retry_attempts
               and int(snapshot.get("retryable") or 0) > 0):
            retry_response = client.post(f"/api/amazon-reviews/jobs/{job_id}/retry", json={})
            retry_response.raise_for_status()
            retry_attempts += 1
            retry_deadline = time.monotonic() + args.timeout_seconds
            while time.monotonic() < retry_deadline:
                snapshot_response = client.get(f"/api/amazon-reviews/jobs/{job_id}")
                snapshot_response.raise_for_status()
                snapshot = snapshot_response.json()
                if snapshot.get("status") in ("completed", "completed_with_errors", "failed", "cancelled"):
                    break
                time.sleep(2)
            else:
                raise TimeoutError("live retry exceeded the configured deadline")
        result["retry_attempts"] = retry_attempts
        review_call_delta = int((snapshot.get("metrics") or {}).get("review_tool_calls") or 0) - metrics_before_retry
        result["review_call_delta_on_resume"] = review_call_delta if resumed else None
        if resumed and review_call_delta != 0:
            raise RuntimeError("resume unexpectedly made additional Sorftime review API calls")
        if snapshot.get("status") != "completed":
            raise RuntimeError(f"live job status={snapshot.get('status')}: {_safe(snapshot.get('error') or snapshot.get('errors'), key)}")
        if [item.get("asin") for item in snapshot.get("items", [])] != asins:
            raise RuntimeError("category filter did not return exactly the requested ASINs")

        raw_response = client.get(f"/api/amazon-reviews/jobs/{job_id}/raw")
        raw_response.raise_for_status()
        raw_bytes = raw_response.content
        raw_path = run_root / "Amazon-Reviews-raw.json"
        raw_path.write_bytes(raw_bytes)
        raw = raw_response.json()
        excel_response = client.get(f"/api/amazon-reviews/jobs/{job_id}/download")
        excel_response.raise_for_status()
        excel_bytes = excel_response.content
        excel_path = run_root / "Amazon-Reviews.xlsx"
        excel_path.write_bytes(excel_bytes)
        if key.encode("utf-8") in raw_bytes or key.encode("utf-8") in excel_bytes:
            raise RuntimeError("an output artifact contains the Sorftime credential")

        reviews = raw.get("reviews") or []
        metrics = snapshot.get("metrics") or {}
        expected_calls = len(asins) * 2
        if metrics.get("review_tool_calls") != expected_calls:
            raise RuntimeError(f"expected {expected_calls} review calls; received {metrics.get('review_tool_calls')}")
        if metrics.get("new_unique_records") != len(reviews):
            raise RuntimeError("deduplicated review count does not reconcile with call metrics")
        for item in snapshot.get("items", []):
            pass_metrics = item.get("review_call_metrics") or []
            if len(pass_metrics) != 2 or any(p.get("review_type") != "Both" for p in pass_metrics):
                raise RuntimeError(f"{item.get('asin')} did not complete two Both calls")
            if any(int(p.get("raw_record_count") or 0) > 100 for p in pass_metrics):
                raise RuntimeError(f"{item.get('asin')} exceeded Sorftime's 100-review response limit")
        if any((row.get("content") or row.get("review") or row.get("text")) and not
               (row.get("t_content") or row.get("translation_zh")) for row in reviews):
            raise RuntimeError("at least one non-empty review body is missing its Chinese translation")
        if any(row.get("title") and not (row.get("t_title") or row.get("title_zh")) for row in reviews):
            raise RuntimeError("at least one non-empty review title is missing its Chinese translation")
        if any(row.get("source_site") != "US" or row.get("source_node") != node_id for row in reviews):
            raise RuntimeError("review source metadata does not match the selected US category")

        book = load_workbook(BytesIO(excel_bytes), data_only=True)
        if book.sheetnames != ["全部评论", "商品评论统计", "统计汇总"]:
            raise RuntimeError(f"unexpected workbook sheets: {book.sheetnames}")
        review_sheet = book["全部评论"]
        if [review_sheet.cell(1, column).value for column in range(1, 15)] != EXPECTED_REVIEW_HEADERS:
            raise RuntimeError("review workbook columns do not match the source Skill")
        if review_sheet.max_row - 1 != len(reviews):
            raise RuntimeError("Excel review row count does not match raw JSON")
        summary = {
            book["统计汇总"].cell(row, 1).value: book["统计汇总"].cell(row, 2).value
            for row in range(2, book["统计汇总"].max_row + 1)
        }
        if summary.get("去重后评论总数") != len(reviews) or summary.get("API 调用次数") != expected_calls:
            raise RuntimeError("summary sheet counts do not reconcile")
        if summary.get("翻译覆盖") != f"{len(reviews)}/{len(reviews)} 条":
            raise RuntimeError("summary sheet reports incomplete translations")

        samples = []
        for row in reviews[:5]:
            samples.append({
                "asin": row.get("asin"),
                "stars": row.get("star_rating"),
                "title_en": str(row.get("title") or "")[:140],
                "title_zh": str(row.get("t_title") or row.get("title_zh") or "")[:140],
                "content_en": str(row.get("content") or "")[:220],
                "content_zh": str(row.get("t_content") or row.get("translation_zh") or "")[:220],
            })
        result.update({
            "status": "passed",
            "job_status": snapshot["status"],
            "retry_attempts": retry_attempts,
            "products_completed": snapshot.get("completed", 0),
            "review_count": len(reviews),
            "metrics": metrics,
            "workbook_sheets": book.sheetnames,
            "excel_rows": review_sheet.max_row - 1,
            "translation_coverage": f"{len(reviews)}/{len(reviews)}",
            "artifacts": {"xlsx": str(excel_path), "raw_json": str(raw_path)},
            "translation_samples": samples,
            "finished_at": datetime.now(timezone.utc).isoformat(),
        })
    except Exception as exc:  # noqa: BLE001
        result.update({
            "status": "failed",
            "error": _safe(exc, key),
            "finished_at": datetime.now(timezone.utc).isoformat(),
        })
    finally:
        if client is not None:
            client.close()
        if server is not None:
            server.should_exit = True
        if server_thread is not None:
            server_thread.join(timeout=20)
        if sorftime_settings is not None:
            try:
                sorftime_settings.clear_key()
                isolated_credential_cleared = not sorftime_settings.status().get("key_configured")
            except Exception:
                isolated_credential_cleared = False
        credential_artifacts = []
        if isolated_credential_cleared:
            secret_bytes = key.encode("utf-8")
            for path in storage.rglob("*"):
                if path.is_file() and path.suffix.casefold() in (".json", ".txt", ".log"):
                    content = path.read_bytes()
                    if secret_bytes in content:
                        path.write_bytes(content.replace(secret_bytes, b"[hidden]"))
                        credential_artifacts.append(str(path.relative_to(run_root)))
        if credential_artifacts:
            result["status"] = "failed"
            result["error"] = "the Sorftime credential was redacted from an unexpected isolated artifact"
            result["redacted_artifacts"] = credential_artifacts
        result["isolated_credential_cleared"] = isolated_credential_cleared
        report_name = ("verification-report-resume-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + ".json"
                       if resumed else "verification-report.json")
        report_path = run_root / report_name
        result["report"] = str(report_path)
        report_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    for sample in result.get("translation_samples", []):
        print(json.dumps(sample, ensure_ascii=False))
    print(json.dumps({key_name: result.get(key_name) for key_name in (
        "status", "job_status", "site", "node_id", "asins_requested", "review_count",
        "retry_attempts", "metrics", "workbook_sheets", "translation_coverage", "isolated_credential_cleared",
        "artifacts", "report", "error")}, ensure_ascii=False, indent=2))
    if not isolated_credential_cleared:
        print("Isolated test credential cleanup could not be verified.", file=sys.stderr)
        return 1
    return 0 if result.get("status") == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
