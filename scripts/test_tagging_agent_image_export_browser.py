"""Run a bounded Agent + image-label + version-export flow in a local DEV browser.

This is a real, paid-capable acceptance flow. Invoke it only after the current DEV
instance is confirmed as live_validation and the current run has been authorized.
The script never changes runtime policy, reads credentials, or contacts non-local
services from the browser. The backend may download the fixture's product-image
URLs and call the configured text/vision models during this explicitly requested run.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


MAX_PRODUCTS_HARD_LIMIT = 3
POLL_SECONDS = 2
ACTIVE_STATUSES = {"queued", "running", "pausing", "cancelling"}
IMAGE_SOURCE_VALUES = {"image", "both"}


class E2EFailure(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _json_response(response, label: str) -> dict:
    if not response.ok:
        raise E2EFailure(f"{label}_http_{response.status}", f"{label} 请求失败（HTTP {response.status}）")
    try:
        value = response.json()
    except Exception as error:
        raise E2EFailure(f"{label}_invalid_json", f"{label} 返回了无效 JSON") from error
    if not isinstance(value, dict):
        raise E2EFailure(f"{label}_invalid_shape", f"{label} 返回格式不正确")
    return value


def _get_json(request, url: str, label: str) -> dict:
    try:
        response = request.get(url, timeout=20_000)
    except Exception as error:
        raise E2EFailure(f"{label}_unavailable", f"无法读取 {label}；请确认本机 DEV 服务正在运行") from error
    return _json_response(response, label)


def _post_inspect(request, base_url: str, fixture: Path, sheet: str | None = None) -> dict:
    multipart: dict = {
        "file": {
            "name": fixture.name,
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": fixture.read_bytes(),
        }
    }
    if sheet:
        multipart["sheet"] = sheet
    try:
        response = request.post(f"{base_url}/api/import/inspect", multipart=multipart, timeout=60_000)
    except Exception as error:
        raise E2EFailure("fixture_inspect_unavailable", "无法在本机 DEV 中检查 fixture 工作簿") from error
    return _json_response(response, "fixture_inspect")


def _normalize_header(value: object) -> str:
    return re.sub(r"[\s_\-]+", "", str(value or "")).casefold()


def _find_header_row(ws, headers: list[str]) -> int:
    wanted = [_normalize_header(value) for value in headers]
    for row_index in range(1, min(ws.max_row, 40) + 1):
        values = [_normalize_header(ws.cell(row_index, col).value) for col in range(1, len(headers) + 1)]
        if values == wanted:
            return row_index
    raise E2EFailure("fixture_header_row_not_found", "无法定位商品表头行；请提供规范的商品工作表")


def _choose_product_sheet(inspection: dict, explicit_sheet: str | None) -> str:
    if explicit_sheet:
        sheets = inspection.get("sheets") or []
        if explicit_sheet not in sheets:
            raise E2EFailure("fixture_sheet_missing", "--sheet 指定的工作表不存在")
        return explicit_sheet
    if not inspection.get("sheet_required") and inspection.get("sheet"):
        return str(inspection["sheet"])
    candidates = [str(value) for value in inspection.get("candidate_sheets") or []
                  if str(value).strip().casefold() != "notes"]
    if len(candidates) == 1:
        return candidates[0]
    raise E2EFailure("fixture_sheet_ambiguous", "无法唯一识别商品工作表；请用 --sheet 指定")


def _create_bounded_fixture(source: Path, sheet_name: str, headers: list[str], max_products: int,
                            temp_dir: Path) -> tuple[Path, dict]:
    """Keep the product sheet and Notes page, trimming imported products to the cap."""
    try:
        workbook = load_workbook(source, data_only=False, read_only=False)
    except Exception as error:
        raise E2EFailure("fixture_open_failed", "无法用 openpyxl 读取 fixture 工作簿") from error
    try:
        if sheet_name not in workbook.sheetnames:
            raise E2EFailure("fixture_sheet_missing", "商品工作表在 fixture 中不存在")
        notes_name = next((name for name in workbook.sheetnames if name.strip().casefold() == "notes"), None)
        if notes_name is None:
            raise E2EFailure("fixture_notes_sheet_missing", "fixture 必须包含 Notes 工作表")
        product_ws = workbook[sheet_name]
        notes_ws = workbook[notes_name]
        header_row = _find_header_row(product_ws, headers)
        data_rows = [row_index for row_index in range(header_row + 1, product_ws.max_row + 1)
                     if any(product_ws.cell(row_index, col).value not in (None, "")
                            for col in range(1, len(headers) + 1))]
        if len(data_rows) < max_products:
            raise E2EFailure("fixture_rows_too_few", f"fixture 商品行不足 {max_products} 行")
        kept_rows = data_rows[:max_products]
        # Parser-backed workbooks use contiguous product rows; reject sparse layouts rather than
        # silently pairing a later image URL with the wrong product after trimming.
        if kept_rows != list(range(header_row + 1, header_row + 1 + max_products)):
            raise E2EFailure("fixture_rows_not_contiguous", "商品行之间有空行；请先整理 fixture 后重试")
        for row_index in range(header_row + 1 + max_products, product_ws.max_row + 1):
            for cell in product_ws[row_index]:
                cell.value = None
                cell._style = product_ws.cell(header_row + max_products, cell.column)._style
                cell.number_format = product_ws.cell(header_row + max_products, cell.column).number_format
        if product_ws.max_row > header_row + max_products:
            product_ws.delete_rows(header_row + max_products + 1,
                                    product_ws.max_row - (header_row + max_products))
        # Remove unrelated tabs so the application sees exactly the selected product sheet and
        # the deliberately non-product Notes tab.
        for name in list(workbook.sheetnames):
            if name not in {sheet_name, notes_name}:
                workbook.remove(workbook[name])
        workbook._active = workbook.index(product_ws)
        output = temp_dir / "bounded-fixture.xlsx"
        workbook.save(output)
    finally:
        workbook.close()
    return output, {"product_sheet": sheet_name, "notes_sheet": notes_name, "row_count": max_products}


def _validate_fixture_content(fixture: Path, inspection: dict, sheet_name: str, max_products: int) -> None:
    if int(inspection.get("count") or 0) != max_products:
        raise E2EFailure("fixture_row_count_mismatch", "裁剪后的商品行数与 max-products 不一致")
    if int(inspection.get("embedded_images") or 0) != 0:
        raise E2EFailure("fixture_expected_url_images", "本流程要求用商品主图 URL 覆盖导入补图路径")
    headers = [str(value) for value in inspection.get("headers") or []]
    mapping = inspection.get("mapping") or {}
    image_header = mapping.get("image_url")
    if not image_header or image_header not in headers:
        raise E2EFailure("fixture_image_url_mapping_missing", "商品工作表未识别到图片 URL 列")
    try:
        workbook = load_workbook(fixture, data_only=True, read_only=True)
        ws = workbook[sheet_name]
        header_row = _find_header_row(ws, headers)
        rows = list(ws.iter_rows(min_row=header_row + 1, max_row=header_row + max_products,
                                 min_col=1, max_col=len(headers), values_only=True))
    except E2EFailure:
        raise
    except Exception as error:
        raise E2EFailure("fixture_validate_failed", "无法校验裁剪后的 fixture 内容") from error
    finally:
        try:
            workbook.close()
        except Exception:
            pass
    positions = {header: index for index, header in enumerate(headers)}
    url_count = 0
    text_input_count = 0
    for row in rows:
        url = str(row[positions[image_header]] or "").strip()
        if urlsplit(url).scheme.lower() not in {"http", "https"} or not urlsplit(url).netloc:
            raise E2EFailure("fixture_image_url_invalid", "每个目标商品都必须有有效的 HTTP(S) 主图 URL")
        url_count += 1
        for role in ("sku", "parameters", "title"):
            header = mapping.get(role)
            if header and str(row[positions[header]] or "").strip():
                text_input_count += 1
                break
    if url_count != max_products:
        raise E2EFailure("fixture_image_url_count", "fixture 中可用主图 URL 数量不足")
    if text_input_count == 0:
        raise E2EFailure("fixture_text_input_missing", "至少一行需要有 SKU、详细参数或标题，才能覆盖导入文字整理")


def _check_runtime(request, base_url: str) -> dict:
    health = _get_json(request, f"{base_url}/api/health", "health")
    if health.get("status") != "ok" or health.get("application") != "product-tagging-workbench":
        raise E2EFailure("dev_health_invalid", "base URL 未返回格物工作台健康状态")
    capability = health.get("capabilities") or {}
    if capability.get("visual_tagging") is not True:
        raise E2EFailure("vision_disabled", "当前 DEV 未启用自动视觉打标")
    runtime = _get_json(request, f"{base_url}/api/runtime/status", "runtime_status")
    if str(runtime.get("profile") or "").casefold() != "dev":
        raise E2EFailure("not_dev_profile", "base URL 的运行身份不是 DEV")
    effects = runtime.get("external_effects") or {}
    services = effects.get("services") or {}
    if effects.get("policy") != "live_validation" or services.get("model_api") is not True:
        raise E2EFailure("dev_not_live_validation", "DEV 未以 live_validation 开启模型调用；脚本不会切换运行策略")
    settings = _get_json(request, f"{base_url}/api/settings/models", "model_settings")
    defaults = settings.get("defaults") or []
    if len(defaults) < 2:
        raise E2EFailure("model_defaults_missing", "DEV 尚未配置文本和视觉模型")
    providers = {str(item.get("name")): item for item in settings.get("providers") or []}
    for kind, model_ref in (("text", defaults[0]), ("vision", defaults[1])):
        provider_name = str(model_ref).partition("/")[0]
        provider = providers.get(provider_name)
        if not provider or not provider.get("base_url"):
            raise E2EFailure(f"{kind}_provider_missing", f"默认{kind}模型服务未配置")
        credentials = provider.get("credentials") or []
        authorized = any(item.get("authorized") is True and item.get("enabled", True) is True
                         for item in credentials)
        if not authorized and provider.get("authorized") is not True:
            raise E2EFailure(f"{kind}_credential_missing", f"默认{kind}模型没有已启用的授权凭据")
        validation = provider.get("validation") or {}
        visual_validation = validation.get("vision") or validation.get("image")
        if kind == "vision":
            vision_validation_status = (visual_validation or {}).get("status") or "unknown"
    return {
        "profile": "dev",
        "policy": "live_validation",
        "instance_id": runtime.get("instance_id"),
        "build_id": (runtime.get("source") or {}).get("build_id"),
        "commit": (runtime.get("source") or {}).get("commit"),
        "visual_tagging": True,
        "text_credential_available": True,
        "vision_credential_available": True,
        "vision_validation_status": vision_validation_status,
    }


def _poll_state(request, base_url: str, dataset_id: str, timeout_seconds: int,
                predicate, timeout_code: str, timeout_message: str) -> dict:
    deadline = time.monotonic() + timeout_seconds
    last_state: dict = {}
    while time.monotonic() < deadline:
        last_state = _get_json(request, f"{base_url}/api/drafts/{dataset_id}", "dataset_state")
        if predicate(last_state):
            return last_state
        time.sleep(POLL_SECONDS)
    raise E2EFailure(timeout_code, timeout_message)


def _wait_import_enrichment(request, base_url: str, dataset_id: str, max_products: int,
                            timeout_seconds: int) -> dict:
    def settled(state: dict) -> bool:
        jobs = state.get("jobs") or []
        return bool(jobs) and not any(job.get("status") in ACTIVE_STATUSES for job in jobs)

    state = _poll_state(request, base_url, dataset_id, timeout_seconds, settled,
                        "import_enrichment_timeout", "导入文字/图片整理任务超时")
    jobs = state.get("jobs") or []
    text_job = next((job for job in reversed(jobs) if (job.get("spec") or {}).get("kind") == "text_facts"), None)
    if not text_job:
        raise E2EFailure("import_enrichment_missing", "没有找到导入文字/图片整理任务")
    metrics = text_job.get("metrics") or {}
    rows = state.get("rows") or []
    if len(rows) != max_products:
        raise E2EFailure("imported_row_count_mismatch", "导入后的商品行数与 max-products 不一致")
    images_available = sum(any(asset.get("path") for asset in row.get("assets") or []) for row in rows)
    return {"status": text_job.get("status"), "image_downloads": int(metrics.get("images_completed") or 0),
            "image_failures": int(metrics.get("images_failed") or 0), "images_available": images_available,
            "text_items_completed": int(metrics.get("attributes_completed") or 0),
            "text_items_expected": int(metrics.get("attributes_expected") or 0)}


def _verify_import_enrichment(state: dict, max_products: int) -> dict:
    jobs = state.get("jobs") or []
    text_job = next((job for job in reversed(jobs) if (job.get("spec") or {}).get("kind") == "text_facts"), None)
    if not text_job or text_job.get("status") != "completed":
        raise E2EFailure("import_enrichment_failed", "重试后导入文字/图片整理仍未完整成功")
    metrics = text_job.get("metrics") or {}
    rows = state.get("rows") or []
    if len(rows) != max_products:
        raise E2EFailure("imported_row_count_mismatch", "导入后的商品行数与 max-products 不一致")
    images_available = sum(any(asset.get("path") for asset in row.get("assets") or []) for row in rows)
    if images_available != max_products or int(metrics.get("images_failed") or 0) != 0:
        raise E2EFailure("local_images_missing", "重试后仍有商品没有可用的本地图片")
    return {"status": text_job.get("status"), "image_downloads": int(metrics.get("images_completed") or 0),
            "image_failures": int(metrics.get("images_failed") or 0), "images_available": images_available,
            "text_items_completed": int(metrics.get("attributes_completed") or 0),
            "text_items_expected": int(metrics.get("attributes_expected") or 0)}


def _visual_proposal(state: dict, max_products: int) -> tuple[dict, str, list[str]]:
    proposals = [item for item in state.get("proposals") or [] if item.get("status") == "pending"]
    for proposal in reversed(proposals):
        changes = proposal.get("changes") or []
        image_defs = []
        has_execute = False
        scope_ids: list[str] = []
        for change in changes:
            definition = change.get("definition") or {}
            if change.get("kind") in {"add", "revise"} and definition.get("source") in IMAGE_SOURCE_VALUES:
                image_defs.append(definition)
            if change.get("kind") == "execute":
                has_execute = True
                ids = change.get("product_ids") or []
                if not ids:
                    for target in change.get("targets") or []:
                        ids.extend(target.get("product_ids") or [])
                scope_ids.extend(str(value) for value in ids)
        if len(image_defs) == 1 and has_execute:
            distinct_scope = list(dict.fromkeys(scope_ids))
            if len(distinct_scope) > max_products:
                raise E2EFailure("agent_proposal_exceeds_cap", "Agent 提案商品范围超过 max-products 上限")
            name = str(image_defs[0].get("name") or "").strip()
            if not name:
                raise E2EFailure("agent_proposal_label_missing", "Agent 图片标签提案没有标签名称")
            return proposal, name, distinct_scope
    raise E2EFailure("agent_visual_proposal_missing", "Agent 未形成包含图片标签和执行范围的待确认提案")


def _wait_label_job(request, base_url: str, dataset_id: str, label_name: str, max_products: int,
                    timeout_seconds: int) -> tuple[dict, dict]:
    def ready(state: dict) -> bool:
        for job in reversed(state.get("jobs") or []):
            spec = job.get("spec") or {}
            if spec.get("kind") == "text_facts":
                continue
            columns = spec.get("columns") or []
            if any((column.get("definition") or {}).get("name") == label_name for column in columns):
                if job.get("status") not in ACTIVE_STATUSES:
                    return True
        return False

    state = _poll_state(request, base_url, dataset_id, timeout_seconds, ready,
                        "image_label_timeout", "图片标签执行超时")
    job = next((item for item in reversed(state.get("jobs") or [])
                if (item.get("spec") or {}).get("kind") != "text_facts"
                and any((column.get("definition") or {}).get("name") == label_name
                        for column in (item.get("spec") or {}).get("columns") or [])), None)
    if not job or job.get("status") != "completed":
        raise E2EFailure("image_label_failed", "图片标签任务未完整成功")
    spec = job.get("spec") or {}
    product_ids = spec.get("product_ids") or []
    if not product_ids:
        product_ids = [product_id for target in spec.get("targets") or []
                       for product_id in target.get("product_ids") or []]
    product_ids = list(dict.fromkeys(str(value) for value in product_ids))
    product_count = int(spec.get("product_count") or job.get("expected_count") or len(product_ids))
    if not product_ids and product_count == len(state.get("rows") or []):
        # The current job contract records the bounded row count rather than persisting row IDs.
        # This fixture contains only the rows submitted for the run, so all rows are in scope.
        product_ids = [str(row.get("id")) for row in state.get("rows") or [] if row.get("id")]
    if not product_ids or len(product_ids) > max_products or product_count > max_products:
        raise E2EFailure("image_label_scope_invalid", "图片标签任务范围为空或超过产品上限")
    metrics = job.get("metrics") or {}
    column = next(column for column in (spec.get("columns") or [])
                  if (column.get("definition") or {}).get("name") == label_name)
    column_id = str(column.get("id") or "")
    rows = state.get("rows") or []
    target_rows = [row for row in rows if str(row.get("id")) in set(product_ids)]
    if len(target_rows) != len(product_ids) or any(
            (row.get("cells") or {}).get(column_id, {}).get("status") not in {"ready", "unknown"}
            for row in target_rows):
        raise E2EFailure("image_label_results_missing", "执行完成但目标商品没有可验证的图片标签结果")
    return job, {"status": job.get("status"), "product_count": len(product_ids),
                 "requests": int(metrics.get("requests") or 0),
                 "vision_requests": int(metrics.get("vision_requests") or 0),
                 "cache_hits": int(metrics.get("cache_hits") or 0)}


def _default_output_dir() -> Path:
    root = Path(os.environ.get("LOCALAPPDATA") or (Path.home() / ".local" / "share"))
    return root / "GeWuWorkbench" / "reports" / "agent-image-export-browser"


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", required=True, help="当前 DEV 页面 URL，例如 http://127.0.0.1:5174")
    parser.add_argument("--fixture", required=True, type=Path, help="含商品主图 URL 和 Notes 页的 .xlsx fixture")
    parser.add_argument("--sheet", help="商品工作表名；不填时要求能从候选工作表中唯一识别")
    parser.add_argument("--prompt", required=True, help="本轮已授权的中文 Agent 指令；报告只记录摘要哈希")
    parser.add_argument("--max-products", type=int, default=2,
                        help="导入并打标的商品数，默认 2，硬上限 3")
    parser.add_argument("--output-dir", type=Path, default=_default_output_dir(),
                        help="报告和导出文件目录；默认写入用户本机 LocalAppData")
    parser.add_argument("--timeout-seconds", type=int, default=1800,
                        help="每个异步步骤的最长等待秒数")
    parser.add_argument("--headful", action="store_true", help="显示 Chromium 窗口")
    args = parser.parse_args()
    parsed = urlsplit(args.base_url)
    if (parsed.scheme != "http" or parsed.hostname not in {"127.0.0.1", "localhost", "::1"}
            or parsed.username or parsed.password or parsed.path not in ("", "/")
            or parsed.query or parsed.fragment):
        parser.error("--base-url 必须是本机 DEV 的无凭据 HTTP origin（localhost 或 127.0.0.1）")
    if not args.fixture.is_file() or args.fixture.suffix.casefold() != ".xlsx":
        parser.error("--fixture 必须是现存的 .xlsx 文件")
    if not 1 <= args.max_products <= MAX_PRODUCTS_HARD_LIMIT:
        parser.error(f"--max-products 必须在 1 到 {MAX_PRODUCTS_HARD_LIMIT} 之间")
    if args.timeout_seconds < 60:
        parser.error("--timeout-seconds 最少为 60")
    if not args.prompt.strip():
        parser.error("--prompt 不能为空")
    if not re.search(r"[\u4e00-\u9fff]", args.prompt):
        parser.error("--prompt 必须是本轮已授权的中文指令")
    args.base_url = f"{parsed.scheme}://{parsed.netloc}"
    args.fixture = args.fixture.resolve()
    args.output_dir = args.output_dir.expanduser().resolve()
    return args


def main() -> int:
    args = _parse_args()
    run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid.uuid4().hex[:8]
    report = {
        "schema": 1,
        "run_id": run_id,
        "started_at": datetime.now(timezone.utc).isoformat(),
        "status": "failed",
        "base_origin": args.base_url,
        "max_products": args.max_products,
        "prompt_sha256": hashlib.sha256(args.prompt.encode("utf-8")).hexdigest(),
        "prompt_characters": len(args.prompt),
        "browser_external_requests_blocked": True,
        "runtime": {},
        "checks": {},
        "artifacts": {},
    }
    stage = "output_setup"
    browser = None
    page = None
    session_active = False
    temp_root = None
    report_write_failed = False
    try:
        args.output_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="gewu-agent-e2e-") as temp_name:
            temp_root = Path(temp_name)
            stage = "browser_startup"
            with sync_playwright() as playwright:
                try:
                    browser = playwright.chromium.launch(headless=not args.headful)
                except Exception as error:
                    raise E2EFailure("playwright_chromium_missing", "Playwright Chromium 不可用；请安装项目浏览器依赖") from error
                context = browser.new_context(accept_downloads=True, viewport={"width": 1500, "height": 1000})
                external_browser_count = 0

                def local_only(route):
                    nonlocal external_browser_count
                    target = urlsplit(route.request.url)
                    if target.scheme in {"data", "blob", "about"}:
                        route.continue_()
                    elif target.hostname in {"127.0.0.1", "localhost", "::1"}:
                        route.continue_()
                    else:
                        external_browser_count += 1
                        route.abort()

                context.route("**/*", local_only)
                page = context.new_page()
                page_errors: list[str] = []
                page.on("pageerror", lambda _error: page_errors.append("page_error"))
                request = context.request

                stage = "runtime_preflight"
                runtime = _check_runtime(request, args.base_url)
                report["runtime"] = runtime
                report["checks"]["dev_profile_live_policy_and_credentials"] = True

                stage = "fixture_preflight"
                original_inspection = _post_inspect(request, args.base_url, args.fixture, args.sheet)
                product_sheet = _choose_product_sheet(original_inspection, args.sheet)
                if original_inspection.get("sheet_required"):
                    original_inspection = _post_inspect(request, args.base_url, args.fixture, product_sheet)
                bounded_fixture, fixture_info = _create_bounded_fixture(
                    args.fixture, product_sheet, original_inspection.get("headers") or [],
                    args.max_products, temp_root,
                )
                inspection = _post_inspect(request, args.base_url, bounded_fixture, product_sheet)
                _validate_fixture_content(bounded_fixture, inspection, product_sheet, args.max_products)
                candidate_sheets = [str(value) for value in inspection.get("candidate_sheets") or []]
                if any(name.strip().casefold() == "notes" for name in candidate_sheets):
                    raise E2EFailure("notes_candidate_not_excluded", "Notes 仍被识别为商品工作表候选")
                report["checks"]["bounded_url_fixture_with_notes_and_text_input"] = True
                report["checks"]["notes_sheet_excluded"] = True
                report["fixture"] = {"product_sheet": fixture_info["product_sheet"],
                                     "notes_sheet_preserved": True,
                                     "product_rows": args.max_products,
                                     "image_url_rows": args.max_products,
                                     "text_enrichment_rows": "at least 1"}

                stage = "open_dev_and_upload"
                page.goto(f"{args.base_url}/#/tagging", wait_until="domcontentloaded", timeout=60_000)
                page.locator("input[type=file]").first.set_input_files(str(bounded_fixture))
                page.get_by_text("核对表头与项目", exact=True).wait_for(state="visible", timeout=30_000)
                with page.expect_response(lambda response: response.url.endswith("/api/import")
                                          and response.request.method == "POST", timeout=60_000) as imported:
                    page.get_by_role("button", name="导入独立工作表", exact=True).click()
                import_result = _json_response(imported.value, "import")
                dataset_id = str(import_result.get("id") or "")
                if not dataset_id:
                    raise E2EFailure("import_id_missing", "导入成功响应没有数据集 ID")
                report["checks"]["workbook_imported"] = True
                report["dataset_id"] = dataset_id
                page.get_by_label("数据集", exact=True).wait_for(state="visible", timeout=30_000)
                page.get_by_label("数据集", exact=True).select_option(dataset_id)
                page.locator(".product-table tbody tr").nth(args.max_products - 1).wait_for(
                    state="visible", timeout=30_000)

                stage = "import_image_and_text_enrichment"
                import_evidence = _wait_import_enrichment(
                    request, args.base_url, dataset_id, args.max_products, args.timeout_seconds)
                if (import_evidence["status"] != "completed"
                        or import_evidence["image_failures"] > 0
                        or import_evidence["text_items_completed"] < import_evidence["text_items_expected"]):
                    stage = "retry_import_enrichment"
                    page.get_by_role("button", name="文字底稿", exact=True).click()
                    with page.expect_response(lambda response: response.url.endswith(
                            f"/api/drafts/{dataset_id}/text-records/extract")
                            and response.request.method == "POST", timeout=60_000) as retry_response:
                        page.get_by_role("button", name="提取／重试属性与图片", exact=True).click()
                    retry_task = _json_response(retry_response.value, "import_enrichment_retry")
                    retry_task_id = str(retry_task.get("id") or "")
                    if not retry_task_id:
                        raise E2EFailure("import_retry_id_missing", "导入整理重试没有返回任务 ID")
                    retry_state = _poll_state(
                        request, args.base_url, dataset_id, args.timeout_seconds,
                        lambda value: any(job.get("id") == retry_task_id
                                          and job.get("status") not in ACTIVE_STATUSES
                                          for job in value.get("jobs") or []),
                        "import_retry_timeout", "导入图片/文字整理重试超时")
                    import_evidence["retry"] = {"attempted": True, "task_id": retry_task_id}
                    import_evidence["retry_result"] = _verify_import_enrichment(retry_state, args.max_products)
                    import_evidence["status"] = import_evidence["retry_result"]["status"]
                    import_evidence["image_downloads"] = import_evidence["retry_result"]["image_downloads"]
                    import_evidence["image_failures"] = import_evidence["retry_result"]["image_failures"]
                    import_evidence["images_available"] = import_evidence["retry_result"]["images_available"]
                    import_evidence["text_items_completed"] = import_evidence["retry_result"]["text_items_completed"]
                    import_evidence["text_items_expected"] = import_evidence["retry_result"]["text_items_expected"]
                else:
                    import_evidence["retry"] = {"attempted": False}
                    import_evidence = {**import_evidence,
                        **_verify_import_enrichment(_get_json(request,
                            f"{args.base_url}/api/drafts/{dataset_id}", "dataset_state"), args.max_products)}
                report["import_enrichment"] = import_evidence
                report["checks"]["url_images_downloaded_and_text_enrichment_completed"] = True
                for row in page.locator(".product-table tbody tr").all()[:args.max_products]:
                    row.locator("td.fixed.image img").wait_for(state="visible", timeout=20_000)
                report["checks"]["product_images_visible_in_worksheet"] = True

                stage = "start_agent"
                panel = page.locator(".agent-panel")
                if panel.count() != 1:
                    raise E2EFailure("agent_panel_missing", "打标页没有唯一的 Agent 会话面板")
                panel = page.locator(".agent-panel")
                start = page.get_by_role("button", name="启动 Agent", exact=True)
                if start.count() == 0:
                    raise E2EFailure("agent_start_button_missing", "Agent 面板没有可用的启动按钮")
                start.click()
                page.get_by_label("Agent 输入", exact=True).wait_for(state="visible", timeout=60_000)
                expect_active = page.locator(".agent-status")
                deadline = time.monotonic() + 60
                while time.monotonic() < deadline and "本轮：idle" not in expect_active.inner_text():
                    page.wait_for_timeout(250)
                if "本轮：idle" not in expect_active.inner_text():
                    raise E2EFailure("agent_session_start_timeout", "Agent 会话没有进入可发送状态")
                session_active = True
                report["checks"]["agent_session_started"] = True

                stage = "send_authorized_prompt"
                bounded_prompt = (
                    args.prompt.strip()
                    + f"\n\n本次验收边界：仅处理当前工作表中的 {args.max_products} 件商品；"
                    "请提出一个 source=image 的图片标签和仅覆盖当前这些商品的执行提案。"
                    "不要调用 inspect_representative_images、view_product_image、research_knowledge，"
                    "不要访问 Amazon、MCP 或外部资料；不要自行确认或执行提案，等待我在工作表中确认。"
                )
                page.get_by_label("Agent 输入", exact=True).fill(bounded_prompt)
                with page.expect_response(lambda response: "/api/agent/sessions/" in response.url
                                          and response.url.endswith("/messages")
                                          and response.request.method == "POST",
                                          timeout=args.timeout_seconds * 1000):
                    page.get_by_role("button", name="发送", exact=True).click()
                deadline = time.monotonic() + args.timeout_seconds
                while time.monotonic() < deadline:
                    if page.locator(".question-card.pending").count():
                        raise E2EFailure("agent_needs_human_answer", "Agent 提出了问题；请先人工回答后重新运行")
                    if "本轮：idle" in page.locator(".agent-status").inner_text():
                        break
                    page.wait_for_timeout(500)
                else:
                    raise E2EFailure("agent_turn_timeout", "Agent 对话超过等待时间")
                report["checks"]["authorized_chinese_prompt_sent"] = True
                state = _get_json(request, f"{args.base_url}/api/drafts/{dataset_id}", "dataset_state")
                proposal, label_name, proposal_scope = _visual_proposal(state, args.max_products)
                if proposal_scope and len(proposal_scope) > args.max_products:
                    raise E2EFailure("agent_proposal_exceeds_cap", "Agent 提案覆盖商品超过上限")
                report["label_name"] = label_name
                report["proposal_id"] = proposal.get("id")
                report["checks"]["agent_created_visual_proposal"] = True
                page.get_by_role("button", name="结束会话", exact=True).click()
                page.locator(".agent-status").wait_for(state="visible", timeout=10_000)
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline and "stopped" not in page.locator(".agent-status").inner_text():
                    page.wait_for_timeout(250)
                if "stopped" not in page.locator(".agent-status").inner_text():
                    raise E2EFailure("agent_stop_timeout", "Agent 对话结束请求没有完成")
                session_active = False

                stage = "human_confirm_proposal_in_worksheet"
                proposal_card = page.locator("section.proposal").filter(has_text=label_name).first
                proposal_card.wait_for(state="visible", timeout=30_000)
                scope_select = proposal_card.get_by_label("确认商品范围", exact=True)
                if scope_select.count():
                    scope_select.select_option("all")
                confirm_button = proposal_card.get_by_role("button", name="确认并执行", exact=True)
                if confirm_button.count() == 0:
                    raise E2EFailure("visual_proposal_not_confirmable", "图片标签提案没有确认并执行入口")
                confirm_button.click()
                report["checks"]["visual_proposal_confirmed_in_worksheet"] = True

                stage = "wait_image_label_job"
                job, job_evidence = _wait_label_job(
                    request, args.base_url, dataset_id, label_name, args.max_products, args.timeout_seconds)
                report["label_job"] = {"id": job.get("id"), **job_evidence}
                report["checks"]["image_label_job_completed_within_cap"] = True

                stage = "save_version"
                state_before_save = _get_json(request, f"{args.base_url}/api/drafts/{dataset_id}", "dataset_state")
                prior_version_count = len(state_before_save.get("versions") or [])
                page.get_by_role("button", name="保存版本", exact=True).click()
                state = _poll_state(request, args.base_url, dataset_id, args.timeout_seconds,
                                    lambda value: len(value.get("versions") or []) > prior_version_count,
                                    "version_save_timeout", "保存正式版本超时")
                version = max(state["versions"], key=lambda item: int(item.get("version") or 0))
                version_id = str(version.get("id") or "")
                if not version_id:
                    raise E2EFailure("saved_version_id_missing", "已保存版本缺少 ID")
                page.get_by_label("查看版本", exact=True).select_option(version_id)
                report["version"] = int(version.get("version") or 0)
                report["checks"]["immutable_version_saved"] = True

                stage = "export_version"
                page.get_by_role("button", name="导出", exact=True).click()
                with page.expect_download(timeout=60_000) as download_info:
                    page.get_by_role("button", name=re.compile(r"^导出正式版本$")).click()
                export_path = args.output_dir / f"agent-image-export-{run_id}.xlsx"
                download_info.value.save_as(str(export_path))
                if not export_path.is_file() or export_path.stat().st_size <= 0:
                    raise E2EFailure("export_file_missing", "浏览器没有生成有效的 Excel 导出文件")

                stage = "verify_export"
                workbook = load_workbook(export_path, data_only=True, read_only=True)
                try:
                    if "Product_Result" not in workbook.sheetnames or "AI_Detail" not in workbook.sheetnames:
                        raise E2EFailure("export_sheets_missing", "导出缺少 Product_Result 或 AI_Detail 工作表")
                    result_ws = workbook["Product_Result"]
                    detail_ws = workbook["AI_Detail"]
                    exported_headers = [str(cell.value or "") for cell in next(result_ws.iter_rows(min_row=1, max_row=1))]
                    if exported_headers[:len(inspection.get("headers") or [])] != inspection["headers"]:
                        raise E2EFailure("export_original_headers_mismatch", "导出没有保留 fixture 的原始商品字段")
                    if label_name not in exported_headers:
                        raise E2EFailure("export_label_column_missing", "导出结果缺少 Agent 提出的图片标签列")
                    if result_ws.max_row != args.max_products + 1:
                        raise E2EFailure("export_product_count_mismatch", "导出商品行数与 max-products 不一致")
                    detail_headers = [str(cell.value or "") for cell in next(detail_ws.iter_rows(min_row=1, max_row=1))]
                    detail_label_index = detail_headers.index("标签") if "标签" in detail_headers else -1
                    detail_labels = {str(row[detail_label_index] or "") for row in detail_ws.iter_rows(
                        min_row=2, values_only=True)} if detail_label_index >= 0 else set()
                    if label_name not in detail_labels:
                        raise E2EFailure("export_detail_label_missing", "AI_Detail 没有图片标签执行记录")
                finally:
                    workbook.close()
                digest = hashlib.sha256(export_path.read_bytes()).hexdigest()
                report["artifacts"] = {"export_filename": export_path.name, "export_sha256": digest,
                                       "export_bytes": export_path.stat().st_size,
                                       "product_result_rows": args.max_products,
                                       "label_column_present": True,
                                       "detail_record_present": True}
                report["checks"]["export_downloaded_and_verified"] = True
                report["browser_page_errors"] = len(page_errors)
                if external_browser_count:
                    raise E2EFailure("browser_external_request_attempted", "浏览器请求了被策略拦截的非本机地址")
                report["checks"]["browser_requests_local_only"] = True
                if page_errors:
                    raise E2EFailure("browser_page_error", "打标页浏览器运行期间出现前端错误")
                report["checks"]["browser_page_error_free"] = True
                report["status"] = "completed"
                session_active = False
                browser.close()
                browser = None
    except E2EFailure as error:
        report["failure"] = {"stage": stage, "code": error.code, "message": str(error)}
    except PlaywrightTimeoutError:
        report["failure"] = {"stage": stage, "code": "browser_timeout", "message": "浏览器步骤超时"}
    except Exception as error:
        # Never persist arbitrary exception text, request bodies, prompts, or provider responses.
        report["failure"] = {"stage": stage, "code": "unexpected_error",
                             "message": f"未处理的 {type(error).__name__}；请查看 DEV 日志"}
    finally:
        if session_active and page is not None:
            try:
                end_session = page.get_by_role("button", name="结束会话", exact=True)
                if end_session.count() and end_session.is_enabled():
                    end_session.click(timeout=3_000)
            except Exception:
                pass
        if browser is not None:
            try:
                browser.close()
            except Exception:
                pass
        report["finished_at"] = datetime.now(timezone.utc).isoformat()
        report_path = args.output_dir / f"agent-image-export-{run_id}.json"
        try:
            args.output_dir.mkdir(parents=True, exist_ok=True)
            temp_report = report_path.with_suffix(".json.tmp")
            temp_report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            temp_report.replace(report_path)
        except Exception:
            print("报告无法写入指定目录。", file=sys.stderr)
            report_write_failed = True
        print(json.dumps({"status": report["status"], "run_id": run_id,
                          "report": report_path.name,
                          "export": (report.get("artifacts") or {}).get("export_filename"),
                          "failure": report.get("failure")}, ensure_ascii=False))
    if report_write_failed:
        return 2
    return 0 if report.get("status") == "completed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
