"""Isolated real-model, real-import integration; no production tags are changed."""
import os,sys,json,time,threading,io,hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
os.environ['TAG_STORAGE']=str(ROOT/'storage'/'skill-acceptance'/('live-'+time.strftime('%Y%m%d-%H%M%S')))
from sqlalchemy import select
from backend.db import *
from backend import draft,batch,text_facts,importer,draft_export,conversation,tagging_rules
from fastapi.testclient import TestClient
from backend.app import app
import openpyxl

def report(x):print(json.dumps(x,ensure_ascii=False),flush=True)
def wait(jid):
    deadline=time.time()+600
    while jid in batch.active:
        if time.time()>deadline:raise RuntimeError('验收任务超时')
        time.sleep(.2)
    with Session() as s:j=obj(require(s,BatchJob,jid))
    report({'job':jid,'kind':j['spec'].get('kind','labels'),'status':j['status'],'metrics':j['metrics'],'errors':j['errors']})
    assert j['status']=='completed',j['errors'];return j
def add(did,d):
    p=draft.propose(did,'验收新增',[{'kind':'add','definition':d}]);draft.confirm(did,p['id'])
    return next(c for c in draft.sheet(did)['columns'] if c['key']==d['key'])
def execute(did,col):
    p=draft.propose(did,'验收执行',[{'kind':'execute','column_ids':[col['id']],'overwrite':True}]);j=draft.confirm(did,p['id'])[0];batch.launch(j['id']);return wait(j['id'])

if __name__=='__main__':
    result={'started_at':now(),'skill':{k:v for k,v in tagging_rules.freeze().items() if k!='content'}};requests=[];original_request=batch.request_json
    def traced(model,messages,metrics,lock):
        requests.append({'model':model,'system_sha256':hashlib.sha256(messages[0]['content'].encode()).hexdigest(),'has_skill_core':'商品标签工程' in messages[0]['content'],'has_judgment':'judgment' in messages[0]['content'],'input_rows':len(json.loads(messages[-1]['content']))})
        result=original_request(model,messages,metrics,lock)
        with open(STORAGE/'model-responses.jsonl','a',encoding='utf-8') as f:f.write(json.dumps({'inputs':json.loads(messages[-1]['content']),'response':result},ensure_ascii=False)+'\n')
        return result
    batch.request_json=traced
    try:
        with TestClient(app) as client:
            sample=ROOT/'tests'/'fixtures'/'sample_10_rows.xlsx'
            response=client.post('/api/import',files={'file':('Skill验收10行.xlsx',sample.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
            assert response.status_code==200,response.text
            imported=response.json();did=imported['id'];result['dataset_id']=did;report({'dataset_id':did,'storage':str(STORAGE)})
            result['first_extraction']=wait(imported['text_job_id'])
            records=text_facts.records(did);assert len(records)==10 and all(r['status']=='ready' for r in records)
            for record in records:text_facts.validate(record['content'],record['source'])
            col=add(did,{'key':'pieces','name':'套装件数','type':'FACT','source':'text','result_shape':'scalar_enum','definition':'根据SKU、详细信息与商品标题明确的套装件数判断。2pcs/2-piece/Two Piece为2件套，3pcs/Three Piece为3件套；年龄尺寸不是件数。来源冲突或未知则null。','options':['2件套','3件套'],'fields':['sku','parameters','title'],'judgment':{'synonyms':{'Two Piece':'2件套','2pcs':'2件套','Three Piece':'3件套'}}})
            result['first_tagging']=execute(did,col);result['repeat_tagging']=execute(did,col)
            assert result['first_tagging']['metrics']['vision_requests']==0 and result['first_tagging']['metrics']['text_record_hits']==10
            result['values']=[{'asin':r['anchor'],'value':r['cells'][col['id']]['value'],'status':r['cells'][col['id']]['status']} for r in draft.sheet(did)['rows']]
            result['v1']=draft.save_version(did);before=draft.sheet(did,result['v1']['id'])
            draft.correct(did,before['rows'][0]['id'],col['id'],'3件套');result['v2']=draft.save_version(did)
            assert draft.sheet(did,result['v1']['id'])==before
            path=draft_export.export(did,[col['id']],result['v1']['id']);book=openpyxl.load_workbook(path);ws=book['Product_Result']
            assert ws.max_column==71 and len(ws._images)==10 and ws.max_row==11
            numeric=0
            for r in before['rows']:
                for index,h in enumerate(before['dataset']['headers'],1):
                    value=r['raw'].get(h)
                    if isinstance(value,(int,float)) and not isinstance(value,bool):
                        c=ws.cell(before['rows'].index(r)+2,index);assert c.value==value and c.data_type=='n';numeric+=1
            result['export']={'path':str(path),'rows':10,'original_fields':70,'images':10,'numeric_cells_preserved':numeric}
            # Second upload proves cache reuse without inheriting row identity.
            second=client.post('/api/import',files={'file':('相同格式重复上传.xlsx',sample.read_bytes())}).json()
            result['repeat_upload']=wait(second['text_job_id']);assert result['repeat_upload']['metrics']['cache_hits']==10
            assert second['id']!=did
            # Clearly synthetic records exercise multiple positive and uncertain alternatives.
            wb=openpyxl.Workbook();sheet=wb.active;sheet.append(['ASIN','SKU','商品标题','详细信息'])
            sheet.append(['SYNTHETIC-A','TEST-A','Floral top and striped shorts; no polka dots','上衣有花卉，短裤有条纹，无波点'])
            sheet.append(['SYNTHETIC-B','TEST-B','Possibly floral or striped; unclear','图案无法确定，未提供其他证据'])
            sheet.append(['SYNTHETIC-C','FL-BK-2','Two Piece swimsuit','共3件；SKU缩写无字典'])
            buf=io.BytesIO();wb.save(buf)
            syn=client.post('/api/import',files={'file':('明确标识的合成验收.xlsx',buf.getvalue())}).json();result['synthetic_dataset_id']=syn['id'];result['synthetic_extraction']=wait(syn['text_job_id'])
            pattern=add(syn['id'],{'key':'pattern','name':'图案','type':'FACT','source':'text','result_shape':'multi_enum','cardinality':'multiple','definition':'识别文字明确出现的所有图案，多个同时存在返回完整命中集合和逐属性证据。可能A或B是未知，没提到不是否定。','options':['花卉','条纹','波点']})
            result['synthetic_tagging']=execute(syn['id'],pattern)
            synsheet=draft.sheet(syn['id']);assert set(synsheet['rows'][0]['cells'][pattern['id']]['value'])=={'花卉','条纹'}
            assert synsheet['rows'][1]['cells'][pattern['id']]['value'] is None
            prop=next(p for p in draft.state(syn['id'])['proposals'] if p['status']=='pending' and p['changes'][0]['kind']=='expand')
            draft.confirm(syn['id'],prop['id']);expanded=draft.sheet(syn['id']);children={c['definition']['attribute']:c for c in expanded['columns'] if c['definition'].get('parent_id')}
            assert expanded['rows'][0]['cells'][children['波点']['id']]['value'] is False
            result['synthetic_values']=expanded
            # Five messages, followed by a process restart on the same native session.
            prompts=[
                '这是隔离验收。请调用 read_tagging_rules(topic="output")，简短说明多属性三态的规则与Skill版本，不要创建提案。',
                '请读取 read_worksheet(offset=0,limit=1)，汇报这件商品文字底稿是否完成，并引用一条底稿原文。不修改数据。',
                '请用 read_product_source 读取刚才商品的 sku、parameters、title，报告三个字段是否有值。不要创建提案。',
                '请调用 read_tasks，报告文字拆解和套装件数任务的完成状态，不创建提案。',
                '请通过 propose_changes 提出新增派生列“件数组合”，key=piece_group,type=INTERPRETATION,source=text,method=mapping，depends_on用当前套装件数真实ID，mapping为2件套映射双件组合、3件套映射三件组合。只提案，不确认。'
            ];turns=[]
            for prompt in prompts:
                t=time.perf_counter();response=conversation.start(did,prompt)
                while did in conversation.busy:time.sleep(.2)
                with Session() as s:m=obj(require(s,Message,response['message_id']))
                cli=conversation.clients[did];item={'status':m['status'],'text':m['text'],'events':m['events'],'session_id':cli.session_id,'pid':cli.proc.pid,'seconds':time.perf_counter()-t};turns.append(item);report({'turn':len(turns),**item})
                assert m['status']=='completed',m['text']
            assert len({t['session_id'] for t in turns})==1 and len({t['pid'] for t in turns})==1
            pending=next(p for p in draft.state(did)['proposals'] if p['status']=='pending' and any(c.get('definition',{}).get('key')=='piece_group' for c in p['changes']))
            draft.confirm(did,pending['id']);assert any(c['key']=='piece_group' for c in draft.sheet(did)['columns'])
            conversation.close();r=conversation.start(did,'重启恢复验收。请调用 read_tagging_rules(topic="attributes")，报告Skill版本和单属性拆列是否需要确认。')
            while did in conversation.busy:time.sleep(.2)
            with Session() as s:m=obj(require(s,Message,r['message_id']))
            result['resume']={'status':m['status'],'session_id':conversation.clients[did].session_id,'text':m['text'],'events':m['events']}
            assert result['resume']['session_id']==turns[0]['session_id'] and m['status']=='completed';result['conversation']=turns
            result['status']='passed'
    except Exception as e:
        result['status']='failed';result['error']=repr(e);raise
    finally:
        result['requests']=requests;result['completed_at']=now();conversation.close()
        (STORAGE/'report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');report({'report':str(STORAGE/'report.json'),'status':result.get('status')})
