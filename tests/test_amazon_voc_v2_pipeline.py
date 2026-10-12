from __future__ import annotations

import copy
import json
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from backend.amazon_voc.runtime.voc_app.config import Settings
from backend.amazon_voc.runtime.voc_app.db import Database
from backend.amazon_voc.runtime.voc_app.full_voc_v2 import (
    FIELD_NAMESPACES,
    _report_aggregate_checks,
    adapt_result,
    build_dictionary_candidates,
    build_assignment_jobs,
    build_payload,
    build_synthesis_payload,
    compact_synthesis_request,
    minimum_formal_segment_reviews,
    normalize_mapping_paths,
    normalize_v2_result,
    repair_synthesis_contract,
    validate_assignment_batch,
    validate_dictionary_normalization,
    validate_result,
)
from backend.amazon_voc.runtime.voc_app.pipeline import Pipeline
from backend.amazon_voc.runtime.voc_app import pipeline as pipeline_module
from backend.amazon_voc.runtime.voc_app.storage import Storage
from backend.local_result_cache import LocalResultCache


def _workbook_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Review data"
    sheet.append(["ASIN", "Product Title", "Rating", "Review Title", "Review Text", "Review Date"])
    sheet.append(["B000000001", "Item 1", 5, "Fits", "Fits well for swimming.", "2026-01-01"])
    sheet.append(["B000000002", "Item 2", 2, "Broken", "The strap broke quickly.", "2026-01-02"])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _pair(zh: str, en: str) -> dict[str, str]:
    return {"zh": zh, "en": en}


def _v2_result(payload: dict) -> dict:
    reviews = payload["reviews"]
    first = reviews[0]
    event = {
        "event_id": "EV-001", "sample_id": first["sample_id"], "primary_for_review": True,
        "audience": {"user": "child"}, "scenario": {"context": "swimming"},
        "motivation": {"explicit": "comfortable fit", "inferred": "", "inference_basis": ""},
        "task": {"core_job": "swim comfortably", "desired_outcome": "comfortable fit"},
        "experience": {"object": "swimsuit", "result": "fits well", "pain": "", "itch": ""},
        "workaround": {"action": ""}, "evidence_quote": "Fits well for swimming.",
        "confidence": "high", "fact_inference_boundary": "The quote states the fit experience.",
    }
    tag_by_value = {}
    mappings = []
    for field_path, namespace in FIELD_NAMESPACES.items():
        value: object = event
        for part in field_path.split("."):
            value = value.get(part, "") if isinstance(value, dict) else ""
        for raw_value in (value if isinstance(value, list) else [value]):
            if not str(raw_value or "").strip():
                continue
            key = (namespace, str(raw_value))
            if key not in tag_by_value:
                tag_id = "T-%02d" % (len(tag_by_value) + 1)
                tag_by_value[key] = tag_id
            mappings.append({"event_id": event["event_id"], "field_path": field_path,
                             "raw_value": raw_value, "tag_id": tag_by_value[key],
                             "mapping_method": "model_normalized"})
    tags = [{"tag_id": tag_id, "namespace": namespace, "canonical_name": value,
             "definition": "Sample value from the review", "normalization_status": "normalized"}
            for (namespace, value), tag_id in tag_by_value.items()]

    segment = {
        "segment_id": "SEG-1", "status": "exploratory",
        "segment_key": {key: _pair(zh, en) for key, zh, en in (
            ("audience", "儿童使用", "Child user"), ("scenario", "游泳", "Swimming"),
            ("motivation", "舒适", "Comfort"), ("task", "舒适游泳", "Swim comfortably"))},
        "name": _pair("舒适游泳需求", "Comfortable swimming need"),
        "synthesis": _pair("照护者希望孩子穿着舒适并能顺利游泳。", "Caregivers want children to swim comfortably."),
        "user": _pair("儿童", "Child"), "buyer": _pair("证据不足", "Insufficient evidence"),
        "relationship": _pair("证据不足", "Insufficient evidence"),
        "physical_constraints": _pair("证据不足", "Insufficient evidence"),
        "role_needs": _pair("舒适穿着", "Comfortable fit"),
        "scenario": _pair("游泳", "Swimming"), "motivation": _pair("舒适", "Comfort"),
        "job": _pair("舒适游泳", "Swim comfortably"),
        "pain": _pair("证据不足", "Insufficient evidence"),
        "itch": _pair("证据不足", "Insufficient evidence"),
        "desired_outcome": _pair("穿着舒适", "Comfortable fit"),
        "caveat": _pair("样本较小，需复核。", "Small sample; verify further."),
        "merge_reason": _pair("相近任务合并。", "Similar tasks are combined."),
        "split_reason": _pair("当前不拆分。", "No split at this stage."),
        "decision_value": _pair("识别舒适度期待。", "Clarifies comfort expectations."),
        "confidence": "medium", "evidence": [{"event_id": "EV-001", "quote": "Fits well for swimming."}],
    }
    stage = {
        "stage_id": "J1", "name": _pair("穿着体验", "Wear experience"),
        "synthesis": _pair("评论反映穿着是否合身。", "Reviews describe whether the fit works."),
        "user_need": _pair("舒适合身", "Comfortable fit"),
        "friction_or_outcome": _pair("本条评论反馈合身。", "This review reports a good fit."),
        "confidence": "medium", "evidence": [{"event_id": "EV-001", "quote": "Fits well for swimming."}],
    }
    return {
        "overview": _pair("本样本主要提到舒适合身。", "The sample mainly mentions comfortable fit."),
        "method_note": _pair("对上传评论逐条覆盖并抽象需求。", "Every uploaded review is covered and needs are synthesized."),
        "review_processing": [
            {"sample_id": row["sample_id"],
             "status": "processed_with_events" if index == 0 else "no_need_signal",
             "reason": "The review describes a fit outcome." if index == 0 else "No VOC signal in fixture",
             "event_count": 1 if index == 0 else 0}
            for index, row in enumerate(reviews)
        ],
        "events": [event], "tag_dictionary": tags, "event_tag_mappings": mappings,
        "segments": [segment],
        "event_assignments": [{"event_id": "EV-001", "segment_id": "SEG-1",
                               "journey_stage_id": "J1", "confidence": "high", "reason": "The event describes fit."}],
        "review_assignments": [
            {"sample_id": row["sample_id"], "status": "processed_with_events" if index == 0 else "no_need_signal",
             "primary_event_id": "EV-001" if index == 0 else "",
             "primary_segment_id": "SEG-1" if index == 0 else "OTHER",
             "journey_stage_id": "J1" if index == 0 else "JOURNEY_UNCLEAR"}
            for index, row in enumerate(reviews)
        ],
        "segment_comparison": [_pair("当前只有一个探索细分。", "Only one exploratory segment is present.")],
        "journey": {"interpretation_type": "Cross-sectional synthesis of experience stages",
                    "summary": _pair("横截面阶段综合。", "Cross-sectional stage synthesis."),
                    "limitations": _pair("不能解释为同一用户的时间线。", "Not a single user's timeline."),
                    "stages": [stage]},
        "cross_cutting": [_pair("舒适合身是当前评论信号。", "Comfortable fit is the current review signal.")],
        "exploratory_signals": [_pair("该细分样本数有限，需进一步验证。", "This segment has limited sample support and needs verification.")],
        "limitations": [_pair("仅代表本测试工作簿。", "Represents only this test workbook.")],
    }


def _make_run(tmp_path: Path, *, analysis_mode: str | None):
    database = Database(tmp_path / "voc.sqlite3")
    storage = Storage(tmp_path / "voc-data")
    project = database.create_project("Pipeline test", "US", "US", "en", "Swimwear")
    workbook_path = tmp_path / "reviews.xlsx"
    workbook_path.write_bytes(_workbook_bytes())
    sources = {"reviews": str(workbook_path)}
    if analysis_mode is not None:
        sources["analysis_mode"] = analysis_mode
    run = database.create_run(project["id"], sources)
    storage.revision_dir(project["id"], run["id"], 1)
    settings = Settings(
        data_dir=tmp_path / "voc-data", database_path=tmp_path / "voc.sqlite3",
        model_base_url="https://model.invalid/v1", model_api_key="fake-key",
        model_protocol="chat", extractor_model="fake-voc-model", reviewer_model="fake-voc-model",
        concurrency=1,
    )
    return database, storage, settings, run


class FakeProvider:
    def __init__(self):
        self.calls: list[str] = []
        self.payloads = []

    def call(self, task, model, instructions, payload, schema):
        self.calls.append(task)
        self.payloads.append(payload)
        if task == "full_voc_v2_extract":
            fixture_payload = {"reviews": [{
                "sample_id": row["sample_id"], "title": row.get("review_title", ""),
                "text": row["review_text"],
            } for row in payload["reviews"]]}
            fixture = _v2_result(fixture_payload)
            events = []
            for row in fixture["events"]:
                source = next(item for item in payload["reviews"] if item["sample_id"] == row["sample_id"])
                if row["evidence_quote"].casefold() in source["review_text"].casefold():
                    events.append({key: value for key, value in row.items() if key != "event_id"})
            counts = {}
            for event in events:
                counts[event["sample_id"]] = counts.get(event["sample_id"], 0) + 1
            processing = [{
                "sample_id": row["sample_id"],
                "status": "processed_with_events" if counts.get(row["sample_id"]) else "no_need_signal",
                "reason": "The review describes a fit outcome." if counts.get(row["sample_id"])
                          else "No VOC signal in fixture",
                "event_count": counts.get(row["sample_id"], 0),
            } for row in payload["reviews"]]
            return {"review_processing": processing, "events": events}, {
                "latency_ms": 1, "input_tokens": 10, "output_tokens": 30}
        if task == "full_voc_v2_synthesis":
            events = payload["events"]
            reviews = list(dict.fromkeys(row["sample_id"] for row in events))
            fixture = _v2_result({"reviews": [{"sample_id": sample_id} for sample_id in reviews]})
            evidence_ids = [row["event_id"] for row in events[:5]]
            segment = dict(fixture["segments"][0])
            segment.pop("evidence", None)
            segment["evidence_event_ids"] = evidence_ids
            stage = dict(fixture["journey"]["stages"][0])
            stage.pop("evidence", None)
            stage["evidence_event_ids"] = evidence_ids
            return {
                "overview": fixture["overview"], "method_note": fixture["method_note"],
                "segments": [segment],
                "segment_comparison": fixture["segment_comparison"],
                "journey": {**fixture["journey"], "stages": [stage]},
                "cross_cutting": fixture["cross_cutting"],
                "exploratory_signals": fixture["exploratory_signals"],
                "limitations": fixture["limitations"],
            }, {
                "latency_ms": 1, "input_tokens": 10, "output_tokens": 30}
        if task == "full_voc_v2_assign":
            return {
                "event_assignments": [
                    {"event_id": row["event_id"], "segment_id": "SEG-1", "journey_stage_id": "J1",
                     "confidence": "high"}
                    for row in payload["events"]
                ]
            }, {"latency_ms": 1, "input_tokens": 10, "output_tokens": 10}
        if task == "extract_events":
            return {"reviews": [{"review_id": row["review_id"], "status": "无可用需求信号",
                                 "no_signal_reason": "fake legacy path", "events": []}
                                for row in payload["reviews"]]}, {"latency_ms": 1}
        raise AssertionError("unexpected fake-provider task: %s" % task)


class DictionaryNormalizationProvider(FakeProvider):
    def call(self, task, model, instructions, payload, schema):
        if task == "full_voc_v2_extract":
            self.calls.append(task)
            self.payloads.append(payload)
            events = []
            processing = []
            for review in payload["reviews"]:
                fixture = _v2_result({"reviews": [{"sample_id": review["sample_id"],
                                                   "text": review["review_text"]}]})
                template = fixture["events"][0]
                processing.append({"sample_id": review["sample_id"], "status": "processed_with_events",
                                   "reason": "The review describes a fit outcome.", "event_count": 2})
                for index, raw_value in enumerate(("Swimsuit", "swim suit")):
                    event = copy.deepcopy(template)
                    event["sample_id"] = review["sample_id"]
                    event["primary_for_review"] = index == 0
                    event["experience"]["object"] = raw_value
                    event["evidence_quote"] = review["review_text"]
                    event.pop("event_id", None)
                    events.append(event)
            return {"review_processing": processing, "events": events}, {
                "latency_ms": 1, "input_tokens": 30, "output_tokens": 50}
        if task == "full_voc_v2_dictionary":
            self.calls.append(task)
            self.payloads.append(payload)
            object_ids = [row["candidate_id"] for row in payload["candidates"]
                          if row["namespace"] == "object"]
            return {"groups": [{
                "namespace": "object", "candidate_ids": object_ids,
                "canonical_name": "泳衣", "canonical_name_en": "Swimsuit",
                "definition": "同义的泳衣对象表达。", "definition_en": "Equivalent swimsuit object terms.",
            }]}, {"latency_ms": 1, "input_tokens": 20, "output_tokens": 20}
        return super().call(task, model, instructions, payload, schema)


class UnknownMappingPathProvider(FakeProvider):
    def call(self, task, model, instructions, payload, schema):
        result, usage = super().call(task, model, instructions, payload, schema)
        if task == "full_voc_v2_assign":
            result["event_assignments"][0]["segment_id"] = "MISSING-SEGMENT"
        return result, usage


class MissingReviewEventCountProvider(FakeProvider):
    def call(self, task, model, instructions, payload, schema):
        result, usage = super().call(task, model, instructions, payload, schema)
        if task == "full_voc_v2_extract":
            result["review_processing"][0].pop("event_count")
        return result, usage


class SchemaEchoThenValidProvider(FakeProvider):
    def __init__(self):
        super().__init__()
        self.echoed_schema = False
        self.instructions_seen = []

    def call(self, task, model, instructions, payload, schema):
        self.instructions_seen.append(instructions)
        if task == "full_voc_v2_extract" and not self.echoed_schema:
            self.echoed_schema = True
            self.calls.append(task)
            self.payloads.append(payload)
            return dict(schema), {"latency_ms": 1, "input_tokens": 10, "output_tokens": 30}
        return super().call(task, model, instructions, payload, schema)


class AssignmentCoverageRepairProvider(FakeProvider):
    def __init__(self):
        super().__init__()
        self.invalidated_assignment = False

    def call(self, task, model, instructions, payload, schema):
        result, usage = super().call(task, model, instructions, payload, schema)
        if task == "full_voc_v2_assign" and not self.invalidated_assignment:
            self.invalidated_assignment = True
            result["event_assignments"].pop()
        return result, usage


def _isolate_local_side_effects(monkeypatch, tmp_path):
    monkeypatch.setattr(pipeline_module, "STAGE_RESULT_CACHE",
                        LocalResultCache("v2-pipeline-test", root=tmp_path / "stage-cache"))
    monkeypatch.setattr(pipeline_module, "record_usage", lambda *_args, **_kwargs: None)


def _sample_payload():
    return {"reviews": [
        {"sample_id": "R-001", "title": "Fits", "text": "Fits well for swimming."},
        {"sample_id": "R-002", "title": "Broken", "text": "The strap broke quickly."},
    ]}


def _adapted_two_event_report():
    cleaned_rows = [
        {"review_id": "R-001", "asin": "B000000001", "review_title": "Fits",
         "content": "Fits well for swimming."},
        {"review_id": "R-002", "asin": "B000000002", "review_title": "Broken",
         "content": "The strap broke quickly."},
    ]
    payload = build_payload({}, cleaned_rows)
    result = _v2_result(payload)
    second_event = {**result["events"][0], "event_id": "EV-002", "primary_for_review": False}
    result["events"].append(second_event)
    result["event_tag_mappings"].extend(
        {**mapping, "event_id": "EV-002"}
        for mapping in list(result["event_tag_mappings"])
        if mapping["event_id"] == "EV-001"
    )
    next(row for row in result["review_processing"] if row["sample_id"] == "R-001")["event_count"] = 2
    result["event_assignments"].append({
        "event_id": "EV-002", "segment_id": "SEG-1", "journey_stage_id": "J1", "confidence": "high",
    })
    return cleaned_rows, result, adapt_result({}, cleaned_rows, result, "fake-model")


def test_compact_synthesis_request_keeps_atomic_evidence_and_omits_full_review_and_inference_boundary():
    event = _v2_result(_sample_payload())["events"][0]
    payload = {
        "sample_count": 2,
        "task_settings": {"marketplace": "US"},
        "reviews": [{"sample_id": "R-001", "asin": "B000000001"},
                    {"sample_id": "R-002", "asin": "B000000002"}],
        "events": [event],
    }

    request = compact_synthesis_request(payload)

    assert request["sample_count"] == 2
    assert request["asin_count"] == 2
    assert request["asins"] == ["B000000001", "B000000002"]
    assert request["events"][0]["event_id"] == "EV-001"
    assert request["events"][0]["experience"]["result"] == "fits well"
    assert request["events"][0]["evidence_quote"] == "Fits well for swimming."
    assert "fact_inference_boundary" not in json.dumps(request)
    assert "review_text" not in json.dumps(request)

    assignment_job = build_assignment_jobs(payload, {"segments": [], "journey": {"stages": []}})[0][1]
    assignment_event = assignment_job["events"][0]
    assert assignment_event["event_id"] == "EV-001"
    assert assignment_event["sample_id"] == "R-001"
    assert assignment_event["experience"]["result"] == "fits well"
    assert "evidence_quote" not in assignment_event


def test_report_aggregation_uses_unique_membership_metrics_and_detects_corruption():
    cleaned_rows, result, report = _adapted_two_event_report()
    analysis = report["analysis"]
    segment = analysis["segments"][0]
    assert segment["member_event_ids"] == ["EV-001", "EV-002"]
    assert segment["metrics"] == {
        "review_count": 1, "review_coverage": 0.5, "event_count": 2,
        "asin_count": 1, "primary_review_count": 1,
    }
    auto_gate = next(gate for gate in analysis["quality"]["gates"] if gate["gate_id"] == "V2-AUTO")
    assert auto_gate["pass"] is True
    assert auto_gate["failed_count"] == 0

    events = report["events_doc"]["events"]
    assignments = result["event_assignments"]
    review_assignments = result["review_assignments"]
    baseline = _report_aggregate_checks(
        events, assignments, analysis["segments"], analysis["journey"], review_assignments, len(cleaned_rows),
    )
    assert all(check["pass"] for check in baseline)

    broken_segments = copy.deepcopy(analysis["segments"])
    broken_journey = copy.deepcopy(analysis["journey"])
    broken_segments[0]["member_event_ids"].remove("EV-002")
    broken_segments[0]["metrics"]["asin_count"] = 99
    broken_journey["cells"][0]["event_count"] += 1
    failed = _report_aggregate_checks(
        events, assignments, broken_segments, broken_journey, review_assignments, len(cleaned_rows),
    )
    failed_by_prefix = {check["name"]: check["pass"] for check in failed}
    assert failed_by_prefix["事件投影与细分成员逐条一致（OTHER 不进入细分成员）"] is False
    assert failed_by_prefix["细分 review/event/ASIN 指标等于唯一成员计数"] is False
    assert failed_by_prefix["旅程单元格与成员、阶段归属及唯一计数一致"] is False


def test_report_evidence_validation_allows_whitespace_normalization_only():
    payload = _sample_payload()
    payload["reviews"][0]["text"] = "Fits well for  swimming."
    result = _v2_result(payload)

    assert validate_result(payload, result) == []


def test_assignment_jobs_use_bounded_event_batches():
    events = [{"event_id": "EV-%03d" % index, "sample_id": "R-%03d" % index}
              for index in range(105)]
    payload = {"sample_count": 105, "events": events, "reviews": [], "task_settings": {}}
    synthesis = {
        "segments": [{"segment_id": "SEG-1", "name": {"zh": "测试", "en": "Test"}}],
        "journey": {"stages": []},
    }

    jobs = build_assignment_jobs(payload, synthesis, batch_size=50)

    assert [len(job[1]["events"]) for job in jobs] == [50, 50, 5]
    assert [job[0] for job in jobs] == ["assign-0001", "assign-0002", "assign-0003"]


def test_unknown_mapping_path_repairs_only_when_one_exact_event_value_and_namespace_match():
    payload = _sample_payload()
    result = _v2_result(payload)
    row = next(item for item in result["event_tag_mappings"] if item["field_path"] == "audience.user")
    original_row = dict(row)
    row["field_path"] = "user.audience"

    normalized = normalize_mapping_paths(result)

    repaired = next(item for item in normalized["event_tag_mappings"]
                    if item["raw_value"] == original_row["raw_value"]
                    and item["tag_id"] == original_row["tag_id"])
    assert repaired["field_path"] == "audience.user"
    assert {key: value for key, value in repaired.items() if key != "field_path"} == {
        key: value for key, value in original_row.items() if key != "field_path"
    }
    assert row["field_path"] == "user.audience"  # The model result is never mutated in place.
    assert validate_result(payload, normalized) == []


def test_unknown_mapping_path_with_two_same_namespace_candidates_fails_closed():
    payload = _sample_payload()
    result = _v2_result(payload)
    event = result["events"][0]
    event["motivation"]["inferred"] = event["motivation"]["explicit"]
    explicit = next(item for item in result["event_tag_mappings"]
                    if item["field_path"] == "motivation.explicit")
    duplicate = dict(explicit)
    explicit["field_path"] = "motivation_explicit"
    duplicate["field_path"] = "motivation_inferred"
    result["event_tag_mappings"].append(duplicate)
    original_paths = [item["field_path"] for item in result["event_tag_mappings"]]

    normalized = normalize_mapping_paths(result)

    assert [item["field_path"] for item in normalized["event_tag_mappings"]] == original_paths
    assert any("每个非空事件字段值" in error for error in validate_result(payload, normalized))
    assert [item["field_path"] for item in result["event_tag_mappings"]] == original_paths


def test_unknown_mapping_path_uses_tag_namespace_to_disambiguate_equal_values():
    payload = _sample_payload()
    result = _v2_result(payload)
    event = result["events"][0]
    event["scenario"]["context"] = event["audience"]["user"]
    mapping = next(item for item in result["event_tag_mappings"] if item["field_path"] == "scenario.context")
    mapping["raw_value"] = event["audience"]["user"]
    mapping["field_path"] = "context_scenario"

    normalized = normalize_mapping_paths(result)

    assert mapping["field_path"] == "context_scenario"
    assert next(item for item in normalized["event_tag_mappings"]
                if item["tag_id"] == mapping["tag_id"] and item["raw_value"] == mapping["raw_value"])["field_path"] == "scenario.context"
    assert validate_result(payload, normalized) == []


def test_mapping_repair_requires_exact_raw_value_including_whitespace():
    payload = _sample_payload()
    result = _v2_result(payload)
    event = result["events"][0]
    event["audience"]["user"] = " child "
    mapping = next(item for item in result["event_tag_mappings"] if item["field_path"] == "audience.user")
    mapping["field_path"] = "unknown_user_path"
    mapping["raw_value"] = "child"

    normalized = normalize_mapping_paths(result)

    assert next(item for item in normalized["event_tag_mappings"] if item["tag_id"] == mapping["tag_id"])["field_path"] == "unknown_user_path"
    assert validate_result(payload, normalized)


def test_missing_review_event_count_is_derived_without_overwriting_reported_counts():
    payload = _sample_payload()
    result = _v2_result(payload)
    result["review_processing"][0].pop("event_count")
    result["review_processing"][1]["event_count"] = 7

    normalized = normalize_v2_result(result)

    assert normalized["review_processing"][0]["event_count"] == 1
    assert normalized["review_processing"][1]["event_count"] == 7
    errors = validate_result(payload, normalized)
    assert any("事件数与事件表不一致" in error for error in errors)


def test_pipeline_rejects_missing_extraction_event_count(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = MissingReviewEventCountProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    with pytest.raises(ValueError, match="event_count"):
        pipeline.process(run["id"])

    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    assert provider.calls == ["full_voc_v2_extract"]
    assert database.get_run(run["id"])["status"] == "failed"
    assert not list((revision / "output").glob("*.html"))


def test_v2_extraction_retries_once_when_model_echoes_schema(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = SchemaEchoThenValidProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    pipeline.process(run["id"])

    assert provider.calls == ["full_voc_v2_extract", "full_voc_v2_extract", "full_voc_v2_synthesis", "full_voc_v2_assign"]
    assert provider.payloads[0] == provider.payloads[1]
    assert "不要复述、复制或输出 JSON Schema" in provider.instructions_seen[1]
    assert database.get_run(run["id"])["status"] == "completed"


def test_v2_extraction_does_not_retry_other_invalid_results(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = MissingReviewEventCountProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    with pytest.raises(ValueError, match="event_count"):
        pipeline.process(run["id"])

    assert provider.calls == ["full_voc_v2_extract"]
    assert database.get_run(run["id"])["status"] == "failed"


def test_v2_synthesis_stays_semantic_and_assignment_runs_in_separate_stage(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = FakeProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    pipeline.process(run["id"])

    assert provider.calls == ["full_voc_v2_extract", "full_voc_v2_synthesis", "full_voc_v2_assign"]
    assert provider.payloads[1]["events"][0]["evidence_quote"] == "Fits well for swimming."
    assignment_event = provider.payloads[2]["events"][0]
    assert assignment_event["event_id"] == provider.payloads[1]["events"][0]["event_id"]
    assert assignment_event["experience"]["result"] == "fits well"
    assert "evidence_quote" not in assignment_event
    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    result = json.loads((revision / "cache" / "full_voc_v2_result.json").read_text(encoding="utf-8"))
    assert result["event_assignments"][0]["segment_id"] == "SEG-1"
    assert result["event_assignments"][0]["journey_stage_id"] == "J1"
    assert any(row["mapping_method"] == "raw_passthrough" for row in result["event_tag_mappings"])
    assert any("未确认语义归并" in row["zh"] for row in result["limitations"])
    assert database.get_run(run["id"])["status"] == "completed"


def test_v2_bounded_dictionary_stage_merges_frequent_terms_and_keeps_raw_variants(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = DictionaryNormalizationProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    pipeline.process(run["id"])

    assert provider.calls == ["full_voc_v2_extract", "full_voc_v2_synthesis",
                              "full_voc_v2_dictionary", "full_voc_v2_assign"]
    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    analysis = json.loads((revision / "cache" / "analysis.json").read_text(encoding="utf-8"))
    dictionary = json.loads((revision / "cache" / "tag_dictionary.json").read_text(encoding="utf-8"))
    object_mappings = [row for row in dictionary["mappings"] if row["field_path"] == "experience.object"]
    assert {row["raw_value"] for row in object_mappings} == {"Swimsuit", "swim suit"}
    assert len({row["tag_id"] for row in object_mappings}) == 1
    assert analysis["dictionary_summary"]["model_merge_groups"] == 1
    assert analysis["dictionary_summary"]["model_merged_values"] == 2
    assert analysis["counts"]["normalized_value_coverage_pct"] > 0
    assert database.get_run(run["id"])["status"] == "completed"


def test_assignment_batch_contract_requires_exact_event_coverage_and_valid_destinations():
    first = _v2_result({"reviews": [{"sample_id": "R-001", "text": "Fits well for swimming."}]})["events"][0]
    second = {**first, "event_id": "EV-002", "sample_id": "R-002"}
    synthesis_payload = {"events": [first, second]}
    synthesis, _ = FakeProvider().call("full_voc_v2_synthesis", "fake", "", synthesis_payload, {})
    assignment_payload = {
        "events": [first, second], "segments": synthesis["segments"],
        "journey_stages": synthesis["journey"]["stages"],
    }
    assignment, _ = FakeProvider().call("full_voc_v2_assign", "fake", "", assignment_payload, {})

    assert "event_assignments" not in synthesis
    assert validate_assignment_batch(assignment_payload, assignment) == []
    assignment["event_assignments"].pop()
    errors = validate_assignment_batch(assignment_payload, assignment)
    assert any("恰好归属一次" in error for error in errors)


def test_full_voc_synthesis_drops_unknown_evidence_ids_and_audits_them():
    first = _v2_result({"reviews": [{"sample_id": "R-001", "text": "Fits well for swimming."}]})["events"][0]
    second = {**first, "event_id": "EV-002", "sample_id": "R-002"}
    payload = {
        "sample_count": 2,
        "task_settings": {},
        "reviews": [{"sample_id": "R-001"}, {"sample_id": "R-002"}],
        "events": [first, second],
    }
    semantic, _ = FakeProvider().call("full_voc_v2_synthesis", "fake", "", payload, {})
    assignment_payload = {
        "events": [first, second], "segments": semantic["segments"],
        "journey_stages": semantic["journey"]["stages"],
    }
    assignment, _ = FakeProvider().call("full_voc_v2_assign", "fake", "", assignment_payload, {})
    semantic["segments"][0]["evidence_event_ids"].append("UNKNOWN-EVENT")
    semantic["journey"]["stages"][0]["evidence_event_ids"].append("UNKNOWN-EVENT")
    semantic["event_assignments"] = assignment["event_assignments"]

    repaired, summary = repair_synthesis_contract(payload, semantic)

    assert "UNKNOWN-EVENT" not in {row["event_id"] for row in repaired["segments"][0]["evidence"]}
    assert "UNKNOWN-EVENT" not in {row["event_id"] for row in repaired["journey"]["stages"][0]["evidence"]}
    assert summary["stale_evidence_removed"] == 2


def test_formal_segment_threshold_uses_five_review_floor_and_one_percent_rate():
    assert minimum_formal_segment_reviews(0) == 5
    assert minimum_formal_segment_reviews(100) == 5
    assert minimum_formal_segment_reviews(501) == 6
    assert minimum_formal_segment_reviews(2000) == 20


def test_formal_segment_repair_demotes_below_threshold_and_keeps_supported_cluster():
    def run_with_support(member_count):
        events = [{"event_id": "E%d" % index, "sample_id": "R%d" % index}
                  for index in range(1, member_count + 1)]
        assignments = [{"event_id": row["event_id"], "segment_id": "SEG-1",
                        "journey_stage_id": "JOURNEY_UNCLEAR"} for row in events]
        segment = {"segment_id": "SEG-1", "status": "formal",
                   "evidence": [{"event_id": row["event_id"], "quote": "evidence"}
                                for row in events[:3]]}
        repaired, summary = repair_synthesis_contract(
            {"sample_count": 5, "events": events},
            {"events": events, "segments": [segment], "event_assignments": assignments,
             "journey": {"stages": []}},
        )
        return repaired["segments"][0]["status"], summary

    below_status, below_summary = run_with_support(4)
    at_status, at_summary = run_with_support(5)

    assert below_status == "exploratory"
    assert below_summary["formal_segments_demoted"] == 1
    assert at_status == "formal"
    assert at_summary["formal_segments_demoted"] == 0


def test_dictionary_candidate_selection_is_frequency_bounded_and_uses_unique_reviews():
    events = []
    for index, (sample_id, value) in enumerate([
        ("R1", "Swimsuit"), ("R2", "Swimsuit"), ("R3", "swim suit"), ("R4", "swim suit"),
        ("R5", "rare term"),
    ], 1):
        events.append({"event_id": "E%d" % index, "sample_id": sample_id,
                       "experience": {"object": value}})
    # A duplicate event from R1 must not inflate the independent-review count.
    events.append({"event_id": "E6", "sample_id": "R1", "experience": {"object": "Swimsuit"}})

    result = build_dictionary_candidates(events)

    by_value = {row["raw_value"]: row for row in result["candidates"]}
    assert by_value["Swimsuit"]["review_count"] == 2
    assert by_value["swim suit"]["review_count"] == 2
    assert "rare term" not in by_value
    assert result["candidate_count"] == 2
    assert result["candidate_coverage"] == 1.0


def test_dictionary_normalization_merges_only_valid_candidate_ids_and_preserves_raw_values():
    events = [
        {"event_id": "E1", "sample_id": "R1", "experience": {"object": "Swimsuit"}},
        {"event_id": "E2", "sample_id": "R2", "experience": {"object": "Swimsuit"}},
        {"event_id": "E3", "sample_id": "R3", "experience": {"object": "swim suit"}},
        {"event_id": "E4", "sample_id": "R4", "experience": {"object": "swim suit"}},
        {"event_id": "E5", "sample_id": "R5", "experience": {"object": "rare term"}},
    ]
    candidate_data = build_dictionary_candidates(events)
    payload = {"sample_count": 5, "events": events}
    assignments = [{"event_id": row["event_id"], "segment_id": "OTHER",
                    "journey_stage_id": "JOURNEY_UNCLEAR"} for row in events]
    ids = [row["candidate_id"] for row in candidate_data["candidates"]]
    dictionary_result = {"groups": [{
        "namespace": "object", "candidate_ids": ids,
        "canonical_name": "泳衣", "canonical_name_en": "Swimsuit",
        "definition": "同义的泳衣对象表达", "definition_en": "Equivalent terms for a swimsuit object",
    }]}

    assert validate_dictionary_normalization({"candidates": candidate_data["candidates"]}, dictionary_result) == []
    result, summary = repair_synthesis_contract(
        payload,
        {"events": events, "segments": [], "event_assignments": assignments, "journey": {"stages": []}},
        dictionary_normalization={"candidates": candidate_data["candidates"],
                                  "candidate_coverage": candidate_data["candidate_coverage"],
                                  "result": dictionary_result},
        dictionary_status="completed",
    )

    object_mappings = [row for row in result["value_tag_mappings"] if row["field_path"] == "experience.object"]
    assert {row["raw_value"] for row in object_mappings} == {"Swimsuit", "swim suit", "rare term"}
    assert len({row["tag_id"] for row in object_mappings if row["raw_value"] != "rare term"}) == 1
    assert next(row for row in object_mappings if row["raw_value"] == "rare term")["mapping_method"] == "raw_passthrough"
    merged = next(row for row in result["tag_dictionary"] if row["tag_id"] == object_mappings[0]["tag_id"])
    assert {row["raw_value"] for row in merged["raw_variants"]} == {"Swimsuit", "swim suit"}
    assert merged["independent_review_count"] == 4
    assert summary["dictionary_model_merged_values"] == 2


def test_surface_normalization_is_conservative_and_keeps_namespace_separate():
    events = [
        {"event_id": "E1", "sample_id": "R1", "experience": {"object": "Swimsuit"}},
        {"event_id": "E2", "sample_id": "R2", "experience": {"object": " swimsuit "}},
        {"event_id": "E3", "sample_id": "R3", "experience": {"pain": "Tight"}},
        {"event_id": "E4", "sample_id": "R4", "experience": {"pain": "tight"}},
    ]
    assignments = [{"event_id": row["event_id"], "segment_id": "OTHER",
                    "journey_stage_id": "JOURNEY_UNCLEAR"} for row in events]
    result, summary = repair_synthesis_contract(
        {"sample_count": 4, "events": events},
        {"events": events, "segments": [], "event_assignments": assignments, "journey": {"stages": []}},
        dictionary_status="not_needed",
    )

    objects = [row for row in result["value_tag_mappings"] if row["field_path"] == "experience.object"]
    pains = [row for row in result["value_tag_mappings"] if row["field_path"] == "experience.pain"]
    assert len({row["tag_id"] for row in objects}) == 1
    assert all(row["mapping_method"] == "surface_normalized" for row in objects)
    assert len({row["tag_id"] for row in pains}) == 1
    assert objects[0]["tag_id"] != pains[0]["tag_id"]
    assert summary["dictionary_surface_merged_values"] == 4


def test_v2_marker_runs_chunked_extraction_and_global_synthesis_report(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = FakeProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    pipeline.process(run["id"])

    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    raw = json.loads((revision / "cache" / "raw_reviews.json").read_text(encoding="utf-8"))
    cleaned = json.loads((revision / "cache" / "cleaned_reviews.json").read_text(encoding="utf-8"))
    analysis = json.loads((revision / "cache" / "analysis.json").read_text(encoding="utf-8"))
    assert len(raw["rows"]) == len(cleaned["rows"]) == 2
    assert provider.calls == ["full_voc_v2_extract", "full_voc_v2_synthesis", "full_voc_v2_assign"]
    assert len(provider.payloads[0]["reviews"]) == 2
    assert analysis["analysis_mode"] == "full_voc_v2"
    assert analysis["counts"]["events"] == 1
    assert analysis["quality"]["acceptance_ready"] is False
    html_path = next((revision / "output").glob("*.html"))
    xlsx_path = next((revision / "output").glob("*.xlsx"))
    html_text = html_path.read_text(encoding="utf-8")
    assert all("module%02d" % index in html_text for index in range(1, 7))
    assert 'class="v2-insight-grid"' in html_text
    assert "报告限制" in html_text and "舒适合身" in html_text
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "任务与方法说明", "产品主表", "清洗后评论", "评论处理覆盖", "语义事件", "动态标签词典",
        "事件标签映射", "正式与探索细分", "细分需求", "动态旅程", "证据索引", "质量门禁",
    ]
    assert workbook["清洗后评论"].max_row == 3
    assert workbook["语义事件"].max_row == 2
    assert "user_need" not in str(workbook["动态旅程"][1][0].value)
    workbook.close()
    assert database.get_run(run["id"])["status"] == "completed"


def test_pipeline_expands_unique_value_mappings_into_event_mappings_and_report(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = FakeProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    pipeline.process(run["id"])

    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    synthesis_batch = json.loads((revision / "cache" / "batches" / "full_voc_v2_synthesis" / "global.json").read_text(encoding="utf-8"))
    assignment_batch = json.loads((revision / "cache" / "batches" / "full_voc_v2_assign" / "assign-0001.json").read_text(encoding="utf-8"))
    accepted_result = json.loads((revision / "cache" / "full_voc_v2_result.json").read_text(encoding="utf-8"))
    assert "event_assignments" not in synthesis_batch
    assert assignment_batch["event_assignments"][0]["event_id"] == accepted_result["events"][0]["event_id"]
    assert provider.calls.count("full_voc_v2_synthesis") == 1
    assert provider.calls.count("full_voc_v2_assign") == 1
    assert accepted_result["event_tag_mappings"]
    assert len(accepted_result["review_assignments"]) == 2
    assert database.get_run(run["id"])["status"] == "completed"


def test_valid_stage_checkpoint_is_reused_after_downstream_run_failure(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = FakeProvider()
    pipeline = Pipeline(settings, database, storage, provider)
    pipeline.process(run["id"])
    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    cache = revision / "cache"
    cleaned = json.loads((cache / "cleaned_reviews.json").read_text(encoding="utf-8"))["rows"]
    master = json.loads((cache / "master.json").read_text(encoding="utf-8"))
    extraction = json.loads((cache / "full_voc_v2_extraction.json").read_text(encoding="utf-8"))
    semantic = json.loads((cache / "full_voc_v2_synthesis_result.json").read_text(encoding="utf-8"))
    payload = build_payload(master, cleaned)
    synthesis_payload = build_synthesis_payload(payload, extraction["events"])
    assignment_id, assignment_payload = build_assignment_jobs(
        synthesis_payload, semantic, batch_size=pipeline_module.ASSIGNMENT_EVENTS_PER_CALL)[0]
    database.execute("UPDATE batches SET status='failed' WHERE run_id=? AND revision=1 AND stage=? AND batch_id=?",
                     (run["id"], "full_voc_v2_assign", assignment_id))
    original_call_count = len(provider.calls)

    reused = pipeline._call(run, 1, revision, "full_voc_v2_assign", assignment_id, assignment_payload)

    assert len(reused["event_assignments"]) == len(assignment_payload["events"])
    assert len(provider.calls) == original_call_count
    assert database.one("SELECT status FROM batches WHERE run_id=? AND revision=1 AND stage=? AND batch_id=?",
                        (run["id"], "full_voc_v2_assign", assignment_id))["status"] == "completed"


def test_v2_multiple_review_chunks_get_one_global_synthesis_and_keep_full_coverage(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    settings = replace(settings, batch_reviews=1)
    provider = FakeProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    pipeline.process(run["id"])

    assert provider.calls == ["full_voc_v2_extract", "full_voc_v2_extract", "full_voc_v2_synthesis", "full_voc_v2_assign"]
    assert len(provider.payloads[0]["reviews"]) == len(provider.payloads[1]["reviews"]) == 1
    global_payload = provider.payloads[2]
    assert global_payload["sample_count"] == 2
    assert "review_text" not in json.dumps(global_payload, ensure_ascii=False)
    assert len(global_payload["events"]) == 1
    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    cleaned = json.loads((revision / "cache" / "cleaned_reviews.json").read_text(encoding="utf-8"))
    extraction = json.loads((revision / "cache" / "full_voc_v2_extraction.json").read_text(encoding="utf-8"))
    assert {row["sample_id"] for row in extraction["review_processing"]} == {
        row["review_id"] for row in cleaned["rows"]
    }
    assert database.get_run(run["id"])["status"] == "completed"


def test_pipeline_falls_back_for_invalid_assignment_and_preserves_mapping_report(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = UnknownMappingPathProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    pipeline.process(run["id"])

    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    assert provider.calls == ["full_voc_v2_extract", "full_voc_v2_synthesis", "full_voc_v2_assign"]
    result = json.loads((revision / "cache" / "full_voc_v2_result.json").read_text(encoding="utf-8"))
    analysis = json.loads((revision / "cache" / "analysis.json").read_text(encoding="utf-8"))
    assignment = result["event_assignments"][0]
    assert assignment["segment_id"] == "OTHER"
    assert assignment["journey_stage_id"] == "JOURNEY_UNCLEAR"
    assert assignment["confidence"] == "low"
    assert any(row["field_path"] == "audience.user" and row["mapping_method"] == "raw_passthrough"
               for row in result["event_tag_mappings"])
    assert analysis["segments"][0]["member_event_ids"] == []
    assert analysis["journey"]["unclear_event_count"] == 1
    assert all(cell["event_count"] == 0 for cell in analysis["journey"]["cells"])
    assert analysis["quality"]["all_automatic_gates_green"] is True
    assert database.get_run(run["id"])["status"] == "completed"
    assert list((revision / "output").glob("*.html"))


def test_incomplete_assignment_is_recovered_locally_without_repeating_global_synthesis(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode="full_voc_v2")
    provider = AssignmentCoverageRepairProvider()
    pipeline = Pipeline(settings, database, storage, provider)

    pipeline.process(run["id"])

    revision = storage.revision_dir(run["project_id"], run["id"], 1)
    result = json.loads((revision / "cache" / "full_voc_v2_result.json").read_text(encoding="utf-8"))
    extraction = json.loads((revision / "cache" / "full_voc_v2_extraction.json").read_text(encoding="utf-8"))
    assert provider.calls == ["full_voc_v2_extract", "full_voc_v2_synthesis", "full_voc_v2_assign"]
    assert provider.calls.count("full_voc_v2_synthesis") == 1
    recovered = result["event_assignments"][0]
    assert recovered["event_id"] == extraction["events"][0]["event_id"]
    assert recovered["segment_id"] == "OTHER"
    assert recovered["journey_stage_id"] == "JOURNEY_UNCLEAR"
    assert recovered["confidence"] == "low"
    repair_audit = json.loads((revision / "cache" / "v2_synthesis_repairs.json").read_text(encoding="utf-8"))
    assert repair_audit["assignment_batch_missing"] == 1
    assert (revision / "cache" / "full_voc_v2_extraction.json").exists()
    assert list((revision / "output").glob("*.html"))
    assert database.get_run(run["id"])["status"] == "completed"


def test_markerless_historical_run_keeps_legacy_pipeline(tmp_path, monkeypatch):
    _isolate_local_side_effects(monkeypatch, tmp_path)
    database, storage, settings, run = _make_run(tmp_path, analysis_mode=None)
    provider = FakeProvider()
    pipeline = Pipeline(settings, database, storage, provider)
    legacy_stages = []
    monkeypatch.setattr(pipeline, "_reconcile_extraction_coverage", lambda *_args: None)
    monkeypatch.setattr(pipeline, "_review_cache_complete", lambda *_args: True)
    monkeypatch.setattr(pipeline, "_withhold_unbounded_inferences", lambda *_args: None)
    for name in ("_normalize", "_segment", "_journey", "_compile"):
        monkeypatch.setattr(pipeline, name, lambda *_args, stage=name: legacy_stages.append(stage))
    monkeypatch.setattr(pipeline, "_export_calls", lambda *_args: None)

    pipeline.process(run["id"])

    assert provider.calls == ["extract_events"]
    assert legacy_stages == ["_normalize", "_segment", "_journey", "_compile"]
    assert database.get_run(run["id"])["status"] == "completed"
    assert "full_voc_v2" not in provider.calls
