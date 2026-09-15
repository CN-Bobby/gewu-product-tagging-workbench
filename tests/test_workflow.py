import hashlib
import io
import json
from pathlib import Path
import pytest
import openpyxl
from sqlalchemy import select
from fastapi.testclient import TestClient
from backend import importer, service
from backend.db import *
from fixture_data import FIXTURE_PATH

DEFS=[{'key':'color','name':'颜色','type':'FACT','definition':'仅提取标题或 SKU 中明确声明的颜色，不确定则留空','options':[],'source':'text'}]
def schema(ds):
    sc=service.propose(ds['id'],'流程测试',DEFS);service.approve_schema(sc['id']);return sc
def finish(run):
    for pid in run['product_ids']:
        service.record_tags(run['id'],pid,[{'key':'color','raw_value':'Blue','normalized_value':'蓝色','confidence':.9,'source':'text','evidence':'测试夹具：SKU Color: Blue；不是实际商品结论','asset_ids':[]}])
    with Session.begin() as s:require(s,Run,run['id']).status='completed'

def test_wps_images_row_binding_and_original_integrity(dataset):
    source=FIXTURE_PATH.read_bytes()
    assert hashlib.sha256(source).hexdigest()==dataset['sha256']
    assert Path(dataset['original_file']).read_bytes()==source
    ps=service.products(dataset['id'])
    assert len(ps)==10 and len(dataset['headers'])==70
    assert [p['row'] for p in ps]==list(range(2,12))
    assert ps[0]['anchor']=='SYNTHETIC-ASIN-001'
    assert ps[-1]['anchor']=='SYNTHETIC-ASIN-010'
    assert ps[0]['raw']['Prime价格($)'] is None
    assert ps[0]['fields']['sales']==1112
    assert all(len([a for a in p['images'] if a['source']=='excel'])==1 for p in ps)
    assert len({next(a['sha256'] for a in p['images'] if a['path']) for p in ps})==10
    assert all('DISPIMG' in p['raw']['图片'] for p in ps)

def test_dataset_and_schema_version_never_overwrite(dataset):
    d2=importer.import_file(FIXTURE_PATH.read_bytes(),'第二期.xlsx',workspace_id=dataset['workspace_id'])
    assert d2['version']==2 and dataset['version']==1
    a=schema(dataset);b=service.propose(dataset['id'],'第二个目标',[{**DEFS[0],'definition':'只分析标题的颜色声明，不用SKU'}])
    assert b['version']==2
    with Session() as s:assert require(s,Schema,a['id']).definitions==DEFS
    with pytest.raises(ValueError):service.create_run(d2['id'],a['id'],'sample','test')

def test_sample_gates_and_same_model(dataset):
    sc=service.propose(dataset['id'],'测试',DEFS)
    with pytest.raises(ValueError):service.create_run(dataset['id'],sc['id'],'sample','test')
    service.approve_schema(sc['id'])
    with pytest.raises(ValueError):service.create_run(dataset['id'],sc['id'],'full','test')
    r=service.create_run(dataset['id'],sc['id'],'sample','test');assert len(r['product_ids'])==10
    with pytest.raises(ValueError):service.approve_sample(r['id'])
    finish(r)
    with pytest.raises(ValueError):service.approve_sample(r['id'])
    ts=service.run_results(r['id'])
    for t in ts:service.review_tag(t['id'],{'color':'冷蓝'},'测试校准')
    service.approve_sample(r['id'])
    with pytest.raises(ValueError):service.create_run(dataset['id'],sc['id'],'full','other-model')
    full=service.create_run(dataset['id'],sc['id'],'full','test')
    assert len(service.run_results(full['id']))==10
    assert service.run_results(full['id'])[0]['corrected']=={'color':'冷蓝'}
    with pytest.raises(ValueError):service.review_tag(ts[0]['id'],{'color':'红色'})

def test_reject_invented_image_evidence(dataset):
    sc=schema(dataset);r=service.create_run(dataset['id'],sc['id'],'sample','test')
    p=service.products(dataset['id'])[0];asset=p['images'][0]['id']
    v={'key':'color','raw_value':'Blue','normalized_value':'蓝色','confidence':.9,'source':'image','evidence':'声称看图','asset_ids':[asset]}
    with pytest.raises(ValueError,match='尚未成功'):service.record_tags(r['id'],p['id'],[v])
    v['source']='unknown'
    with pytest.raises(ValueError,match='留空'):service.record_tags(r['id'],p['id'],[v])
    v.update(source='text',asset_ids=['does-not-exist'])
    with pytest.raises(ValueError):service.record_tags(r['id'],p['id'],[v])

def test_transform_chain_keeps_original_and_reference_only(dataset):
    sc=schema(dataset);r=service.create_run(dataset['id'],sc['id'],'sample','test');finish(r)
    first=service.propose_transform(r['id'],'color','family','蓝色归入蓝色系',{'蓝色':'蓝色系'})
    assert 'family' not in service.distribution(r['id'])
    service.approve_transform(first['id'])
    second=service.propose_transform(r['id'],'family','temperature','当前分析中蓝色系视为冷色',{'蓝色系':'冷色'})
    service.approve_transform(second['id'])
    assert service.distribution(r['id'])=={'color':{'蓝色':10},'family':{'蓝色系':10},'temperature':{'冷色':10}}
    with pytest.raises(ValueError):service.propose_transform(r['id'],'color','family','不允许覆盖旧维度',{'蓝色':'新值'})
    with Session() as s:assert all(e.solution['reference_only'] for e in s.scalars(select(Experience)))

def test_export_three_sheets_images_and_nulls(dataset):
    from backend.exporter import export_run
    sc=schema(dataset);r=service.create_run(dataset['id'],sc['id'],'sample','test')
    with pytest.raises(ValueError):export_run(r['id'])
    finish(r);target=export_run(r['id']);w=openpyxl.load_workbook(target)
    assert w.sheetnames==['Product_Result','AI_Detail','Tag_Dictionary']
    ws=w['Product_Result'];assert ws.max_row==11 and len(ws._images)==10
    assert [c.value for c in ws[1]][:70]==dataset['headers']
    assert ws.cell(2,dataset['headers'].index('Prime价格($)')+1).value is None
    assert ws.cell(2,dataset['headers'].index('月销量')+1).value==1112
    assert all(c.data_type!='f' for row in ws for c in row)
    assert w['AI_Detail'].max_row==11
    assert Path(dataset['original_file']).read_bytes()==FIXTURE_PATH.read_bytes()

def test_http_import_inspection_and_approval_surface(dataset):
    from backend.app import app
    with TestClient(app) as c:
        assert c.get('/api/health').status_code==200
        p=c.post('/api/import/inspect',files={'file':(FIXTURE_PATH.name,FIXTURE_PATH.read_bytes())})
        assert p.status_code==200 and p.json()['embedded_images']==10
        assert c.get('/api/datasets/'+dataset['id']).json()['dataset']['stage']=='DATA_READY'
        assert c.post('/api/datasets/'+dataset['id']+'/schemas',json={'goal':'test','definitions':DEFS},headers={'origin':'https://evil.example'}).status_code==403

def test_float_images_and_sparse_rows():
    from PIL import Image
    from openpyxl.drawing.image import Image as XImage
    w=openpyxl.Workbook();ws=w.active;ws.append(['ASIN','Title','Price']);ws.append(['A1','first',None]);ws.append([None,None,None]);ws.append(['A2','second',0])
    pic=io.BytesIO();Image.new('RGB',(10,10),'red').save(pic,format='PNG');pic.seek(0);ws.add_image(XImage(pic),'B4')
    buf=io.BytesIO();w.save(buf)
    parsed=importer.parse(buf.getvalue())
    assert [x['row'] for x in parsed['rows']]==[2,4]
    assert len(parsed['rows'][1]['images'])==1 and parsed['rows'][1]['fields']['price']==0

def test_retry_keeps_successes(dataset,monkeypatch):
    from backend import jobs
    sc=schema(dataset);r=service.create_run(dataset['id'],sc['id'],'sample','test')
    first=r['product_ids'][0]
    service.record_tags(r['id'],first,[{'key':'color','raw_value':None,'normalized_value':None,'confidence':None,'source':'unknown','evidence':'测试：缺少输入','asset_ids':[]}])
    called=[]
    class Fake:
        def __init__(self,ds,model,rid,pid,role=None,allow_business_questions=None):
            assert role=='worker'
            assert allow_business_questions is False   # a worker is not the question owner
            self.pid=pid;called.append(pid)
        def prompt(self,text):raise RuntimeError('controlled network failure')
        def close(self):pass
    monkeypatch.setattr(jobs,'ACP',Fake)
    jobs.active[r['id']]={'stop':False,'client':None};jobs._execute(r['id'])
    assert first not in called and len(called)==9
    assert len(service.run_results(r['id']))==1
    with Session() as s:assert require(s,Run,r['id']).status=='partial'
