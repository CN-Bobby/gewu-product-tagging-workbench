"""Phase 2 integration contract for the single-workbook Amazon VOC workbench page."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.amazon_voc import router as voc_router
from backend.amazon_voc import service as voc_service_module
from backend.amazon_voc.runtime.voc_app.pipeline import (
    Pipeline, build_journey_stage_proposal_input, build_segment_proposal_batches,
)
from backend.amazon_voc.service import AmazonVOCService


def review_workbook() -> bytes:
    from io import BytesIO
    book = Workbook()
    sheet = book.active
    sheet.title = "Reviews"
    sheet.append(["ASIN", "Title", "Rating", "Review", "Date"])
    for i in range(3):
        sheet.append(["B001TEST", "Fixture swim top", 2,
                      f"The support feels unstable and takes too long to adjust. Review {i}.",
                      "2026-09-01"])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class DeterministicVOCProvider:
    """Small schema-compatible fake; keeps integration tests independent of paid APIs."""
    def call(self, task, model, instructions, payload, schema):
        from backend.amazon_voc.runtime.voc_app.quality import iter_values

        def event_for(review):
            return {
                "review_id": review["review_id"], "evidence_quote": "support feels unstable",
                "audience": {"user": "明确使用者", "buyer": "本人", "relationship": "自购自用",
                             "physical_constraints": [], "skill_level": "普通使用者", "role_needs": []},
                "scenario": {"context": "日常使用", "place": "", "environmental_stressors": [],
                             "social_emotional_context": "", "trigger": "需要完成任务",
                             "journey_action": "执行核心动作"},
                "motivation": {"explicit": "完成任务", "inferred": "", "inference_basis": []},
                "task": {"core_job": "稳定完成任务", "functional_job": "稳定完成任务",
                         "emotional_job": "", "social_job": "", "desired_outcomes": ["结果稳定"]},
                "experience": {"object": "测试产品", "component": "核心部件", "attribute": "稳定性",
                               "result": "结果不稳定", "pain": ["容易失败"], "itch": ["希望更省时"],
                               "sentiment": "负面", "severity": "中"},
                "workaround": {"exists": False, "raw_action": "", "standardized_action": "",
                               "strategy": "", "recovery_goal": "", "costs": [], "effect": ""},
                "fact_inference": "事实", "confidence": "高",
            }

        if task == "extract_events":
            result = {"reviews": [{"review_id": row["review_id"], "status": "已生成事件",
                                   "no_signal_reason": "", "events": [event_for(row)]}
                                  for row in payload["reviews"]]}
        elif task == "review_events":
            result = {"decisions": [{"event_id": row["event_id"], "action": "accept",
                                     "reason": "证据匹配", "replacement": None}
                                    for row in payload["events"]]}
        elif task == "build_dictionary":
            grouped = {}
            for event in payload["events"]:
                for field, raw, namespace in iter_values(event):
                    grouped.setdefault((namespace, raw), []).append(event["event_id"])
            result = {"tags": [
                {"tag_id": f"TAG-{i:03d}", "namespace": namespace, "canonical_name": raw,
                 "definition": f"样本中的{namespace}表达", "parent_tag_id": "",
                 "raw_variants": [raw], "include_examples": [raw], "exclude_examples": [],
                 "support_event_ids": sorted(set(event_ids)), "confidence": "高"}
                for i, ((namespace, raw), event_ids) in enumerate(sorted(grouped.items()), 1)
            ]}
        elif task == "map_dictionary":
            lookup = {(item["namespace"], item["canonical_name"]): item["tag_id"]
                      for item in payload["tag_dictionary"]}
            result = {"mappings": [
                {"event_id": event["event_id"], "field_path": value["field_path"],
                 "raw_value": value["raw_value"],
                 "tag_id": lookup[(value["namespace"], value["raw_value"])]}
                for event in payload["events"] for value in event["values"]
            ]}
        elif task == "propose_segments":
            result = {"segments": [{"segment_id": "SEG-DAILY", "name": "日常稳定任务者",
                "segment_key": {"audience": "明确使用者", "scenario": "日常使用",
                                "motivation": "完成任务", "task": "稳定完成任务"},
                "persona_summary": "在日常使用中关注任务稳定完成。",
                "jtbd_statement": "明确使用者在日常使用时，为完成任务而稳定完成任务。",
                "merge_reason": "任务与成功标准一致。", "split_reason": "当前没有稳定的相反约束。",
                "need_priorities": []}]}
        elif task == "assign_segments":
            result = {"assignments": [{"event_id": event["event_id"], "segment_id": "SEG-DAILY",
                                       "reason": "四维一致"} for event in payload["events"]]}
        elif task == "build_journey":
            stages = [{"stage_id": "J-USE", "name": "执行任务", "order": 1,
                       "definition": "执行核心动作并获得结果"}]
            assignments = [] if not payload["fixed_stages"] else [
                {"event_id": event["event_id"], "stage_id": "J-USE",
                 "reason": "动作发生于执行阶段"} for event in payload["events"]]
            result = {"stages": stages, "assignments": assignments}
        else:
            raise AssertionError(f"unexpected deterministic model task: {task}")
        return result, {"latency_ms": 1, "input_tokens": 10, "output_tokens": 10}


@pytest.fixture
def voc(tmp_path, monkeypatch):
    instance = AmazonVOCService(root=tmp_path / "amazon-voc")
    monkeypatch.setattr(voc_router, "service", instance)
    monkeypatch.setattr(instance, "resolve_shared_model", lambda model_ref=None, credential_id=None: {
        "model_ref": model_ref or "shared/fake-text", "credential_id": credential_id or "cred-1",
        "base_url": "https://model.invalid/v1", "api_key": "test-secret-never-persist",
        "model": "fake-text",
    })
    monkeypatch.setattr(instance, "public_model_settings", lambda: {
        "configured": True, "default_model": "shared/fake-text", "default_credential_id": "cred-1",
        "providers": [{"name": "shared", "model_refs": ["shared/fake-text"],
                       "credentials": [{"id": "cred-1", "name": "Shared key",
                                        "enabled": True, "authorized": True}]}],
    })
    return instance


def client():
    from backend.app import app
    return TestClient(app)


def create_project(api):
    response = api.post("/api/amazon-voc/projects", json={
        "name": "Phase 2 fixture", "marketplace": "US", "country": "US",
        "language": "en", "category": "Swimwear",
    })
    assert response.status_code == 200, response.text
    return response.json()


def test_single_workbook_upload_uses_shared_text_model_and_isolates_storage(voc):
    api = client()
    settings = api.get("/api/amazon-voc/settings")
    assert settings.status_code == 200
    assert settings.json()["default_model"] == "shared/fake-text"
    assert "test-secret-never-persist" not in settings.text

    project = create_project(api)
    uploaded = api.post(f"/api/amazon-voc/projects/{project['id']}/runs",
        data={"model": "shared/fake-text", "credential_id": "cred-1"},
        files={"reviews": ("reviews.xlsx", review_workbook(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert uploaded.status_code == 200, uploaded.text
    run = uploaded.json()
    assert run["status"] == "uploaded"
    assert run["sources"]["reviews_sha256"]
    assert set(run["sources"]) == {
        "reviews", "reviews_sha256", "review_source_context"
    }
    assert run["sources"]["review_source_context"] is None
    snapshot = voc.db.one("SELECT * FROM workbench_model_snapshots WHERE run_id=?", (run["id"],))
    assert snapshot == {"run_id": run["id"], "model_ref": "shared/fake-text", "credential_id": "cred-1"}
    assert "test-secret-never-persist" not in voc.db.path.read_text(encoding="utf-8", errors="ignore")
    assert run["storage_scope"] == "amazon-voc"

    duplicate = api.post(f"/api/amazon-voc/projects/{project['id']}/runs",
        files=[("reviews", ("one.xlsx", review_workbook(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
               ("extra", ("two.xlsx", review_workbook(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))])
    assert duplicate.status_code == 422


def test_deterministic_pipeline_reports_artifacts_and_revision_rerun(voc):
    api = client()
    project = create_project(api)
    response = api.post(f"/api/amazon-voc/projects/{project['id']}/runs",
        files={"reviews": ("reviews.xlsx", review_workbook(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200, response.text
    run_id = response.json()["id"]

    voc.run_pipeline(run_id, provider=DeterministicVOCProvider())
    run = api.get(f"/api/amazon-voc/runs/{run_id}").json()
    assert run["status"] == "completed"
    assert run["model_ref"] == "shared/fake-text"
    assert run["model_calls"]
    html = api.get(f"/api/amazon-voc/runs/{run_id}/artifacts/html")
    excel = api.get(f"/api/amazon-voc/runs/{run_id}/artifacts/excel")
    assert html.status_code == excel.status_code == 200
    html_text = html.content.decode("utf-8")
    assert "01 · 数据与方法概览" in html_text
    assert html_text.count('class="module"') == 6
    assert "Alexa" not in html_text
    from io import BytesIO
    workbook = load_workbook(BytesIO(excel.content), read_only=True)
    assert len(workbook.sheetnames) == 12
    assert "外部补全" not in workbook.sheetnames
    assert "AI产品建议" not in workbook.sheetnames
    workbook.close()

    events = api.get(f"/api/amazon-voc/runs/{run_id}/events").json()["events"]
    revised = dict(events[0])
    revised["experience"] = dict(revised["experience"])
    revised["experience"]["pain"] = ["用户人工修订痛点"]
    patched = api.patch(f"/api/amazon-voc/runs/{run_id}/events/{revised['event_id']}",
                        json={"event": revised, "note": "phase-2 test"})
    assert patched.status_code == 200, patched.text
    assert patched.json()["revision"] == 2
    voc.run_pipeline(run_id, provider=DeterministicVOCProvider())
    run = api.get(f"/api/amazon-voc/runs/{run_id}").json()
    assert run["status"] == "completed"
    assert run["current_revision"] == 2
    assert "test-secret-never-persist" not in json.dumps(run, ensure_ascii=False)


def test_non_excel_input_is_rejected_and_secret_never_returned(voc):
    api = client()
    project = create_project(api)
    response = api.post(f"/api/amazon-voc/projects/{project['id']}/runs",
        files={"reviews": ("reviews.csv", b"review text", "text/csv")})
    assert response.status_code == 422
    assert "test-secret-never-persist" not in response.text


def test_cancel_retry_and_startup_recovery_keep_the_model_snapshot(voc):
    api = client()
    project = create_project(api)
    response = api.post(f"/api/amazon-voc/projects/{project['id']}/runs",
        files={"reviews": ("reviews.xlsx", review_workbook(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200, response.text
    run_id = response.json()["id"]
    cancelled = api.post(f"/api/amazon-voc/runs/{run_id}/cancel", json={})
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    retried = api.post(f"/api/amazon-voc/runs/{run_id}/retry", json={})
    assert retried.status_code == 200
    assert retried.json()["status"] == "uploaded"
    voc.db.update_run(run_id, status="extracting", worker_id="stale-worker")
    voc.recover()
    assert voc.db.get_run(run_id)["worker_id"] == ""
    assert voc.model_snapshot(run_id)["model_ref"] == "shared/fake-text"
    completed = voc.run_pipeline(run_id, provider=DeterministicVOCProvider())
    assert completed["status"] == "completed"


def test_pipeline_tuning_env_reaches_workbench_without_overriding_selected_model(voc, monkeypatch):
    tuning = {
        "VOC_BATCH_REVIEWS": "16",
        "VOC_BATCH_CHARS": "20000",
        "VOC_CONCURRENCY": "3",
        "VOC_MAX_REPAIRS": "0",
    }
    for name, value in tuning.items():
        monkeypatch.setenv(name, value)
    monkeypatch.setenv("VOC_MODEL_BASE_URL", "https://untrusted.invalid/v1")
    monkeypatch.setenv("VOC_MODEL_API_KEY", "env-secret-must-not-be-used")
    monkeypatch.setenv("VOC_MODEL_PROTOCOL", "responses")
    monkeypatch.setenv("VOC_EXTRACTOR_MODEL", "env-model-must-not-be-used")
    monkeypatch.setenv("VOC_REVIEWER_MODEL", "env-reviewer-must-not-be-used")

    project = voc.create_project({
        "name": "Runtime tuning fixture", "marketplace": "US", "country": "US",
        "language": "en", "category": "Swimwear",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    captured = {}

    class PipelineSpy:
        def __init__(self, settings, db, storage, provider):
            captured["settings"] = settings

        def process(self, run_id):
            captured["run_id"] = run_id

    monkeypatch.setattr(voc_service_module, "Pipeline", PipelineSpy)
    result = voc.run_pipeline(run["id"], provider=object())

    settings = captured["settings"]
    assert captured["run_id"] == run["id"]
    assert result["id"] == run["id"]
    assert (settings.batch_reviews, settings.batch_chars, settings.concurrency, settings.max_repairs) == (
        16, 20_000, 3, 0,
    )
    assert settings.data_dir == voc.root
    assert settings.database_path == voc.db.path
    assert settings.model_base_url == "https://model.invalid/v1"
    assert settings.model_api_key == "test-secret-never-persist"
    assert settings.model_protocol == "chat"
    assert settings.extractor_model == settings.reviewer_model == "fake-text"
    assert "env-secret-must-not-be-used" not in str(settings)


def test_persistently_unverifiable_review_events_are_excluded_after_two_rounds(voc):
    project = voc.create_project({
        "name": "Evidence exclusion fixture", "marketplace": "US", "country": "US",
        "language": "en", "category": "Swimwear",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)
    cache = workdir / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    source = "The real review says the fit feels secure during normal activity."
    invalid = {
        "event_id": "event-invalid-quote", "review_id": "review-1", "asin": "B001TEST",
        "rating": 5, "review_date": "2026-09-01", "review_text": source,
        "evidence_quote": "The product is surprisingly comfortable for every possible occasion.",
        "confidence": "低", "motivation": {"inferred": "", "inference_basis": []},
    }
    valid = {
        "event_id": "event-valid-quote", "review_id": "review-1", "asin": "B001TEST",
        "rating": 5, "review_date": "2026-09-01", "review_text": source,
        "evidence_quote": "fit feels secure", "confidence": "高",
        "motivation": {"inferred": "", "inference_basis": []},
    }
    (cache / "events.json").write_text(json.dumps({
        "schema": "amazon-voc-events/v4", "metadata": {}, "events": [invalid, valid],
    }, ensure_ascii=False), encoding="utf-8")
    (cache / "review_processing.json").write_text(json.dumps({
        "reviews": [{"review_id": "review-1", "status": "已生成事件", "event_count": 2, "reason": ""}],
    }, ensure_ascii=False), encoding="utf-8")
    (cache / "semantic_review.json").write_text(json.dumps({
        "selected_event_ids": [invalid["event_id"], valid["event_id"]],
        "selection_rule": "低置信/无依据推断/无效原文证据 + ASIN×评分10%分层抽样",
        "decisions": [
            {"event_id": invalid["event_id"], "action": "accept", "reason": "复核fixture", "replacement": None},
            {"event_id": valid["event_id"], "action": "accept", "reason": "复核fixture", "replacement": None},
        ],
        "unresolved_quote_event_ids": [invalid["event_id"]], "missing_event_ids": [], "complete": False,
    }, ensure_ascii=False), encoding="utf-8")
    (cache / "coverage_rounds.json").write_text(json.dumps({"review": 2}), encoding="utf-8")

    class NoCallProvider:
        def call(self, *args, **kwargs):
            raise AssertionError("already retried evidence must be excluded without another model call")

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
    )
    pipeline = Pipeline(settings, voc.db, voc.storage, NoCallProvider())
    pipeline._review(voc.db.get_run(run["id"]), 1, workdir)
    events = json.loads((cache / "events.json").read_text(encoding="utf-8"))["events"]
    review = json.loads((cache / "semantic_review.json").read_text(encoding="utf-8"))
    processing = json.loads((cache / "review_processing.json").read_text(encoding="utf-8"))["reviews"]
    assert [event["event_id"] for event in events] == [valid["event_id"]]
    assert review["complete"] is True
    assert review["unresolved_quote_event_ids"] == []
    assert review["dropped_unverifiable_event_ids"] == [invalid["event_id"]]
    assert review["dropped_unverifiable_event_count"] == 1
    assert processing[0]["event_count"] == 1


def test_extraction_resume_recovers_cached_events_and_repairs_only_missing_reviews(voc):
    project = voc.create_project({
        "name": "Extraction coverage recovery", "marketplace": "US", "country": "US",
        "language": "en", "category": "Swimwear",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)
    cache = workdir / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    source = "The support feels unstable and takes too long to adjust."
    cleaned = [{"review_id": "review-1", "asin": "B001TEST", "content": source, "rating": 2,
                "review_date": "2026-09-01"},
               {"review_id": "review-2", "asin": "B001TEST", "content": source, "rating": 2,
                "review_date": "2026-09-02"},
               {"review_id": "review-3", "asin": "B001TEST", "content": source, "rating": 2,
                "review_date": "2026-09-03"}]
    (cache / "cleaned_reviews.json").write_text(json.dumps({"rows": cleaned}), encoding="utf-8")
    (cache / "master.json").write_text(json.dumps({"marketplace": "US", "category": "Swimwear"}), encoding="utf-8")
    (cache / "events.json").write_text(json.dumps({"schema": "amazon-voc-events/v4", "metadata": {},
        "events": [{"event_id": "review-1-事01", "review_id": "review-1"}]}), encoding="utf-8")
    duplicate_no_signal = {"review_id": "review-2", "status": "无可用需求信号",
                           "reason": "明确无需求信号", "event_count": 0}
    (cache / "review_processing.json").write_text(json.dumps({"reviews": [
        duplicate_no_signal, dict(duplicate_no_signal)]}), encoding="utf-8")

    class ExplicitNoSignalProvider:
        calls = []
        def call(self, task, model, instructions, payload, schema):
            self.calls.append([row["review_id"] for row in payload["reviews"]])
            return {"reviews": [{"review_id": row["review_id"], "status": "无可用需求信号",
                                  "no_signal_reason": "评论未出现可识别的产品需求线索", "events": []}
                                 for row in payload["reviews"]]}, {"latency_ms": 1}

    provider = ExplicitNoSignalProvider()
    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
    )
    pipeline = Pipeline(settings, voc.db, voc.storage, provider)
    pipeline._reconcile_extraction_coverage(voc.db.get_run(run["id"]), 1, workdir)

    processing = json.loads((cache / "review_processing.json").read_text(encoding="utf-8"))["reviews"]
    assert {row["review_id"]: row["status"] for row in processing} == {
        "review-1": "已生成事件", "review-2": "无可用需求信号", "review-3": "无可用需求信号",
    }
    assert provider.calls == [["review-3"]]
    assert len(processing) == 3
    assert next(row for row in processing if row["review_id"] == "review-1")["event_count"] == 1


def test_raw_unmerged_tags_preserve_dictionary_coverage_without_claiming_normalization(voc):
    project = voc.create_project({
        "name": "Raw vocabulary fallback", "marketplace": "US", "country": "US",
        "language": "en", "category": "Pet Toys",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)
    cache = workdir / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    event = {"event_id": "event-1", "review_id": "review-1", "asin": "B001TEST", "rating": 5,
             "review_date": "2026-09-01", "review_text": "The squeaker is too loud but the toy is fun.",
             "evidence_quote": "squeaker is too loud", "confidence": "高",
             "audience": {"user": "dog owner"}, "scenario": {}, "motivation": {}, "task": {},
             "experience": {"object": "squeaky toy", "pain": ["too loud"]}, "workaround": {}}
    (cache / "events.json").write_text(json.dumps({"schema": "amazon-voc-events/v4",
        "metadata": {"model": "fake-text", "prompt_version": "v1", "generated_at": "2026-10-03"},
        "events": [event]}), encoding="utf-8")
    (cache / "master.json").write_text(json.dumps({"marketplace": "US", "category": "Pet Toys"}), encoding="utf-8")

    class PainOnlyDictionaryProvider:
        calls = []
        def call(self, task, model, instructions, payload, schema):
            self.calls.append(task)
            from backend.amazon_voc.runtime.voc_app.quality import iter_values
            if task == "build_dictionary":
                values = {}
                for item in payload["events"]:
                    for path, raw, namespace in iter_values(item):
                        if namespace == "pain":
                            values.setdefault(raw, []).append(item["event_id"])
                return {"tags": [{"tag_id": "TAG-PAIN", "namespace": "pain", "canonical_name": raw,
                    "definition": "模型归一痛点", "parent_tag_id": "", "raw_variants": [raw],
                    "include_examples": [raw], "exclude_examples": [],
                    "support_event_ids": sorted(set(event_ids)), "confidence": "高"}
                    for raw, event_ids in values.items()]}, {"latency_ms": 1}
            if task == "map_dictionary":
                lookup = {(tag["namespace"], tag["canonical_name"]): tag["tag_id"]
                          for tag in payload["tag_dictionary"]}
                return {"mappings": [{"event_id": item["event_id"], "field_path": value["field_path"],
                    "raw_value": value["raw_value"], "tag_id": lookup[(value["namespace"], value["raw_value"])]}
                    for item in payload["events"] for value in item["values"]
                    if (value["namespace"], value["raw_value"]) in lookup]}, {"latency_ms": 1}
            raise AssertionError(task)

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
        max_repairs=0,
    )
    provider = PainOnlyDictionaryProvider()
    pipeline = Pipeline(settings, voc.db, voc.storage, provider)
    pipeline._normalize(voc.db.get_run(run["id"]), 1, workdir)
    dictionary = json.loads((cache / "tag_dictionary.json").read_text(encoding="utf-8"))
    assert dictionary["metadata"]["mapping_missing_count"] == 0
    assert dictionary["metadata"]["normalized_tag_count"] == 1
    assert dictionary["metadata"]["raw_unmerged_tag_count"] > 0
    assert all(tag["normalization_status"] == "raw_unmerged" and "未执行同义词归并" in tag["definition"]
               for tag in dictionary["tags"] if tag["normalization_status"] == "raw_unmerged")
    assert any(row["mapping_method"] == "raw_passthrough" for row in dictionary["mappings"])
    assert dictionary["metadata"]["mapping_model_unmapped_count"] > 0
    provider_call_count = len(provider.calls)
    pipeline._normalize(voc.db.get_run(run["id"]), 1, workdir)
    assert len(provider.calls) == provider_call_count

    (cache / "cleaned_reviews.json").write_text(json.dumps({"rows": [{"review_id": "review-1",
        "asin": "B001TEST", "content": event["review_text"], "rating": 5, "review_date": "2026-09-01"}]}),
        encoding="utf-8")
    (cache / "master.json").write_text(json.dumps({"marketplace": "US", "category": "Pet Toys",
        "source_files": {"reviews": "reviews.xlsx"}}), encoding="utf-8")
    (cache / "review_processing.json").write_text(json.dumps({"reviews": [{"review_id": "review-1",
        "status": "已生成事件", "reason": "", "event_count": 1}]}), encoding="utf-8")
    (cache / "semantic_review.json").write_text(json.dumps({"selected_event_ids": ["event-1"],
        "decisions": [{"event_id": "event-1", "action": "accept"}], "complete": True,
        "missing_event_ids": [], "unresolved_quote_event_ids": []}), encoding="utf-8")
    (cache / "segment_model.json").write_text(json.dumps({"segments": [{"segment_key": {
        "audience": "dog owner", "scenario": "", "motivation": "", "task": ""},
        "member_event_ids": ["event-1"]}], "journey": {"event_assignments": [
        {"event_id": "event-1", "stage_id": "stage-1"}]}}), encoding="utf-8")
    from backend.amazon_voc.runtime.voc_app.quality import evaluate
    report = evaluate(workdir)
    dictionary_gate = next(gate for gate in report["gates"] if gate["gate_id"] == 5)
    assert dictionary_gate["pass"] is True
    assert report["counts"]["normalized_tags"] == 1
    assert report["counts"]["raw_unmerged_tags"] > 0


def test_unbounded_inference_is_withheld_from_reused_review_cache(voc):
    project = voc.create_project({
        "name": "Inference boundary", "marketplace": "US", "country": "US",
        "language": "en", "category": "Swimwear",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)
    cache = workdir / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    (cache / "events.json").write_text(json.dumps({"events": [{"event_id": "event-1",
        "motivation": {"explicit": "明确动机", "inferred": "无法核验的推断", "inference_basis": []}}]}),
        encoding="utf-8")
    (cache / "semantic_review.json").write_text(json.dumps({"complete": True, "decisions": []}), encoding="utf-8")

    withheld_count = Pipeline._withhold_unbounded_inferences(workdir)
    events = json.loads((cache / "events.json").read_text(encoding="utf-8"))["events"]
    review = json.loads((cache / "semantic_review.json").read_text(encoding="utf-8"))
    assert withheld_count == 1
    assert events[0]["motivation"] == {"explicit": "明确动机", "inferred": "", "inference_basis": []}
    assert review["unbounded_inference_withheld_count"] == 1


def test_journey_cache_deduplicates_repeated_model_assignments(voc):
    project = voc.create_project({
        "name": "Journey duplicate protection", "marketplace": "US", "country": "US",
        "language": "en", "category": "Swimwear",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)
    cache = workdir / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    (cache / "events.json").write_text(json.dumps({"events": [
        {"event_id": "event-1", "review_id": "review-1"},
    ]}), encoding="utf-8")
    (cache / "segment_model.json").write_text(json.dumps({"segments": [
        {"segment_id": "segment-1", "member_event_ids": ["event-1"]},
    ], "journey": {}}), encoding="utf-8")

    class DuplicateJourneyProvider:
        def call(self, task, model, instructions, payload, schema):
            if not payload["fixed_stages"]:
                self.stage_payload = payload
                return {"stages": [{"stage_id": "stage-1", "name": "使用", "order": 1,
                                    "definition": "用户实际使用产品"}], "assignments": []}, {"latency_ms": 1}
            duplicate = {"event_id": "event-1", "stage_id": "stage-1", "reason": "同一条阶段分配"}
            return {"stages": [], "assignments": [duplicate, dict(duplicate)]}, {"latency_ms": 1}

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
    )
    provider = DuplicateJourneyProvider()
    pipeline = Pipeline(settings, voc.db, voc.storage, provider)
    pipeline._journey(voc.db.get_run(run["id"]), 1, workdir)
    model = json.loads((cache / "segment_model.json").read_text(encoding="utf-8"))
    assert provider.stage_payload["event_count"] == 1
    assert "events" not in provider.stage_payload
    assert provider.stage_payload["segments"][0]["member_event_count"] == 1
    assert model["journey"]["event_assignments"] == [
        {"event_id": "event-1", "stage_id": "stage-1"},
    ]
    assert model["journey"]["stage_generation"]["sampled_event_count"] == 1


def test_journey_stage_proposal_is_bounded_and_stratified_without_raw_reviews():
    events = []
    segments = []
    for segment_index in range(3):
        member_ids = []
        for event_index in range(30):
            event_id = f"event-{segment_index}-{event_index:03d}"
            member_ids.append(event_id)
            events.append({
                "event_id": event_id,
                "review_id": f"review-{segment_index}-{event_index}",
                "asin": f"B00SECRET{segment_index}{event_index}",
                "review_date": f"2026-01-{event_index + 1:02d}",
                "review_text": "raw review content must not be included " * 100,
                "scenario": {"journey_action": f"动作 {segment_index}-{event_index}",
                             "context": "使用玩具"},
                "motivation": {"explicit": "安抚狗狗"},
                "task": {"core_job": "让狗狗玩耍"},
                "experience": {"pain": ["容易损坏"]},
                "workaround": {"standardized_action": "监督使用"},
                "evidence_quote": f"短证据 {segment_index}-{event_index}",
            })
        segments.append({
            "segment_id": f"SEG-{segment_index}", "name": f"细分{segment_index}",
            "persona_summary": "画像说明" * 100, "jtbd_statement": "任务说明" * 100,
            "member_event_ids": member_ids,
        })

    payload = build_journey_stage_proposal_input(events, segments)
    encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    assert len(encoded) <= 192_000
    assert payload["event_count"] == 90
    assert payload["sampling"]["sampled_event_count"] == 6
    assert payload["sampling"]["active_segment_count"] == 3
    assert [row["member_event_count"] for row in payload["segments"]] == [30, 30, 30]
    sample_ids = {signal["event_id"] for segment in payload["segments"]
                  for signal in segment["review_signals"]}
    assert sample_ids == {f"event-{segment_index}-{event_index:03d}"
                          for segment_index in range(3) for event_index in (0, 29)}
    assert "review_text" not in encoded.decode("utf-8")
    assert "B00SECRET" not in encoded.decode("utf-8")
    assert payload == build_journey_stage_proposal_input(events, segments)


def test_report_compiler_redacts_full_review_copied_into_derived_fields():
    import importlib.util

    script = Path(__file__).resolve().parents[1] / "backend" / "amazon_voc" / "runtime" / "scripts" / "compile_analysis.py"
    spec = importlib.util.spec_from_file_location("amazon_voc_compile_analysis_redaction_test", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    source = "The squeaker broke after one short play session."
    payload = {"segments": [{
        "pains": [{"label": source}, {"label": "A shorter safe summary."}],
        "need_priorities": [{"basis": "Verbatim: the squeaker broke\nafter one short play session."}],
        "evidence_snippets": [{"quote": source, "quote_omitted": False}],
    }]}

    redacted = module._redact_full_review_text(payload, [{"review_text": source}])

    assert redacted == 3
    assert source.casefold() not in json.dumps(payload, ensure_ascii=False).casefold()
    assert payload["segments"][0]["evidence_snippets"][0]["quote_omitted"] is True
    assert all("完整评论原文未嵌入报告" in value for value in (
        payload["segments"][0]["pains"][0]["label"],
        payload["segments"][0]["need_priorities"][0]["basis"],
        payload["segments"][0]["evidence_snippets"][0]["quote"],
    ))


def test_safe_provider_preserves_nonretryable_error_and_redacts_secret(voc):
    from backend.amazon_voc.runtime.voc_app.provider import ModelError

    project = voc.create_project({
        "name": "Safe provider error metadata", "marketplace": "US", "country": "US",
        "language": "en", "category": "Pet Toys",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)

    class PermanentErrorProvider:
        calls = 0
        def call(self, *args, **kwargs):
            self.calls += 1
            raise ModelError("bad request test-secret-never-persist", retryable=False, status_code=400)

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="test-secret-never-persist", model_protocol="chat", extractor_model="fake-text",
        reviewer_model="fake-text", max_repairs=2,
    )
    safe_provider = voc_service_module.SafeProvider(settings)
    inner = PermanentErrorProvider()
    safe_provider.inner = inner
    pipeline = Pipeline(settings, voc.db, voc.storage, safe_provider)

    with pytest.raises(ValueError) as error:
        pipeline._call(voc.db.get_run(run["id"]), 1, workdir, "build_journey", "journey-stages",
                       {"event_count": 0, "segments": [], "event_assignments": [], "fixed_stages": []})

    assert str(error.value).endswith("bad request [redacted]")
    assert inner.calls == 1
    calls = voc.db.all("SELECT status,retries,error FROM model_calls WHERE run_id=? AND stage=?",
                       (run["id"], "build_journey"))
    assert len(calls) == 1
    assert calls[0]["retries"] == 0
    assert "test-secret-never-persist" not in calls[0]["error"]
    with pytest.raises(ModelError) as wrapped:
        safe_provider.call("build_journey", "fake-text", "", {}, {})
    assert wrapped.value.retryable is False
    assert wrapped.value.status_code == 400


def test_oversized_dictionary_uses_explicit_exact_passthrough_without_model_calls(voc):
    project = voc.create_project({
        "name": "Oversized dictionary fallback", "marketplace": "US", "country": "US",
        "language": "en", "category": "Pet Toys",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)
    cache = workdir / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    event = {"event_id": "event-large", "review_id": "review-large", "audience": {"user": "x" * 210_000}}
    (cache / "events.json").write_text(json.dumps({"events": [event]}), encoding="utf-8")
    (cache / "master.json").write_text(json.dumps({"marketplace": "US", "category": "Pet Toys"}),
                                       encoding="utf-8")

    class NoCallProvider:
        calls = 0
        def call(self, *args, **kwargs):
            self.calls += 1
            raise AssertionError("oversized dictionary must not issue a model request")

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
        max_repairs=0,
    )
    provider = NoCallProvider()
    pipeline = Pipeline(settings, voc.db, voc.storage, provider)
    pipeline._normalize(voc.db.get_run(run["id"]), 1, workdir)
    dictionary = json.loads((cache / "tag_dictionary.json").read_text(encoding="utf-8"))
    assert provider.calls == 0
    assert dictionary["metadata"]["dictionary_generation_mode"] == "raw_passthrough_input_too_large"
    assert dictionary["metadata"]["dictionary_input_bytes"] > 192_000
    assert dictionary["metadata"]["mapping_missing_count"] == 0
    assert dictionary["metadata"]["raw_unmerged_tag_count"] == 1
    assert dictionary["mappings"][0]["mapping_method"] == "raw_passthrough"


def test_segment_proposals_batch_all_profiles_then_merge_duplicate_candidates(voc):
    project = voc.create_project({
        "name": "Segment profile batches", "marketplace": "US", "country": "US",
        "language": "en", "category": "Pet Toys",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)
    cache = workdir / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    events = []
    for index in range(30):
        events.append({"event_id": f"event-{index:03d}", "review_id": f"review-{index:03d}",
            "audience": {"user": f"dog owner {index} " + ("x" * 3_500)}, "scenario": {}, "motivation": {},
            "task": {}, "experience": {}, "workaround": {}, "evidence_quote": ""})
    (cache / "events.json").write_text(json.dumps({"events": events}), encoding="utf-8")
    (cache / "tag_dictionary.json").write_text(json.dumps({"tags": []}), encoding="utf-8")
    (cache / "cleaned_reviews.json").write_text(json.dumps({"rows": [{"review_id": e["review_id"]}
        for e in events]}), encoding="utf-8")
    (cache / "master.json").write_text(json.dumps({"analysis_config": {}}), encoding="utf-8")

    class ProposalProvider:
        def __init__(self):
            self.proposals = []
        def call(self, task, model, instructions, payload, schema):
            if task == "propose_segments":
                self.proposals.append(payload)
                return {"segments": [{"segment_id": "SEG-ALL", "name": "狗主人",
                    "segment_key": {"audience": "狗主人", "scenario": "日常", "motivation": "陪伴", "task": "玩耍"},
                    "persona_summary": "狗主人让狗狗玩耍。", "jtbd_statement": "陪狗狗玩耍。",
                    "merge_reason": "共同人群", "split_reason": "无稳定差异", "need_priorities": []},
                    {"segment_id": "SEG-INACTIVE", "name": "未命中细分", "segment_key": {
                        "audience": "新手犬主", "scenario": "训练", "motivation": "纠正行为", "task": "行为训练"},
                     "persona_summary": "没有获得样本事件的独立画像。", "jtbd_statement": "没有事件支持的任务。",
                     "merge_reason": "独立候选", "split_reason": "待验证", "need_priorities": []}]}, {"latency_ms": 1}
            if task == "assign_segments":
                return {"assignments": [{"event_id": event["event_id"], "segment_id": "SEG-ALL", "reason": "同一候选"}
                    for event in payload["events"]]}, {"latency_ms": 1}
            raise AssertionError(task)

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
        concurrency=1,
    )
    provider = ProposalProvider()
    pipeline = Pipeline(settings, voc.db, voc.storage, provider)
    pipeline._segment(voc.db.get_run(run["id"]), 1, workdir)
    assert len(provider.proposals) > 1
    assert sum(len(payload["profiles"]) for payload in provider.proposals) == len(events)
    assert sum(payload["coverage"]["event_count"] for payload in provider.proposals) == len(events)
    submitted = {profile["d"][0][0] for payload in provider.proposals for profile in payload["profiles"]}
    assert len(submitted) == len(events)
    model = json.loads((cache / "segment_model.json").read_text(encoding="utf-8"))
    assert len(model["segments"]) == 1
    assert len(model["segments"][0]["member_event_ids"]) == len(events)
    assert model["metadata"]["proposed_segment_count"] == 2
    assert model["metadata"]["inactive_candidate_count"] == 1
    assert model["metadata"]["proposal_mode"] == "profile_batches"
    assert model["metadata"]["proposal_batch_count"] == len(provider.proposals)


def test_nonretryable_model_error_stops_pipeline_attempts_immediately(voc):
    from backend.amazon_voc.runtime.voc_app.provider import ModelError

    project = voc.create_project({
        "name": "Permanent provider error", "marketplace": "US", "country": "US",
        "language": "en", "category": "Pet Toys",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)

    class PermanentErrorProvider:
        calls = 0
        def call(self, *args, **kwargs):
            self.calls += 1
            raise ModelError("模型接口 HTTP 400", retryable=False, status_code=400)

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
        max_repairs=2,
    )
    provider = PermanentErrorProvider()
    pipeline = Pipeline(settings, voc.db, voc.storage, provider)
    with pytest.raises(ValueError, match="模型接口 HTTP 400"):
        pipeline._call(voc.db.get_run(run["id"]), 1, workdir, "build_dictionary", "dictionary", {"events": []})
    assert provider.calls == 1


def test_dynamic_assignment_schema_repairs_unknown_segment_id_before_accepting(voc):
    from backend.amazon_voc.runtime.voc_app.schemas import SCHEMAS

    project = voc.create_project({
        "name": "Dynamic assignment ID validation", "marketplace": "US", "country": "US",
        "language": "en", "category": "Pet Toys",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)

    class InvalidThenValidProvider:
        def __init__(self):
            self.payloads = []
            self.schemas = []
        def call(self, task, model, instructions, payload, schema):
            self.payloads.append(payload)
            self.schemas.append(schema)
            chosen_id = "OUTSIDE-CANDIDATES" if len(self.payloads) == 1 else "SEG-VALID"
            return {"assignments": [{"event_id": "event-1", "segment_id": chosen_id, "reason": "匹配"}]}, {"latency_ms": 1}

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
        max_repairs=0,
    )
    provider = InvalidThenValidProvider()
    pipeline = Pipeline(settings, voc.db, voc.storage, provider)
    assignment_schema = json.loads(json.dumps(SCHEMAS["assign_segments"]))
    assignment_schema["properties"]["assignments"]["items"]["properties"]["segment_id"]["enum"] = ["SEG-VALID"]
    result = pipeline._call(voc.db.get_run(run["id"]), 1, workdir, "assign_segments", "strict-coverage",
                            {"events": [{"event_id": "event-1"}], "required_event_ids": ["event-1"]},
                            schema_override=assignment_schema)

    assert len(provider.payloads) == 2
    assert provider.schemas[0]["properties"]["assignments"]["items"]["properties"]["segment_id"]["enum"] == ["SEG-VALID"]
    assert "SEG-VALID" in provider.payloads[1]["repair_instruction"]
    assert result["assignments"][0]["segment_id"] == "SEG-VALID"


def test_segment_retry_reuses_base_batches_and_adds_only_a_new_bounded_repair(voc):
    project = voc.create_project({
        "name": "Stable assignment retry keys", "marketplace": "US", "country": "US",
        "language": "en", "category": "Pet Toys",
    })
    run = voc.create_run(project["id"], "reviews.xlsx", review_workbook())
    workdir = voc.storage.revision_dir(project["id"], run["id"], 1)
    cache = workdir / "cache"
    cache.mkdir(parents=True, exist_ok=True)
    event = {"event_id": "event-1", "review_id": "review-1", "audience": {"user": "dog owner"},
             "scenario": {}, "motivation": {}, "task": {}, "experience": {}, "workaround": {}, "evidence_quote": ""}
    (cache / "events.json").write_text(json.dumps({"events": [event]}), encoding="utf-8")
    (cache / "tag_dictionary.json").write_text(json.dumps({"tags": []}), encoding="utf-8")
    (cache / "cleaned_reviews.json").write_text(json.dumps({"rows": [{"review_id": "review-1"}]}),
                                                 encoding="utf-8")
    (cache / "master.json").write_text(json.dumps({"analysis_config": {}}), encoding="utf-8")

    class RetryCoverageProvider:
        def __init__(self):
            self.assignment_calls = 0
            self.recovery_mode = False
        def call(self, task, model, instructions, payload, schema):
            if task == "propose_segments":
                return {"segments": [{"segment_id": "SEG-VALID", "name": "狗主人", "segment_key": {
                    "audience": "狗主人", "scenario": "日常", "motivation": "陪伴", "task": "玩耍"},
                    "persona_summary": "狗主人带狗狗玩耍。", "jtbd_statement": "陪狗狗玩耍。",
                    "merge_reason": "人群一致", "split_reason": "无稳定分歧", "need_priorities": []}]}, {"latency_ms": 1}
            if task == "assign_segments":
                self.assignment_calls += 1
                event_ids = [row["event_id"] for row in payload.get("events", [])]
                if self.recovery_mode:
                    return {"assignments": [{"event_id": event_id, "segment_id": "SEG-VALID", "reason": "唯一候选"}
                                            for event_id in event_ids]}, {"latency_ms": 1}
                if self.assignment_calls == 1:
                    return {"assignments": [{"event_id": "event-1", "segment_id": "UNKNOWN", "reason": "错误候选"}]}, {"latency_ms": 1}
                return {"assignments": []}, {"latency_ms": 1}
            raise AssertionError(task)

    settings = voc_service_module.Settings(
        data_dir=voc.root, database_path=voc.db.path, model_base_url="https://model.invalid/v1",
        model_api_key="", model_protocol="chat", extractor_model="fake-text", reviewer_model="fake-text",
        concurrency=1, max_repairs=0,
    )
    provider = RetryCoverageProvider()
    pipeline = Pipeline(settings, voc.db, voc.storage, provider)
    with pytest.raises(ValueError, match="细分归属仍缺少 1 个事件"):
        pipeline._segment(voc.db.get_run(run["id"]), 1, workdir)
    first_attempt_calls = provider.assignment_calls
    assert first_attempt_calls == 4  # base + three bounded coverage attempts

    provider.recovery_mode = True
    pipeline._segment(voc.db.get_run(run["id"]), 1, workdir)
    assert provider.assignment_calls == first_attempt_calls + 1
    assignment_batches = voc.db.all(
        "SELECT batch_id,status FROM batches WHERE run_id=? AND revision=? AND stage='assign_segments' ORDER BY batch_id",
        (run["id"], 1))
    assert not any(row["batch_id"].startswith("assignment-r01-") for row in assignment_batches)
    segment_model = json.loads((cache / "segment_model.json").read_text(encoding="utf-8"))
    assert segment_model["segments"][0]["member_event_ids"] == ["event-1"]
