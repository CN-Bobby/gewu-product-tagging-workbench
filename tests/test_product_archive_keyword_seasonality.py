"""Read-only, archive-scoped keyword trend contract."""

from __future__ import annotations

from openpyxl import Workbook
import pytest

from backend.product_archive import keyword_seasonality


def report(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "03_24月搜索量_宽表"
    sheet.append(["关键词", "中文翻译", "类型", "2025-01", "2025-02", "2025-03", "2025-04"])
    sheet.append(["work pants for men", "男士工装裤", "核心", 100, 120, 180, None])
    sheet.append(["linen pants", "亚麻裤", "长尾", 80, -5, "unknown", 230])
    path = tmp_path / "keywords.xlsx"
    book.save(path)
    return path


def test_report_reads_keywords_in_original_order_and_preserves_missing_months(tmp_path):
    value = keyword_seasonality.read_report(report(tmp_path))
    assert value["status"] == "available"
    assert value["months"] == ["2025-01", "2025-02", "2025-03", "2025-04"]
    assert [item["keyword"] for item in value["keywords"]] == ["work pants for men", "linen pants"]
    assert value["keywords"][0]["values"] == [100, 120, 180, None]
    assert value["keywords"][1]["values"] == [80, None, None, 230]


def test_archive_trend_endpoint_only_accepts_completed_owned_report(tmp_path, monkeypatch):
    path = report(tmp_path)
    monkeypatch.setattr("backend.product_archive.service._read", lambda _archive: {
        "id": "archive-one", "workspace_id": "workspace-one", "marketplace": "US"})
    monkeypatch.setattr("backend.product_archive.overview_drafts.list_sources", lambda _archive: [
        {"id": "artifact:owned", "type": "keywords", "available": True,
         "name": "keyword-report.xlsx", "date": "2026-10-01"},
        {"id": "artifact:unfinished", "type": "keywords", "available": False},
    ])
    calls = []

    def artifact_path(workspace_id, artifact_id):
        calls.append((workspace_id, artifact_id))
        return path

    monkeypatch.setattr("backend.project_files.artifact_path", artifact_path)
    value = keyword_seasonality.for_archive("archive-one", "artifact:owned")
    assert value["status"] == "available" and value["source"]["id"] == "artifact:owned"
    assert calls == [("workspace-one", "owned")]
    for source in ("artifact:foreign", "artifact:unfinished"):
        with pytest.raises(ValueError, match="不属于当前档案"):
            keyword_seasonality.for_archive("archive-one", source)
    assert calls == [("workspace-one", "owned")]
    model_input = keyword_seasonality.from_selected_materials(
        {"workspace_id": "workspace-one"}, [{"source_id": "artifact:owned", "type": "keywords"}])
    assert model_input["months"] == value["months"]
    assert model_input["source_id"] == "artifact:owned"


def test_missing_trend_sheet_has_explained_empty_state(tmp_path):
    book = Workbook()
    path = tmp_path / "old-report.xlsx"
    book.save(path)
    value = keyword_seasonality.read_report(path)
    assert value == {"status": "trend_sheet_missing", "months": [], "keywords": []}


def test_model_brief_prefers_selected_report_with_more_trend_terms(tmp_path, monkeypatch):
    rich = report(tmp_path)
    book = Workbook()
    sheet = book.active
    sheet.title = "03_24月搜索量_宽表"
    sheet.append(["关键词", "2025-01", "2025-02", "2025-03"])
    sheet.append(["one term", 10, 20, 30])
    sparse = tmp_path / "sparse.xlsx"
    book.save(sparse)
    paths = {"sparse": sparse, "rich": rich}
    monkeypatch.setattr("backend.project_files.artifact_path",
                        lambda _workspace, artifact_id: paths[artifact_id])
    value = keyword_seasonality.from_selected_materials({"workspace_id": "workspace-one"}, [
        {"source_id": "artifact:sparse", "type": "keywords"},
        {"source_id": "artifact:rich", "type": "keywords"},
    ])
    assert value["source_id"] == "artifact:rich"
    assert len(value["keywords"]) == 2
