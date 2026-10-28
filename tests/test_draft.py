import copy
import threading
import pytest
import openpyxl
from sqlalchemy import select
from backend import draft,batch,draft_export
from backend.db import *

def add(did,key='pieces',source='text',**extras):
    d={'key':key,'name':{'pieces':'套装件数','color':'主色系','warm':'冷暖色'}.get(key,key),'type':'FACT','source':source,'definition':'只按可见证据判断商品标签，不确定返回未知','options':[],**extras}
    p=draft.propose(did,'新增',[{'kind':'add','definition':d}]);draft.confirm(did,p['id'])
    return next(c for c in draft.sheet(did)['columns'] if c['key']==key)

def job(did,col,ids=None,**kwargs):
    change={'kind':'execute','column_ids':[col['id']],**kwargs}
    if ids:change['product_ids']=ids
    p=draft.propose(did,'执行',[change]);return draft.confirm(did,p['id'])[0]

def run(j):
    batch.active[j['id']]=threading.Event();batch.execute(j['id'])
    with Session() as s:return obj(require(s,BatchJob,j['id']))

def test_continuous_versions_and_dependency(dataset):
    did=dataset['id'];c=add(did,'color');pid=draft.sheet(did)['rows'][0]['id']
    draft.correct(did,pid,c['id'],'红色')
    w=add(did,'warm',method='mapping',depends_on=[c['id']],mapping={'红色':'暖色','蓝色':'冷色'})
    v1=draft.save_version(did);frozen=copy.deepcopy(draft.sheet(did,v1['id']))
    draft.correct(did,pid,c['id'],'蓝色');v2=draft.save_version(did)
    assert draft.sheet(did)['rows'][0]['cells'][w['id']]['value']=='冷色'
    assert draft.sheet(did,v1['id'])==frozen
    assert len(draft.compare(did,v1['id'],v2['id']))==2

def test_manual_conflict_and_rule_revision(dataset):
    did=dataset['id'];c=add(did);p=draft.sheet(did)['rows'][0]['id'];j=job(did,c,[p])
    draft.correct(did,p,c['id'],'人工值')
    with draft.WRITE_LOCK,Session.begin() as s:
        draft.write_result(s,require(s,BatchJob,j['id']),p,c,{'normalized_value':'后台值','raw_value':'raw'})
    cell=draft.sheet(did)['rows'][0]['cells'][c['id']]
    assert cell['value']=='人工值' and cell['conflict']['normalized_value']=='后台值'
    with pytest.raises(ValueError):draft.save_version(did)

def test_proposal_group_deduplicates_merges_and_supersedes(dataset):
    did=dataset['id'];a=add(did,'a');b=add(did,'b');pids=[row['id'] for row in draft.sheet(did)['rows'][:2]]
    with Session.begin() as session:session.add(WorkbenchMeta(id='active-chat:'+did,dataset_id=did,kind='active_chat_turn',data={'turn_id':'turn-a'}))
    first=draft.propose(did,'第一列',[{'kind':'execute','column_ids':[a['id']],'product_ids':[pids[0]]}],use_active_turn=True)
    duplicate=draft.propose(did,'重复',[{'kind':'execute','column_ids':[a['id']],'product_ids':[pids[0]]}],use_active_turn=True)
    assert duplicate['id']==first['id'] and duplicate['merge_action']=='reused'
    combined=draft.propose(did,'第二列',[{'kind':'execute','column_ids':[b['id']],'product_ids':[pids[1]]}],use_active_turn=True)
    assert combined['group_id']==first['group_id'] and combined['group_revision']==2
    assert first['id']!=combined['id'] and len(combined['changes'][0]['column_ids'])==2
    with Session.begin() as session:require(session,WorkbenchMeta,'active-chat:'+did).data={'turn_id':'turn-b'}
    replacement=draft.propose(did,'修改第一列范围',[{'kind':'execute','column_ids':[a['id']],'product_ids':[pids[1]]}],use_active_turn=True)
    assert replacement['group_revision']==3
    targets={item['column_id']:item['product_ids'] for item in replacement['changes'][0]['targets']}
    assert targets=={a['id']:[pids[1]],b['id']:[pids[1]]}
    with pytest.raises(ValueError,match='已替换'):draft.confirm(did,combined['id'])

def test_confirmation_key_is_idempotent(dataset):
    did=dataset['id'];column=add(did,'idempotent');pid=draft.sheet(did)['rows'][0]['id']
    proposal=draft.propose(did,'执行',[{'kind':'execute','column_ids':[column['id']],'product_ids':[pid]}])
    first=draft.confirm(did,proposal['id'],confirmation_key='same-request')
    second=draft.confirm(did,proposal['id'],confirmation_key='same-request')
    assert [x['id'] for x in first]==[x['id'] for x in second]
    with Session() as session:assert len(list(session.scalars(select(BatchJob).where(BatchJob.proposal_id==proposal['id']))))==1

def test_batch_text_no_images_partial_split_retry(dataset,monkeypatch):
    import json
    did=dataset['id'];c=add(did);calls=[]
    def fake(model,messages,metrics,lock):
        items=json.loads(messages[-1]['content']);calls.append(len(items))
        if len(items)>2:raise ValueError('truncated')
        return {'results':[{'product_id':p['product_id'],'normalized_value':'2件套','raw_value':'2pcs','evidence':'标题'} for p in items]}
    monkeypatch.setattr(batch,'request_json',fake)
    monkeypatch.setattr(batch,'observation',lambda *a:pytest.fail('text called vision'))
    j=run(job(did,c));assert j['status']=='completed' and len(j['completed'])==10
    assert 5 in calls and 1 in calls
    j2=run(job(did,c));assert not j2['spec']['expected']

def test_export_preserves_70_fields_images_and_selected_order(dataset):
    did=dataset['id'];c=add(did);w=add(did,'color');data=draft.sheet(did)
    path=draft_export.export(did,[w['id'],c['id']]);wb=openpyxl.load_workbook(path)
    assert wb.sheetnames==['Product_Result','AI_Detail','Tag_Dictionary']
    ws=wb['Product_Result'];assert ws.max_column==72 and len(ws._images)==10
    assert list(next(ws.values))[:70]==data['dataset']['headers']
    assert ws.cell(2,data['dataset']['headers'].index('价格($)')+1).value==data['rows'][0]['raw']['价格($)']
    assert ws.cell(2,71).value=='【待处理】'
    assert wb['AI_Detail'].max_row==21 and wb['Tag_Dictionary'].max_row==3

def test_export_unknown_as_blank_keeps_unknown_details(dataset):
    did=dataset['id'];column=add(did,'unknown_blank');row=draft.sheet(did)['rows'][0]
    draft.correct(did,row['id'],column['id'],None)
    with Session.begin() as session:
        cell=draft.get_cell(session,row['id'],column['id'])
        cell.raw_value='unresolved raw observation'
        cell.evidence={'explanation':'test evidence remains in the detail sheet'}
    path=draft_export.export(did,[column['id']],unknown_as_blank=True)
    workbook=openpyxl.load_workbook(path,data_only=True)
    result=workbook['Product_Result']
    detail=workbook['AI_Detail']
    assert result.cell(2,len(dataset['headers'])+1).value is None
    assert detail.cell(2,5).value=='未知'
    assert detail.cell(2,6).value=='unresolved raw observation'
    assert 'test evidence remains in the detail sheet' in detail.cell(2,9).value


def test_project_dataset_version_endpoint_exports_data_x(dataset):
    from fastapi.testclient import TestClient
    from backend import project_files
    from backend.app import app
    did=dataset['id'];column=add(did,'project_unknown');row=draft.sheet(did)['rows'][0]
    draft.correct(did,row['id'],column['id'],None)
    with TestClient(app) as client:
        response=client.post(f'/api/drafts/{did}/versions',json={})
    assert response.status_code==200,response.text
    saved=response.json()
    assert saved['version']==1
    assert saved['project_artifact']['kind']=='data_x'
    artifact_path=project_files.project_dir(dataset['workspace_id'])/saved['project_artifact']['relative_path']
    workbook=openpyxl.load_workbook(artifact_path,data_only=True)
    result=workbook['Product_Result']
    detail=workbook['AI_Detail']
    assert result.cell(2,len(dataset['headers'])+1).value is None
    assert detail.cell(2,5).value=='未知'


def test_scope_approval_staleness_and_cross_dataset(dataset):
    did=dataset['id'];c=add(did);ids=[x['id'] for x in draft.sheet(did)['rows'][:3]]
    p=draft.propose(did,'执行',[{'kind':'execute','column_ids':[c['id']],'product_ids':ids}])
    j=draft.confirm(did,p['id'],scope=ids[:1],sample=True)[0]
    assert j['spec']['product_ids']==ids[:1] and j['spec']['sample']
    with pytest.raises(ValueError):draft.confirm(did,p['id'])
    p=draft.propose(did,'旧执行',[{'kind':'execute','column_ids':[c['id']]}])
    rev=draft.propose(did,'修改',[{'kind':'revise','column_id':c['id'],'definition':{**c['definition'],'definition':'新的规则只读取标题信息'}}]);draft.confirm(did,rev['id'])
    with pytest.raises(ValueError):draft.confirm(did,p['id'])

def test_pause_and_retry_only_incomplete(dataset,monkeypatch):
    import json
    did=dataset['id'];c=add(did);ids=[r['id'] for r in draft.sheet(did)['rows']];broken=ids[0]
    def fake(model,messages,metrics,lock):
        items=json.loads(messages[-1]['content'])
        return {'results':[{'product_id':p['product_id'],'normalized_value':'2件套','raw_value':'2pcs','evidence':'标题'} for p in items if p['product_id']!=broken]}
    monkeypatch.setattr(batch,'request_json',fake)
    j=run(job(did,c));assert len(j['completed'])==9 and j['status']=='partial'
    seen=[];broken=None
    original=fake
    def retry(model,messages,metrics,lock):seen.extend(json.loads(messages[-1]['content']));return original(model,messages,metrics,lock)
    monkeypatch.setattr(batch,'request_json',retry)
    j=run(j);assert j['status']=='completed' and len(seen)==1

def test_cycle_rejected(dataset):
    did=dataset['id'];c=add(did);w=add(did,'warm',method='mapping',depends_on=[c['id']],mapping={})
    with pytest.raises(ValueError):draft.propose(did,'循环',[{'kind':'revise','column_id':c['id'],'definition':{**c['definition'],'method':'mapping','depends_on':[w['id']],'mapping':{}}}])
