import io
import json
import threading
from pathlib import Path

import openpyxl
from sqlalchemy import select

from backend import batch, execution_contract, importer, text_facts
from backend.db import Asset, BatchJob, Dataset, Product, Session, obj, require


def workbook_bytes(*, extra_product_sheet=False):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = 'Product-US'
    ws.append(['ASIN', 'SKU', '详细参数', '商品主图', '商品标题'])
    ws.append(['B012345678', 'Color: Aqua | Size: 12',
               'Color: Blue | Fabric type: 100% Polyester',
               'https://images.example/item.jpg', 'Girls Rash Guard'])
    ws.append(['Export note: variant rows may affect export count'])
    if extra_product_sheet:
        second = book.create_sheet('Product-CA')
        second.append(['ASIN', 'SKU', '商品标题'])
        second.append(['B087654321', 'Color: Red', 'Girls Shirt'])
    notes = book.create_sheet('Notes')
    notes.append(['Exported from research tool'])
    notes.append(['Contact information omitted from test'])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_source_attribute_prompt_preserves_source_key_language_and_namespaces():
    assert 'Color统一为Color' in execution_contract.TEXT_PREFIX
    assert 'Size统一为Size' in execution_contract.TEXT_PREFIX
    assert '不得跨sku与parameters两个来源合并' in execution_contract.TEXT_PREFIX


def test_inspect_ignores_notes_and_skips_sparse_trailing_export_note():
    result = importer.inspect(workbook_bytes())
    assert result['candidate_sheets'] == ['Product-US']
    assert result['sheet'] == 'Product-US'
    assert result['count'] == 1
    assert result['sheet_required'] is False
    assert any('已跳过疑似非商品尾部行' in warning for warning in result['warnings'])


def test_inspect_without_product_candidates_requests_manual_sheet_selection():
    book = openpyxl.Workbook()
    book.active.title = 'Notes'
    book.active.append(['Read me'])
    book.active.append(['Source export notes'])
    stream = io.BytesIO()
    book.save(stream)
    result = importer.inspect(stream.getvalue())
    assert result['candidate_sheets'] == []
    assert result['sheet_required'] is True
    assert result['sheet'] is None
    assert any('手动选择工作表' in warning for warning in result['warnings'])


def test_inspect_requires_choice_when_multiple_product_sheets_exist():
    data = workbook_bytes(extra_product_sheet=True)
    preview = importer.inspect(data)
    assert preview['sheet_required'] is True
    assert preview['candidate_sheets'] == ['Product-US', 'Product-CA']
    assert preview['count'] == 0
    selected = importer.inspect(data, 'Product-CA')
    assert selected['sheet'] == 'Product-CA'
    assert selected['count'] == 1


def test_import_inspect_endpoint_exposes_sheet_choice_and_refreshes_preview():
    from fastapi.testclient import TestClient
    from backend.app import app
    data = workbook_bytes(extra_product_sheet=True)
    with TestClient(app) as client:
        preview = client.post('/api/import/inspect', files={'file': ('sample.xlsx', data)})
        assert preview.status_code == 200
        assert preview.json()['candidate_sheets'] == ['Product-US', 'Product-CA']
        assert preview.json()['sheet_required'] is True
        selected = client.post('/api/import/inspect', files={'file': ('sample.xlsx', data)},
                               data={'sheet': 'Product-CA'})
        assert selected.status_code == 200
        assert selected.json()['sheet'] == 'Product-CA'
        assert selected.json()['count'] == 1
        assert selected.json()['sample'][0]['fields']['asin'] == 'B087654321'


def test_model_attributes_keep_source_namespace_and_leave_missing_fields_out():
    source = {
        'sku': 'Color: Aqua | Size: 12',
        'parameters': 'Color: Blue | Fabric type: 100% Polyester',
    }
    value = {
        'facts': [],
        'source_attributes': [
            {'source': 'sku', 'attribute': 'Color', 'value': 'Aqua', 'evidence': 'Color: Aqua'},
            {'source': 'sku', 'attribute': 'Size', 'value': '12', 'evidence': 'Size: 12'},
            {'source': 'parameters', 'attribute': 'Color', 'value': 'Blue', 'evidence': 'Color: Blue'},
            {'source': 'parameters', 'attribute': 'Fabric type', 'value': '100% Polyester',
             'evidence': 'Fabric type: 100% Polyester'},
        ],
    }
    result = text_facts.validate(value, source, require_source_attributes=True)
    assert result['source_attribute_columns'] == {
        'SKU-Color属性': 'Aqua',
        'SKU-Size属性': '12',
        '详细参数-Color属性': 'Blue',
        '详细参数-Fabric type属性': '100% Polyester',
    }
    assert result['source_attributes_complete'] is True


def test_unverifiable_model_attribute_is_discarded_and_warned():
    result = text_facts.validate({
        'facts': [],
        'source_attributes': [
            {'source': 'sku', 'attribute': 'Color', 'value': 'Red', 'evidence': 'Color: Aqua'},
        ],
    }, {'sku': 'Color: Aqua'}, salvage=True, require_source_attributes=True)
    assert result['source_attribute_columns'] == {}
    assert result['source_attributes_complete'] is False
    assert any(item['code'] == 'discarded_source_attribute' for item in result['validation_warnings'])


def test_imported_attributes_become_visible_raw_columns(dataset):
    with Session.begin() as session:
        product = session.scalar(select(Product).where(Product.dataset_id == dataset['id']).order_by(Product.row))
        assert text_facts._write_source_attribute_columns(session, dataset['id'], product, {
            'source_attribute_columns': {'SKU-Color属性': 'Aqua', '详细参数-Color属性': 'Blue'},
        })
    with Session() as session:
        product = session.scalar(select(Product).where(Product.dataset_id == dataset['id']).order_by(Product.row))
        stored = require(session, Product, product.id)
        current_dataset = require(session, Dataset, dataset['id'])
        assert stored.raw['SKU-Color属性'] == 'Aqua'
        assert stored.raw['详细参数-Color属性'] == 'Blue'
        assert current_dataset.headers[-2:] == ['SKU-Color属性', '详细参数-Color属性']


def test_import_task_splits_attributes_and_downloads_image_without_touching_source(monkeypatch):
    from backend.import_image_download import ImageDownloadResult
    from backend import import_image_download
    from PIL import Image
    image = io.BytesIO()
    Image.new('RGB', (4, 3), 'blue').save(image, format='PNG')
    payload = image.getvalue()
    original = workbook_bytes()
    dataset = importer.import_file(original, '属性和图片.xlsx')

    def request_json(_model, messages, _metrics, _lock, *_args):
        inputs = json.loads(messages[-1]['content'])
        return {'results': [{
            'product_id': item['product_id'],
            'facts': [],
            'source_attributes': [
                {'source': 'sku', 'attribute': 'Color', 'value': 'Aqua', 'evidence': 'Color: Aqua'},
                {'source': 'sku', 'attribute': 'Size', 'value': '12', 'evidence': 'Size: 12'},
                {'source': 'parameters', 'attribute': 'Color', 'value': 'Blue', 'evidence': 'Color: Blue'},
                {'source': 'parameters', 'attribute': 'Fabric type', 'value': '100% Polyester',
                 'evidence': 'Fabric type: 100% Polyester'},
            ],
        } for item in inputs]}

    monkeypatch.setattr(batch, 'request_json', request_json)
    monkeypatch.setattr(import_image_download, 'download_image',
                        lambda _url: ImageDownloadResult(data=payload, content_type='image/png'))
    job = text_facts.prepare(dataset['id'])
    batch.active[job['id']] = threading.Event()
    batch.execute(job['id'])
    with Session() as session:
        result = obj(require(session, BatchJob, job['id']))
        product = session.scalar(select(Product).where(Product.dataset_id == dataset['id']))
        stored_dataset = require(session, Dataset, dataset['id'])
        asset = session.scalar(select(Asset).where(Asset.product_id == product.id, Asset.url.is_not(None)))
        assert result['status'] == 'completed'
        assert result['metrics']['attributes_completed'] == 1
        assert result['metrics']['images_completed'] == 1
        assert result['metrics']['images_failed'] == 0
        assert product.raw['SKU-Color属性'] == 'Aqua'
        assert product.raw['详细参数-Color属性'] == 'Blue'
        assert 'SKU-Color属性' in stored_dataset.headers
        assert asset.path and Path(asset.path).is_file()
        assert asset.url == 'https://images.example/item.jpg'
    assert Path(dataset['original_file']).read_bytes() == original


def test_image_failure_is_reported_by_row_and_retried_without_repeating_model(monkeypatch):
    from backend.import_image_download import ImageDownloadResult
    from backend import import_image_download
    from PIL import Image
    image = io.BytesIO()
    Image.new('RGB', (4, 3), 'green').save(image, format='PNG')
    payload = image.getvalue()
    dataset = importer.import_file(workbook_bytes(), '可重试图片.xlsx')
    model_calls = []

    def request_json(_model, messages, _metrics, _lock, *_args):
        model_calls.append(True)
        inputs = json.loads(messages[-1]['content'])
        return {'results': [{
            'product_id': item['product_id'], 'facts': [],
            'source_attributes': [
                {'source': 'sku', 'attribute': 'Color', 'value': 'Aqua', 'evidence': 'Color: Aqua'},
                {'source': 'parameters', 'attribute': 'Fabric type', 'value': '100% Polyester',
                 'evidence': 'Fabric type: 100% Polyester'},
            ],
        } for item in inputs]}

    monkeypatch.setattr(batch, 'request_json', request_json)
    monkeypatch.setattr(import_image_download, 'download_image',
                        lambda _url: ImageDownloadResult(error_code='timeout', error_message='下载图片超时'))
    first = text_facts.prepare(dataset['id'])
    batch.active[first['id']] = threading.Event()
    batch.execute(first['id'])
    with Session() as session:
        failed = obj(require(session, BatchJob, first['id']))
        assert failed['status'] == 'partial'
        assert any('第2行' in message and '下载图片超时' in message
                   for key, message in failed['errors'].items() if key.startswith('image:'))
    assert len(model_calls) == 1

    monkeypatch.setattr(import_image_download, 'download_image',
                        lambda _url: ImageDownloadResult(data=payload, content_type='image/png'))
    with Session.begin() as session:
        require(session, BatchJob, first['id']).status = 'queued'
    batch.active[first['id']] = threading.Event()
    text_facts.execute(first['id'])
    with Session() as session:
        retried = obj(require(session, BatchJob, first['id']))
        product = session.scalar(select(Product).where(Product.dataset_id == dataset['id']))
        asset = session.scalar(select(Asset).where(Asset.product_id == product.id, Asset.url.is_not(None)))
        assert retried['status'] == 'completed'
        assert retried['metrics']['images_completed'] == 1
        assert retried['metrics']['images_failed'] == 0
        assert retried['metrics']['images_pending'] == 0
        assert not retried['errors']
        assert asset.path and Path(asset.path).is_file()
    assert len(model_calls) == 1

