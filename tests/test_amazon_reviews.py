import json
import time
import threading

from openpyxl import load_workbook

from backend.amazon_reviews import service, sorftime
from backend.keyword_research.llm_client import LLMError


ASIN_A = "B0D5TQJXRN"
ASIN_B = "B0GY6VX976"
NODE = "17269957011"


class FakeResponse:
    def __init__(self, value):
        self.value = value

    def json(self):
        return self.value


class FakeTranslator:
    def __init__(self, *_args):
        pass

    def complete(self, messages, **_kwargs):
        inputs = json.loads(messages[-1]["content"])
        return FakeResponse({"translations": [
            {"index": row["index"], "title_zh": "译：" + row["title"],
             "text_zh": "译：" + row["text"]}
            for row in inputs
        ]})


class OneBasedTranslator:
    def complete(self, messages, **_kwargs):
        inputs = json.loads(messages[-1]["content"])
        return FakeResponse({"translations": [
            {"index": row["index"] + 1, "title_zh": "译：" + row["title"],
             "text_zh": "译：" + row["text"]}
            for row in inputs
        ]})


class GenericFieldTranslator:
    def complete(self, messages, **_kwargs):
        inputs = json.loads(messages[-1]["content"])
        return FakeResponse({"translations": [
            {"index": row["index"], "title": "译：" + row["title"],
             "text": "译：" + row["text"]}
            for row in inputs
        ]})


class InvalidThenValidTranslator:
    def __init__(self, *_args):
        self.calls = 0

    def complete(self, messages, **_kwargs):
        self.calls += 1
        if self.calls == 1:
            class InvalidResponse:
                @staticmethod
                def json():
                    raise LLMError("LLM 输出不是合法 JSON")
            return InvalidResponse()
        assert "上一次回答不是合法 JSON" in messages[-1]["content"]
        inputs = json.loads(messages[-2]["content"])
        return FakeResponse({"translations": [
            {"index": row["index"], "title_zh": "译：" + row["title"],
             "text_zh": "译：" + row["text"]}
            for row in inputs
        ]})


def test_extract_sorftime_category_and_review_envelopes():
    category = {"data": {"top100_products": [{"asin": ASIN_A, "title": "A"}]}}
    reviews = {"content": [{"type": "text", "text": json.dumps({
        "data": [{"review_date": "20260901", "star_rating": 5, "title": "Good", "content": "Works"}]
    })}]}

    assert service._extract_products(category) == [{
        "asin": ASIN_A, "title": "A", "category": "", "rank": None,
    }]
    assert service._extract_reviews(reviews) == [{
        "review_date": "20260901", "star_rating": 5, "title": "Good", "content": "Works",
    }]


def test_review_identity_uses_source_composite_key():
    first = {"asin": ASIN_A, "review_date": "20260901", "title": "Good", "content": "Works", "review_id": "one"}
    second = {**first, "review_id": "different"}
    distinct = {**first, "review_date": "20260902"}

    assert service._review_key(first) == service._review_key(second)
    assert service._review_key(first) != service._review_key(distinct)


def test_translation_accepts_one_based_model_indexes(monkeypatch, tmp_path):
    monkeypatch.setattr(service, "ROOT", tmp_path / "amazon-reviews")
    job_id = "b" * 32
    job_dir = service._job_dir(job_id)
    (job_dir / "raw").mkdir(parents=True)
    job = {"id": job_id, "status": "running", "stage": "translating"}
    service._save(job)
    item = {"asin": ASIN_A}
    reviews = [
        {"asin": ASIN_A, "title": "Easy to use", "content": "It works well."},
        {"asin": ASIN_A, "title": "Good value", "content": "Worth the price."},
    ]

    service._translate_reviews(OneBasedTranslator(), job, item, reviews, threading.Event())

    assert [row["t_title"] for row in reviews] == ["译：Easy to use", "译：Good value"]
    assert [row["t_content"] for row in reviews] == ["译：It works well.", "译：Worth the price."]


def test_translation_accepts_generic_title_and_text_fields(monkeypatch, tmp_path):
    monkeypatch.setattr(service, "ROOT", tmp_path / "amazon-reviews")
    job_id = "c" * 32
    job_dir = service._job_dir(job_id)
    (job_dir / "raw").mkdir(parents=True)
    job = {"id": job_id, "status": "running", "stage": "translating"}
    service._save(job)
    item = {"asin": ASIN_A}
    reviews = [{"asin": ASIN_A, "title": "Worth buying", "content": "It works as promised."}]

    service._translate_reviews(GenericFieldTranslator(), job, item, reviews, threading.Event())

    assert reviews[0]["t_title"] == "译：Worth buying"
    assert reviews[0]["t_content"] == "译：It works as promised."


def test_translation_retries_one_malformed_json_completion(monkeypatch, tmp_path):
    monkeypatch.setattr(service, "ROOT", tmp_path / "amazon-reviews")
    job_id = "d" * 32
    job_dir = service._job_dir(job_id)
    (job_dir / "raw").mkdir(parents=True)
    job = {"id": job_id, "status": "running", "stage": "translating"}
    service._save(job)
    translator = InvalidThenValidTranslator()
    reviews = [{"asin": ASIN_A, "title": "Good", "content": "Works well."}]

    service._translate_reviews(translator, job, {"asin": ASIN_A}, reviews, threading.Event())

    assert translator.calls == 2
    assert reviews[0]["t_title"] == "译：Good"
    assert reviews[0]["t_content"] == "译：Works well."


def test_sorftime_exception_groups_keep_leaf_cause_without_secrets():
    error = ExceptionGroup("transport failed", [
        RuntimeError("GET https://mcp.example.test/mcp?key=private-value failed"),
        ValueError("Bearer private-value"),
    ])

    message = sorftime._scrub(sorftime._diagnostic(error), "private-value")

    assert "private-value" not in message
    assert "?key=[hidden]" in message
    assert "RuntimeError" in message and "ValueError" in message


def test_live_shaped_node_job_deduplicates_translates_and_exports(monkeypatch, tmp_path):
    monkeypatch.setattr(service, "ROOT", tmp_path / "amazon-reviews")
    monkeypatch.setattr(service, "SHUTTING_DOWN", False)
    monkeypatch.setattr(service.sorftime_settings, "connection", lambda: ("https://mcp.invalid/mcp", "test-secret"))
    monkeypatch.setattr(service, "model_target", lambda *_args: ("https://model.invalid/v1", "model-secret", "mock/model"))
    monkeypatch.setattr(service, "require_validated", lambda *_args: True)
    monkeypatch.setattr(service, "LLMClient", FakeTranslator)

    calls_by_asin = {}
    api_calls = []

    def fake_sorftime(_endpoint, _key, tool, arguments):
        api_calls.append((tool, arguments))
        if tool == "category_report":
            return {"data": {"top100_products": [
                {"asin": ASIN_A, "title": "Product A", "rank": 1},
                {"asin": ASIN_B, "title": "Product B", "rank": 2},
                {"asin": "B000000000", "title": "Not selected", "rank": 3},
            ]}}
        assert tool == "product_reviews"
        assert arguments["amz_site"] == "US"
        assert arguments["review_type"] == "Both"
        asin = arguments["asin"]
        calls_by_asin[asin] = calls_by_asin.get(asin, 0) + 1
        shared = {
            "variant_attribute": "Blue",
            "review_date": "20260901",
            "star_rating": 5,
            "title": "=Works well",
            "content": "The light is bright and easy to use.",
        }
        new_review = {
            "variant_attribute": "Blue",
            "review_date": "20260902",
            "star_rating": 1,
            "title": "Stopped working",
            "content": "It stopped after one day.",
        }
        return {"data": [shared] if calls_by_asin[asin] == 1 else [shared, new_review]}

    monkeypatch.setattr(service.sorftime, "call", fake_sorftime)
    created = service.create({
        "mode": "node_asins", "site": "US", "node_id": NODE, "asins": [ASIN_A, ASIN_B],
        "rank_limit": 50, "model": "mock/model", "credential_id": "mock-key",
    })

    deadline = time.monotonic() + 10
    result = service.get(created["id"])
    while result["status"] not in ("completed", "completed_with_errors", "failed"):
        if time.monotonic() >= deadline:
            raise AssertionError("评论任务未在测试时限内结束")
        time.sleep(0.02)
        result = service.get(created["id"])
    worker = service.ACTIVE.get(created["id"])
    if worker:
        worker.join(timeout=2)
    result = service.get(created["id"])

    assert result["status"] == "completed"
    assert [item["asin"] for item in result["items"]] == [ASIN_A, ASIN_B]
    assert result["metrics"] == {"review_tool_calls": 4, "raw_records": 6, "new_unique_records": 4}
    assert len(api_calls) == 5  # one category report and two identical Both calls per ASIN
    assert all(calls_by_asin[asin] == 2 for asin in (ASIN_A, ASIN_B))

    raw = json.loads(service.raw_path(created["id"]).read_text(encoding="utf-8"))
    assert len(raw["reviews"]) == 4
    assert all(row["source_site"] == "US" and row["source_node"] == NODE for row in raw["reviews"])
    assert all(row["t_content"] and row["t_content"] == row["translation_zh"] for row in raw["reviews"])
    assert all(row["t_title"] and row["t_title"] == row["title_zh"] for row in raw["reviews"])
    assert all(row["sentiment"] in ("Positive", "Negative") for row in raw["reviews"])
    assert all(len(raw["review_call_responses"][asin]) == 2 for asin in (ASIN_A, ASIN_B))

    workbook = load_workbook(service.output_path(created["id"]), data_only=True)
    assert workbook.sheetnames == ["全部评论", "商品评论统计", "统计汇总"]
    all_reviews = workbook["全部评论"]
    assert all_reviews.max_row == 5
    assert all_reviews.max_column == 14
    assert all_reviews.cell(1, 1).value == "#"
    assert all_reviews.cell(2, 10).value == "'=Works well"  # Excel formula injection is neutralized
    assert workbook["商品评论统计"].max_row == 3
    summary = {workbook["统计汇总"].cell(row, 1).value: workbook["统计汇总"].cell(row, 2).value
               for row in range(2, workbook["统计汇总"].max_row + 1)}
    assert summary["去重后评论总数"] == 4
    assert summary["API 调用次数"] == 4
    assert summary["翻译覆盖"] == "4/4 条"


def test_create_rejects_unconfigured_sorftime(monkeypatch):
    monkeypatch.setattr(service.sorftime_settings, "connection", lambda: (_ for _ in ()).throw(ValueError("not configured")))
    monkeypatch.setattr(service, "model_target", lambda *_args: ("https://model.invalid/v1", "secret", "mock/model"))
    monkeypatch.setattr(service, "require_validated", lambda *_args: True)

    try:
        service.create({
            "mode": "asins", "site": "US", "asins": [ASIN_A],
            "model": "mock/model", "credential_id": "mock-key",
        })
    except ValueError as exc:
        assert str(exc) == "not configured"
    else:
        raise AssertionError("未配置 Sorftime 时不应创建任务")


def test_retry_resumes_the_missing_review_pass(monkeypatch, tmp_path):
    monkeypatch.setattr(service, "ROOT", tmp_path / "amazon-reviews")
    monkeypatch.setattr(service, "SHUTTING_DOWN", False)
    monkeypatch.setattr(service.sorftime_settings, "connection", lambda: ("https://mcp.invalid/mcp", "test-secret"))
    monkeypatch.setattr(service, "model_target", lambda *_args: ("https://model.invalid/v1", "model-secret", "mock/model"))
    monkeypatch.setattr(service, "require_validated", lambda *_args: True)
    monkeypatch.setattr(service, "LLMClient", FakeTranslator)
    requested = []
    row = {"review_date": "20260901", "star_rating": 5, "title": "Good", "content": "Works well"}

    def fake_sorftime(_endpoint, _key, tool, arguments):
        requested.append((tool, arguments))
        return {"data": [row]}

    monkeypatch.setattr(service.sorftime, "call", fake_sorftime)
    job_id = "a" * 32
    job = {
        "id": job_id, "status": "running", "mode": "asins", "site": "US", "node_id": "",
        "asins": [ASIN_A], "rank_limit": 50, "model": "mock/model", "credential_id": "mock-key",
        "created_at": time.time(), "updated_at": time.time(), "attempt": 1, "stage": "interrupted",
        "review_collection_contract": service.REVIEW_CALL_CONTRACT, "review_count": 1, "errors": [],
        "metrics": {"review_tool_calls": 1, "raw_records": 1, "new_unique_records": 1},
        "items": [{
            "asin": ASIN_A, "title": "Product A", "status": "pending", "review_count": 1, "error": "",
            "review_call_metrics": [{"round": 1, "review_type": "Both", "raw_record_count": 1, "new_unique_count": 1}],
        }],
    }
    service._job_dir(job_id).mkdir(parents=True)
    (service._job_dir(job_id) / "raw").mkdir()
    service._save(job)
    first_call = [{"round": 1, "review_type": "Both", "payload": {"data": [row]}}]
    service._atomic_json(service._call_payloads_path(job_id, ASIN_A), first_call)
    reviews, _ = service._reviews_from_calls(first_call, ASIN_A)
    service._annotate_reviews(reviews, job, job["items"][0])
    service._write_reviews(job_id, ASIN_A, reviews)

    service._run(job_id, threading.Event())
    result = service.get(job_id)
    raw = json.loads(service.raw_path(job_id).read_text(encoding="utf-8"))

    assert result["status"] == "completed"
    assert len(requested) == 1  # only missing pass 2 was called
    assert requested[0][0] == "product_reviews"
    assert requested[0][1] == {"amz_site": "US", "asin": ASIN_A, "review_type": "Both"}
    assert result["metrics"] == {"review_tool_calls": 2, "raw_records": 2, "new_unique_records": 1}
    assert len(raw["reviews"]) == 1
    assert raw["reviews"][0]["source_review_rounds"] == [1, 2]
    assert raw["reviews"][0]["translation_zh"] == "译：Works well"
