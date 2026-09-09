import copy,json,threading
import openpyxl
import pytest
from sqlalchemy import select
from backend.db import *
from backend import draft,batch,text_facts,tagging_rules,attributes,draft_export,importer
from test_draft import add,job,run

def fake_facts(model,messages,metrics,lock):
    inputs=json.loads(messages[-1]['content'])
    return {'results':[{'product_id':p['product_id'],'facts':[{'attribute':'原始标题','value':str(p['text']['title']), 'raw_value':p['text']['title'],'state':'observed','evidence':[{'field':'title','quote':p['text']['title']}]}]} for p in inputs]}

def test_rules_frozen_and_enter_request(dataset,monkeypatch):
    did=dataset['id'];col=add(did,judgment={'synonyms':{'Two Piece':'2件套'}});j=job(did,col)
    sha=j['spec']['columns'][0]['definition']['skill']['sha256'];seen=[]
    def fake(model,messages,metrics,lock):
        seen.append(messages[0]['content'])
        return {'results':[{'product_id':p['product_id'],'normalized_value':'2件套','evidence':'Two Piece'} for p in json.loads(messages[-1]['content'])]}
    monkeypatch.setattr(batch,'request_json',fake)
    monkeypatch.setattr(tagging_rules,'freeze',lambda:pytest.fail('worker must use frozen rules'))
    assert run(j)['status']=='completed'
    assert all('Two Piece' in s and '命中及排除' in s for s in seen)
    assert draft.sheet(did)['rows'][0]['cells'][col['id']]['evidence']['skill_sha256']==sha

def test_text_extraction_cache_and_no_vision(dataset,monkeypatch):
    did=dataset['id'];monkeypatch.setattr(batch,'request_json',fake_facts)
    monkeypatch.setattr(batch,'observation',lambda *a:pytest.fail('must not call vision'))
    j=run(text_facts.prepare(did));assert j['status']=='completed' and len(j['completed'])==10
    assert len(text_facts.records(did))==10 and not draft.sheet(did)['columns']
    other=importer.import_file((ROOT/'测试10行.xlsx').read_bytes(),'重复文件.xlsx')['id']
    monkeypatch.setattr(batch,'request_json',lambda *a:pytest.fail('cache should satisfy upload'))
    repeated=run(text_facts.prepare(other));assert repeated['metrics']['cache_hits']==10
    assert all(r['product_id'] not in j['spec']['product_ids'] for r in text_facts.records(other))
    with Session.begin() as s:
        p=require(s,Product,j['spec']['product_ids'][0]);p.fields={**p.fields,'title':'新的商品标题'}
    with pytest.raises(ValueError,match='原文已变化'):text_facts.for_batch(j['spec']['product_ids'][0],threading.Event())
    nextjob=text_facts.prepare(did);assert len(nextjob['spec']['expected'])==1

def test_text_partial_retry_and_pause(dataset,monkeypatch):
    j=text_facts.prepare(dataset['id']);bad=j['spec']['product_ids'][0];seen=[]
    def partial(*args):
        r=fake_facts(*args);r['results']=[x for x in r['results'] if x['product_id']!=bad];return r
    monkeypatch.setattr(batch,'request_json',partial)
    result=run(j);assert result['status']=='partial' and len(result['completed'])==9
    def retry(*args):seen.extend(json.loads(args[1][-1]['content']));return fake_facts(*args)
    monkeypatch.setattr(batch,'request_json',retry)
    assert run(result)['status']=='completed' and len(seen)==1
    with Session.begin() as s:
        p=require(s,Product,bad);p.fields={**p.fields,'title':'变化'}
    j=text_facts.prepare(dataset['id']);event=threading.Event();event.set();batch.active[j['id']]=event;batch.execute(j['id'])
    with Session() as s:assert require(s,BatchJob,j['id']).status=='paused'
    with pytest.raises(ValueError,match='暂停'):text_facts.for_batch(bad,threading.Event())

def test_exact_evidence_and_conflicts():
    value={'facts':[{'attribute':'件数','value':None,'state':'conflict','reason':'标题与参数冲突','evidence':[{'field':'title','quote':'Two Piece'},{'field':'parameters','quote':'共3件'}]}]}
    source={'title':'Two Piece set','parameters':'共3件'}
    assert text_facts.validate(value,source)
    invalid=copy.deepcopy(value);invalid['facts'][0]['value']='2件套'
    with pytest.raises(ValueError):text_facts.validate(invalid,source)
    invalid=copy.deepcopy(value);invalid['facts'][0]['evidence'][0]['quote']='Three Piece'
    with pytest.raises(ValueError,match='定位'):text_facts.validate(invalid,source)
    invalid={'facts':[{'attribute':'尺寸','value':'7-8 Years','state':'observed','evidence':[{'field':'title','quote':'7-8 Years'}]}]}
    with pytest.raises(ValueError,match='中文'):text_facts.validate(invalid,{'title':'7-8 Years'})
    combined=copy.deepcopy(value);combined['facts'].append({'attribute':'件数','value':'3件','state':'observed','evidence':[{'field':'parameters','quote':'共3件'}]})
    normalized=text_facts.validate(combined,source)
    assert len(normalized['facts'])==1 and normalized['facts'][0]['value'] is None and len(normalized['facts'][0]['evidence'])==2

def test_text_fact_salvage_normalizes_model_format_without_inventing_evidence():
    source={
        'sku':'Color: A-black 1 | Size: 13-14 Years',
        'parameters':'Fabric type: Shell:82%Nylon+18%Elastane',
        '商品标题':'Girls Swimsuit 3 Piece Tankini for Teen Girls',
    }
    value={'facts':[
        {'attribute':'颜色','value':'黑色','state':'observed','evidence':[{'field':'text[\'sku\']','quote':'Black'}]},
        {'attribute':'件数','value':'3件套','state':'observed','evidence':[{'field':'text[\'商品标题\']','quote':'3 Piece'}]},
        {'attribute':'面料','value':'82%锦纶','state':'observed','evidence':[{'field':'text[\'参数\']','quote':'82%Nylon'}]},
        {'attribute':'适用人群','value':'少女','state':'observed','evidence':[{'field':'商品标题','quote':'Teen Girls'},{'field':'sku','quote':''}]},
        {'attribute':'颜色款名','value':'待确认','raw_value':'Blue Coconut','state':'unknown','reason':'款名语义不明确','evidence':[{'field':'sku','quote':'Color: A-black 1'}]},
        {'attribute':'款式','value':'分体泳衣','state':'observed','evidence':[{'field':'商品标题','quote':'Tankini missing words'}]},
    ]}
    result=text_facts.validate(value,source,salvage=True)
    assert [f['attribute'] for f in result['facts']]==['颜色','件数','面料','适用人群','颜色款名']
    assert result['facts'][0]['evidence'][0]['quote']=='black'
    assert result['facts'][1]['evidence'][0]['field']=='商品标题'
    assert result['facts'][2]['evidence'][0]['field']=='parameters'
    assert len(result['facts'][3]['evidence'])==1
    assert result['facts'][4]['value'] is None and result['facts'][4]['discarded_value']=='待确认'
    assert {w['code'] for w in result['validation_warnings']} >= {'normalized_evidence','discarded_evidence','discarded_fact','discarded_nondefinitive_value'}
    with pytest.raises(ValueError,match='定位'):
        text_facts.validate({'facts':[value['facts'][-1]]},source)

def test_text_cancel_recovery_and_stale_result(dataset,monkeypatch):
    did=dataset['id'];j=text_facts.prepare(did)
    event=threading.Event();batch.active[j['id']]=event;batch.stop(j['id'],True);batch.execute(j['id'])
    with Session() as s:assert require(s,BatchJob,j['id']).status=='cancelled'
    monkeypatch.setattr(batch,'request_json',fake_facts)
    j=text_facts.prepare(did);batch.recover()
    with Session() as s:assert require(s,BatchJob,j['id']).status=='interrupted'
    assert run(j)['status']=='completed'
    p=draft.sheet(did)['rows'][0]
    with Session.begin() as s:row=require(s,Product,p['id']);row.fields={**row.fields,'title':'重新命名的标题'}
    nextjob=text_facts.prepare(did)
    def changed(*args):
        with Session.begin() as s:row=require(s,Product,p['id']);row.fields={**row.fields,'title':'请求后发生修改'}
        return fake_facts(*args)
    monkeypatch.setattr(batch,'request_json',changed)
    assert run(nextjob)['status']=='partial'
    assert text_facts.records(did)[0]['status']!='ready'

def test_text_wait_and_dependency_resume(dataset,monkeypatch):
    did=dataset['id'];pre=text_facts.prepare(did);c=add(did);tag=job(did,c)
    started=threading.Event();release=threading.Event();calls=[]
    def dispatch(model,messages,metrics,lock):
        if '这是上传后的文字底稿提取' in messages[0]['content']:
            started.set();assert release.wait(5);return fake_facts(model,messages,metrics,lock)
        calls.append(1)
        return {'results':[{'product_id':p['product_id'],'normalized_value':'2件套','evidence':'标题'} for p in json.loads(messages[-1]['content'])]}
    monkeypatch.setattr(batch,'request_json',dispatch)
    batch.launch(tag['id']);batch.launch(pre['id']);assert started.wait(5)
    assert not calls;release.set()
    import time
    deadline=time.time()+10
    while (tag['id'] in batch.active or pre['id'] in batch.active) and time.time()<deadline:time.sleep(.05)
    with Session() as s:assert require(s,BatchJob,tag['id']).status=='completed'

def test_manual_multiple_keeps_candidate_and_requires_confirmation(dataset):
    did=dataset['id'];c=add(did,'pattern');pid=draft.sheet(did)['rows'][0]['id']
    draft.correct(did,pid,c['id'],['花卉','条纹'])
    state=draft.state(did);assert len(state['columns'])==1 and state['rows'][0]['cells'][c['id']]['status']=='review'
    prop=next(p for p in state['proposals'] if p['status']=='pending' and p['changes'][0]['kind']=='expand')
    draft.confirm(did,prop['id'])
    assert all(r['manual'] for cid,r in draft.sheet(did)['rows'][0]['cells'].items() if cid!=c['id'])

def multi_value():
    return {'normalized_value':['花卉','条纹'],'raw_value':'floral and stripes, no dots','evidence':'花卉和条纹，无波点',
            'attributes':[{'attribute':'花卉','value':True,'evidence':'floral'},{'attribute':'条纹','value':True,'evidence':'stripes'},{'attribute':'波点','value':False,'evidence':'no dots'}]}

def test_multi_expand_manual_migration_snapshot_and_export(dataset):
    did=dataset['id'];col=add(did,'pattern',options=['花卉','条纹','波点']);rows=draft.sheet(did)['rows'];pid=rows[0]['id'];manual=rows[1]['id']
    draft.correct(did,manual,col['id'],'花卉');v1=draft.save_version(did);frozen=copy.deepcopy(draft.sheet(did,v1['id']))
    j=job(did,col,[pid])
    with Session.begin() as s:
        task=require(s,BatchJob,j['id']);draft.write_result(s,task,pid,col,multi_value());task.status='completed'
    state=draft.state(did);assert state['rows'][0]['cells'][col['id']]['status']=='review'
    proposal=next(p for p in state['proposals'] if p['changes'][0]['kind']=='expand');draft.confirm(did,proposal['id'])
    state=draft.sheet(did);children={c['definition']['attribute']:c for c in state['columns'] if c['definition'].get('parent_id')==col['id']}
    assert len(children)==3 and len({c['id'] for c in children.values()})==3
    assert state['rows'][0]['cells'][children['波点']['id']]['value'] is False
    assert state['rows'][1]['cells'][children['花卉']['id']]['manual']
    assert state['rows'][1]['cells'][children['条纹']['id']]['value'] is None
    draft.correct(did,pid,children['花卉']['id'],False);v2=draft.save_version(did)
    assert draft.sheet(did,v1['id'])==frozen and draft.compare(did,v1['id'],v2['id'])
    selected=[children[n]['id'] for n in ['花卉','波点','条纹']]
    wb=openpyxl.load_workbook(draft_export.export(did,selected))
    ws=wb['Product_Result'];assert ws.max_row==11 and ws.max_column==73 and len(ws._images)==10
    assert [ws.cell(2,i).value for i in range(71,74)]==['❌','❌','✓']
    assert ws.cell(3,73).value is None and ws.cell(3,73).comment
    assert wb['AI_Detail'].max_row==31 and wb['Tag_Dictionary'].max_row==4

def test_invalid_multi_and_unknown_not_expanded(dataset):
    d={'options':['花卉','条纹']}
    with pytest.raises(ValueError):attributes.validate_result({'normalized_value':['花卉','条纹']},d)
    unknown={'normalized_value':None,'evidence':'可能花卉或条纹，无法判断','attributes':[{'attribute':'花卉','value':None},{'attribute':'条纹','value':None}]}
    attributes.validate_result(unknown,d)
    did=dataset['id'];col=add(did);pid=draft.sheet(did)['rows'][0]['id'];j=job(did,col,[pid])
    with Session.begin() as s:draft.write_result(s,require(s,BatchJob,j['id']),pid,col,unknown)
    assert not any(p['changes'][0]['kind']=='expand' for p in draft.state(did)['proposals'])

def test_attribute_edit_conflict_dependency_and_new_attribute(dataset):
    did=dataset['id'];parent=add(did,'pattern',options=['花卉']);p=draft.propose(did,'拆列',[{'kind':'expand','column_id':parent['id'],'attributes':['花卉','条纹']}]);draft.confirm(did,p['id'])
    c=next(c for c in draft.sheet(did)['columns'] if c['definition'].get('attribute')=='花卉');pid=draft.sheet(did)['rows'][0]['id']
    derived=add(did,'group',method='mapping',depends_on=[c['id']],mapping={'true':'花卉组','false':'其他组'})
    j=job(did,c,[pid]);draft.correct(did,pid,c['id'],False)
    assert draft.sheet(did)['rows'][0]['cells'][derived['id']]['value']=='其他组'
    with Session.begin() as s:draft.write_result(s,require(s,BatchJob,j['id']),pid,c,{'normalized_value':True,'evidence':'旧判断'})
    assert draft.sheet(did)['rows'][0]['cells'][c['id']]['conflict']
    with pytest.raises(ValueError):draft.correct(did,pid,c['id'],'false')
    p=draft.propose(did,'追加',[{'kind':'expand','column_id':parent['id'],'attributes':['波点']}]);draft.confirm(did,p['id'])
    assert next(x for x in draft.sheet(did)['columns'] if x['id']==c['id'])['id']==c['id']
    with pytest.raises(ValueError):job(did,parent)
