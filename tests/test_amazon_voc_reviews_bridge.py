from __future__ import annotations

import hashlib
import json
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.amazon_reviews import service as reviews_service
from backend.amazon_voc import router as voc_router
import backend.amazon_voc.service as voc_service_module
from backend.amazon_voc.service import AmazonVOCService
from backend.runtime_isolation import CredentialUnavailable


def review_workbook() -> bytes:
    from io import BytesIO

    book = Workbook()
    sheet = book.active
    sheet.title = "全部评论"
    sheet.append(["ASIN", "商品标题", "星级", "评论内容(EN)", "评论日期"])
    sheet.append(["B001TEST", "Fixture product", 2, "The support feels unstable.", "2026-09-01"])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.fixture
def bridge(tmp_path, monkeypatch):
    instance = AmazonVOCService(root=tmp_path / "amazon-voc")
    monkeypatch.setattr(voc_router, "service", instance)
    monkeypatch.setattr(instance, "start_worker", lambda **_kwargs: None)
    monkeypatch.setattr(instance, "public_model_settings", lambda: {
        "configured": True, "default_model": "shared/fake-text",
        "default_credential_id": "cred-1", "providers": [],
    })
    monkeypatch.setattr(instance, "resolve_shared_model", lambda model_ref=None, credential_id=None: {
        "model_ref": model_ref or "shared/fake-text", "credential_id": credential_id or "cred-1",
        "base_url": "https://model.invalid/v1", "api_key": "test-secret",
        "model": "fake-text",
    })
    monkeypatch.setattr(reviews_service, "ROOT", tmp_path / "amazon-reviews")
    api = TestClient(__import__("backend.app", fromlist=["app"]).app)
    return instance, api


def add_source_job(job_id: str, status="completed", review_count=1) -> Path:
    folder = reviews_service._job_dir(job_id)
    (folder / "raw").mkdir(parents=True)
    reviews_service._save({
        "id": job_id, "status": status, "mode": "asins", "site": "GB", "node_id": "",
        "asins": ["B001TEST"], "rank_limit": 50, "model": "source/model",
        "credential_id": "source-key-id", "created_at": time.time(), "updated_at": time.time(),
        "stage": "已完成", "attempt": 1, "items": [], "errors": [],
        "review_count": review_count, "metrics": {},
    })
    (folder / "result.xlsx").write_bytes(review_workbook())
    return folder


def handoff(api: TestClient, job_id: str, confirmed=True):
    return api.post(f"/api/amazon-voc/from-amazon-reviews/{job_id}",
                    json={"confirmed_reviewed": confirmed})


def test_reviewed_workbook_creates_one_project_and_idempotent_voc_run(bridge):
    voc, api = bridge
    job_id = "a" * 32
    folder = add_source_job(job_id, status="completed_with_errors")

    first = handoff(api, job_id)
    assert first.status_code == 200, first.text
    result = first.json()
    assert result["duplicate"] is False
    assert result["project"]["marketplace"] == result["project"]["country"] == "GB"
    assert result["project"]["language"] == "auto"
    assert result["project"]["category"] == ""
    assert job_id[:8] in result["project"]["name"]
    assert result["run"]["status"] == "uploaded"
    assert result["run"]["sources"]["analysis_mode"] == "full_voc_v2"
    digest = hashlib.sha256((folder / "result.xlsx").read_bytes()).hexdigest()
    assert result["run"]["sources"]["reviews_sha256"] == digest
    assert result["source_job"]["voc_handoff"] == {
        "source_job_id": job_id, "workbook_sha256": digest,
        "project_id": result["project"]["id"], "run_id": result["run"]["id"],
        "created_at": result["source_job"]["voc_handoff"]["created_at"], "status": "started",
    }
    snapshot = voc.model_snapshot(result["run"]["id"])
    assert snapshot["model_ref"] == "shared/fake-text"
    assert snapshot["credential_id"] == "cred-1"
    assert "test-secret" not in voc.db.path.read_text(encoding="utf-8", errors="ignore")

    second = handoff(api, job_id)
    assert second.status_code == 200, second.text
    assert second.json()["duplicate"] is True
    assert second.json()["run"]["id"] == result["run"]["id"]
    assert second.json()["project"]["id"] == result["project"]["id"]
    assert voc.db.one("SELECT COUNT(*) AS n FROM workbench_reviews_handoffs")["n"] == 1
    assert voc.db.one("SELECT COUNT(*) AS n FROM runs")["n"] == 1
    saved_job = json.loads((folder / "job.json").read_text(encoding="utf-8"))
    assert saved_job["voc_handoff"]["run_id"] == result["run"]["id"]


def test_combined_reviews_auto_handoff_creates_full_voc_v2_run(bridge, monkeypatch):
    voc, _api = bridge
    monkeypatch.setattr(voc_service_module, "service", voc)
    job_id = "7" * 32
    add_source_job(job_id, status="completed", review_count=1)
    source_job = reviews_service.get(job_id)
    source_job["auto_voc_policy"] = "any_exported"
    reviews_service._save(source_job)

    reviews_service._auto_voc_handoff(reviews_service.get(job_id))

    saved_job = reviews_service.get(job_id)
    assert saved_job["voc_handoff"]["status"] == "started"
    run = voc.db.get_run(saved_job["voc_handoff"]["run_id"])
    assert run["sources"]["analysis_mode"] == "full_voc_v2"
    assert run["sources"]["reviews_sha256"] == saved_job["voc_handoff"]["workbook_sha256"]


@pytest.mark.parametrize("payload,status", [
    ({}, 422), ({"confirmed_reviewed": False}, 400), ({"confirmed_reviewed": "true"}, 422),
])
def test_confirmation_must_be_explicit_boolean_true(bridge, payload, status):
    _, api = bridge
    job_id = "b" * 32
    add_source_job(job_id)
    response = api.post(f"/api/amazon-voc/from-amazon-reviews/{job_id}", json=payload)
    assert response.status_code == status


@pytest.mark.parametrize("job_status,review_count", [
    ("queued", 1), ("running", 1), ("failed", 1), ("completed", 0),
])
def test_unfinished_unavailable_or_empty_source_job_is_rejected(bridge, job_status, review_count):
    _, api = bridge
    job_id = "c" * 32
    add_source_job(job_id, status=job_status, review_count=review_count)
    response = handoff(api, job_id)
    assert response.status_code == 400


def test_changed_workbook_conflicts_with_existing_handoff(bridge):
    _, api = bridge
    job_id = "d" * 32
    folder = add_source_job(job_id)
    created = handoff(api, job_id)
    assert created.status_code == 200, created.text
    (folder / "result.xlsx").write_bytes(review_workbook() + b"changed")
    conflict = handoff(api, job_id)
    assert conflict.status_code == 409


def test_missing_shared_credential_does_not_leave_an_empty_project(bridge, monkeypatch):
    voc, api = bridge
    job_id = "f" * 32
    add_source_job(job_id)

    def unavailable(*_args, **_kwargs):
        raise CredentialUnavailable("请先配置已授权的文本模型凭据")

    monkeypatch.setattr(voc, "resolve_shared_model", unavailable)
    response = handoff(api, job_id)

    assert response.status_code == 400
    assert voc.db.one("SELECT COUNT(*) AS n FROM projects")["n"] == 0
    assert voc.db.one("SELECT COUNT(*) AS n FROM runs")["n"] == 0
    assert voc.db.one("SELECT COUNT(*) AS n FROM workbench_reviews_handoffs")["n"] == 0


def test_oversize_is_rejected_before_workbook_read(bridge, monkeypatch):
    _, api = bridge
    job_id = "e" * 32
    add_source_job(job_id)

    class OversizePath:
        name = "result.xlsx"

        def stat(self):
            class Stat:
                st_size = voc_service_module.MAX_REVIEW_BYTES + 1
            return Stat()

        def read_bytes(self):
            raise AssertionError("oversize workbook must be rejected before reading")

    monkeypatch.setattr(reviews_service, "output_path", lambda _job_id: OversizePath())
    response = handoff(api, job_id)
    assert response.status_code == 400
    assert "50 MiB" in response.text


def test_review_source_scope_is_written_into_html_and_excel(tmp_path):
    from openpyxl import Workbook, load_workbook
    from backend.amazon_voc.service import _append_review_scope, _review_source_context

    context = _review_source_context({
        "id": "review-job", "status": "completed_with_errors", "review_count": 12,
        "items": [
            {"asin": "B000000001", "status": "completed", "review_count": 12},
            {"asin": "B000000002", "status": "failed", "review_count": 0},
            {"asin": "B000000003", "status": "completed", "review_count": 0},
        ],
    })
    assert context["requested_asin_count"] == 3
    assert context["comment_asin_count"] == 1
    assert context["failed_asins"] == ["B000000002"]
    assert "最近一年" in context["analysis_scope"]

    html_path = tmp_path / "report.html"
    html_path.write_text("<html><body><main>Report</main></body></html>", encoding="utf-8")
    excel_path = tmp_path / "report.xlsx"
    workbook = Workbook()
    workbook.active.title = "任务与方法说明"
    for index in range(11):
        workbook.create_sheet(f"VOC证据{index + 1}")
    workbook.save(excel_path)
    _append_review_scope(html_path, excel_path, context)

    html_report = html_path.read_text(encoding="utf-8")
    assert "评论采集范围与 ASIN 覆盖" in html_report
    assert "B000000002" in html_report
    workbook = load_workbook(excel_path, read_only=True)
    assert len(workbook.sheetnames) == 12
    assert "采集范围" not in workbook.sheetnames
    values = list(workbook["任务与方法说明"].values)
    metadata = {row[0]: row[1] for row in values if len(row) >= 2 and row[0]}
    assert metadata["有评论的 ASIN 数"] == 1
    assert metadata["失败或未完成 ASIN"] == "B000000002"
    _append_review_scope(html_path, excel_path, context)
    workbook = load_workbook(excel_path, read_only=True)
    values = list(workbook["任务与方法说明"].values)
    assert len(workbook.sheetnames) == 12
    assert sum(row[0] == "评论采集范围与 ASIN 覆盖" for row in values) == 1
