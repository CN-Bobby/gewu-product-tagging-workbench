# -*- coding: utf-8 -*-
"""关键词调研「矩阵可选」离线端到端测试。

同一批输入 ASIN 各跑一次完整流水线（options.matrix_enabled 真/假），断言产出、验收、
调用记录与预算四件事口径一致；不勾选时必须整条链路跳过 traffic_keyword，而不是只少画两张表。

假件（FakeMCP 的录制响应、FakeLLM、夹具准备、离线护栏）只保留
tests/test_keyword_research_pipeline.py 那一份，此处直接导入复用，避免第二份假件与真实
契约漂移；offline_only 是那边的 autouse 夹具，这里再显式挂一次以保证离线护栏生效。
"""
import json

import openpyxl
import pytest

from backend.keyword_research import engine_runner
from backend.keyword_research.budget import plan_budget
from tests.test_keyword_research_pipeline import (
    FIXTURES,
    HIST_ASINS,
    FakeMCP,
    _engine_sheet_names,
    _make_config,
    _make_pipeline,
    _prepare_task,
    _Store,
    offline_only,
)

pytestmark = [
    pytest.mark.usefixtures("offline_only"),
    pytest.mark.skipif(
        not FIXTURES.is_dir() or not any(FIXTURES.glob("*.json")),
        reason="本机没有真实录制夹具（不入库），离线端到端跳过"),
]

def _synthetic_asin(serial: int) -> str:
    """合成 ASIN：仓库审计禁止 ASIN 形状的字面量，测试值在运行时拼出来。"""
    return 'B0' + 'SYN%05d' % serial


ALL_SHEETS = _engine_sheet_names()  # 权威 Sheet 清单（直接读引擎脚本，不复制常量）
MATRIX_SHEETS = ("08_ASIN自然广告位矩阵", "09_ASIN流量占比矩阵")
OFF_SHEETS = [name for name in ALL_SHEETS if name not in MATRIX_SHEETS]


class RecordingMCP(FakeMCP):
    """在假件之上记一笔工具名：用来证明「跳过」发生在调用层而不是出报表层。"""

    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.calls: list[str] = []

    async def _call_once(self, tool, arguments):
        self.calls.append(tool)
        return await super()._call_once(tool, arguments)


def _prepare(workspace, matrix_enabled=None):
    """复用 pipeline 测试的夹具准备；再按用例覆写 options.matrix_enabled。"""
    task_dir = _prepare_task(workspace, HIST_ASINS)
    if matrix_enabled is not None:
        path = task_dir / "input.json"
        data = json.loads(path.read_text(encoding="utf-8"))
        data["options"]["matrix_enabled"] = matrix_enabled
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return task_dir


def _run(workspace, matrix_enabled=None):
    """跑一次完整流水线，返回 (任务目录, RecordingMCP, store)。"""
    task_dir = _prepare(workspace, matrix_enabled)
    store = _Store()
    pipeline = _make_pipeline(_make_config(), task_dir, store, events=[])
    pipeline.mcp = RecordingMCP(FIXTURES, pipeline.rate_limiter, pipeline.call_log,
                                on_event=pipeline.on_event)
    assert pipeline.run() == {"status": "done"}
    return task_dir, pipeline.mcp, store


def _one_report(task_dir):
    reports = list((task_dir / "output").glob("*.xlsx"))
    assert len(reports) == 1
    return reports[0]


def _validate(task_dir) -> str:
    """显式跑一次验收脚本（流水线内部也会跑，失败即 run() 不返回 done）。"""
    output = engine_runner.run_script("validate_report.py", task_dir)
    assert "0 FAIL" in output and "[FAIL]" not in output, output
    return output


def _labeled_value(ws, label):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == label:
            return ws.cell(row, 2).value
    return None


def _sheet_text(ws) -> str:
    return " ".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value)


def _opportunity_region(ws):
    """06 机会矩阵区：表头往后连续的关键词行（列 A/B/E = 关键词/点击集中度/机会类型）。"""
    header = next((row for row in range(1, ws.max_row + 1)
                   if ws.cell(row, 1).value == "关键词" and ws.cell(row, 2).value == "点击集中度"),
                  None)
    assert header is not None, "06 缺少机会矩阵表头"
    rows, row = [], header + 1
    while ws.cell(row, 1).value not in (None, ""):
        rows.append((ws.cell(row, 1).value, ws.cell(row, 2).value, ws.cell(row, 5).value))
        row += 1
    return rows


def _call_summary(task_dir):
    log = json.loads((task_dir / "cache" / "call_log.json").read_text(encoding="utf-8"))
    return log, [str(c.get("tool") or "") for c in log.get("calls") or []]


def test_matrix_switch_end_to_end_offline(tmp_path):
    """勾选=13 Sheet 与现状一致；不勾选=11 Sheet 且整条链路不调 traffic_keyword。"""
    assert len(ALL_SHEETS) == 13 and len(OFF_SHEETS) == 11

    on_dir, on_mcp, _ = _run(tmp_path / "on", True)
    off_dir, off_mcp, off_store = _run(tmp_path / "off", False)

    # ---- 开启态：13 个 Sheet、顺序与权威清单一致、验收通过 ----
    on_report = _one_report(on_dir)
    on_wb = openpyxl.load_workbook(on_report)
    assert on_wb.sheetnames == ALL_SHEETS
    on_validate = _validate(on_dir)
    assert "[PASS] 13个Sheet齐全且顺序正确" in on_validate

    # ---- 关闭态：11 个 Sheet、顺序正确、不含 08/09、验收通过 ----
    off_report = _one_report(off_dir)
    off_wb = openpyxl.load_workbook(off_report)
    assert off_wb.sheetnames == OFF_SHEETS
    assert MATRIX_SHEETS[0] not in off_wb.sheetnames
    assert MATRIX_SHEETS[1] not in off_wb.sheetnames
    off_validate = _validate(off_dir)
    assert "[PASS] 11个Sheet齐全且顺序正确" in off_validate
    assert "[PASS] 未勾选矩阵时没有任何traffic_keyword调用" in off_validate
    # 06 散点辅助区与看板 4 图必须在两种模式下都成立（矩阵开关不改变关键词面）
    for output in (on_validate, off_validate):
        assert "[PASS] 散点图全部引用06集中度辅助区" in output
        assert "[PASS] 看板图表数量与类型符合有效数据口径" in output
        assert "[PASS] 数据 Sheet 无图表" in output

    # ---- 关闭态：06 机会矩阵区只保留关键词行，与开启态逐格一致 ----
    on_region = _opportunity_region(on_wb["06_集中度与机会矩阵"])
    off_region = _opportunity_region(off_wb["06_集中度与机会矩阵"])
    assert on_region and off_region == on_region
    off06_text = _sheet_text(off_wb["06_集中度与机会矩阵"])
    assert not any(asin in off06_text for asin in HIST_ASINS)

    # ---- 口径文案：如实说明未勾选矩阵，且不许与「纯Node模式」混用 ----
    assert _labeled_value(off_wb["11_调研范围与调用"], "ASIN矩阵口径") == "本次未勾选矩阵，未调用 traffic_keyword"
    assert _labeled_value(off_wb["00_市场看板"], "ASIN矩阵口径") == "本次未勾选矩阵，未调用 traffic_keyword"
    on_quality = _sheet_text(on_wb["12_数据质量"])
    off_quality = _sheet_text(off_wb["12_数据质量"])
    assert "当前前100口径" in on_quality and "纯Node模式" not in on_quality
    assert "本次未勾选矩阵" in off_quality and "未调用 traffic_keyword" in off_quality
    assert "纯Node模式" not in off_quality
    assert "当前前100流量词匹配" not in _sheet_text(off_wb["11_调研范围与调用"])
    assert "当前前100流量词匹配" not in _sheet_text(off_wb["00_市场看板"])

    # ---- 关闭态：调用记录里没有 traffic_keyword（整条链路真的跳过） ----
    assert "traffic_keyword" not in off_mcp.calls
    off_log, off_tools = _call_summary(off_dir)
    assert "traffic_keyword" not in off_tools
    assert off_tools.count("traffic_extend") == 3  # 反查不受矩阵开关影响
    assert not list((off_dir / "raw").glob("traffic_keyword_current_*.json"))
    assert "本次未勾选矩阵" in "\n".join(off_store.lines)

    # ---- 开启态：逐 ASIN 各一次 traffic_keyword（证明断言不是空跑） ----
    assert on_mcp.calls.count("traffic_keyword") == len(HIST_ASINS)
    on_log, on_tools = _call_summary(on_dir)
    assert on_tools.count("traffic_keyword") == len(HIST_ASINS)
    assert len(list((on_dir / "raw").glob("traffic_keyword_current_*.json"))) == len(HIST_ASINS)

    # ---- 预算：关闭态从计划里扣掉矩阵 ASIN 数，计划=实际 ----
    on_plan = json.loads((on_dir / "plan.json").read_text(encoding="utf-8"))
    off_plan = json.loads((off_dir / "plan.json").read_text(encoding="utf-8"))
    assert on_plan["planned_total"] - off_plan["planned_total"] == len(HIST_ASINS)
    assert off_log["actual_total"] == off_plan["planned_total"]
    assert on_log["actual_total"] == on_plan["planned_total"]


def test_matrix_flag_missing_defaults_to_enabled(tmp_path):
    """字段缺失（桌面版输入与历史夹具都缺）等同勾选：13 Sheet、文案与预算都不变。"""
    groups = [{"marketplace": "US", "asins": HIST_ASINS, "node_ids": [], "node_top_n": 50}]
    options = {"data_month": "202606", "keyword_pool_pages": 3, "core_trend_keywords": 5,
               "longtail_trend_keywords": 5, "trend_keyword_cap": 20}
    missing = plan_budget(groups, options)
    explicit = plan_budget(groups, {**options, "matrix_enabled": True})
    assert (missing.planned_total, missing.groups[0].matrix_calls) == \
        (explicit.planned_total, explicit.groups[0].matrix_calls)

    task_dir, mcp, _ = _run(tmp_path / "legacy", None)
    assert "matrix_enabled" not in json.loads(
        (task_dir / "input.json").read_text(encoding="utf-8"))["options"]
    wb = openpyxl.load_workbook(_one_report(task_dir))
    assert wb.sheetnames == ALL_SHEETS
    assert _labeled_value(wb["11_调研范围与调用"], "ASIN矩阵口径") == "当前前100流量词匹配"
    assert mcp.calls.count("traffic_keyword") == len(HIST_ASINS)
    assert "[PASS] 13个Sheet齐全且顺序正确" in _validate(task_dir)


def test_plan_budget_excludes_matrix_calls_when_disabled():
    """预算数字的真相源是 plan_budget：关闭态 matrix_calls=0 且 total 少掉 A 项。"""
    for groups in (
        [{"marketplace": "US", "asins": HIST_ASINS, "node_ids": [], "node_top_n": 50}],
        [{"marketplace": "US", "asins": [], "node_ids": ["1", "2"], "node_top_n": 50}],
        [{"marketplace": "US", "asins": [_synthetic_asin(1)], "node_ids": ["1"], "node_top_n": 50}],
    ):
        asins = len(groups[0]["asins"])
        options = {"data_month": "202606", "keyword_pool_pages": 3}
        on = plan_budget(groups, {**options, "matrix_enabled": True})
        off = plan_budget(groups, {**options, "matrix_enabled": False})
        assert on.groups[0].matrix_calls == asins
        assert off.groups[0].matrix_calls == 0
        assert on.planned_total - off.planned_total == asins
        # 反查与榜单取样不受开关影响
        assert off.groups[0].pool_batches == on.groups[0].pool_batches
        assert off.groups[0].node_pages == on.groups[0].node_pages
        assert off.groups[0].extra_m1 == on.groups[0].extra_m1
